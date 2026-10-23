"""Deterministic, non-AI contract-template platform."""
from __future__ import annotations
import copy,hashlib,re,zipfile,uuid
from abc import ABC,abstractmethod
from datetime import date
from decimal import Decimal,InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any,Callable
from xml.etree import ElementTree as ET
from fastapi import APIRouter,Depends,File,HTTPException,UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel,Field
from sqlalchemy.orm import Session as DBS
from docx import Document
from app.conditions import ConditionConfigurationError,ConditionEngine
from app.calculations import CalculationEngine,CalculationError,ValueFormatter
from app.db import ConditionalDocumentRegion,ContractProductVariant,ContractProductVersion,DocumentBinding,DocumentDraft,DocumentTemplate,GeneratedDocument,QuestionnaireField,QuestionnaireSchema,RepeatingGroup,RepeatingGroupField,SessionLocal,TemplateAuditEvent,TemplateVersion,now,settings

DOCX_MIME='application/vnd.openxmlformats-officedocument.wordprocessingml.document'; DOCX_MIMES={DOCX_MIME,'application/octet-stream','application/zip'}
WORD_NS='http://schemas.openxmlformats.org/wordprocessingml/2006/main'; NS={'w':WORD_NS}; FIELD_TYPES={'TEXT','MULTILINE_TEXT','INTEGER','DECIMAL','MONEY','DATE','BOOLEAN','SINGLE_SELECT','LIST','COMPOSITE'}
REGION_TYPES={'CONDITIONAL_RUN','CONDITIONAL_PARAGRAPH','CONDITIONAL_PARAGRAPH_GROUP','CONDITIONAL_TABLE_ROW','CONDITIONAL_TABLE','CONDITIONAL_TABLE_COLUMN'}

def db():
 s=SessionLocal()
 try: yield s
 finally: s.close()
def roles(u): return set(u.roles or [])
def require_manager(u):
 if not roles(u)&{'dpo_specialist','administrator'}: raise HTTPException(403,detail='Управление шаблонами доступно специалистам ДПО')
def require_publisher(u):
 if not roles(u)&{'template_publisher','administrator'}: raise HTTPException(403,detail='Нет разрешения на публикацию шаблонов')
def require_draft(v):
 if v.status!='DRAFT': raise HTTPException(409,detail='Опубликованная версия неизменяема. Загрузите новую версию DOCX.')
def owned_template(id,u,s):
 x=s.get(DocumentTemplate,id)
 if not x or (x.owner_id!=u.id and 'administrator' not in roles(u)): raise HTTPException(404,detail='Шаблон не найден')
 return x
def owned_version(id,u,s):
 v=s.get(TemplateVersion,id)
 if not v: raise HTTPException(404,detail='Версия не найдена')
 return v,owned_template(v.template_id,u,s)
def own_draft(id,u,s):
 x=s.get(DocumentDraft,id)
 if not x or x.owner_id!=u.id: raise HTTPException(404,detail='Черновик не найден')
 return x
def safe_file(value,root_name,checksum):
 root=(Path(settings.storage)/root_name).resolve(); path=Path(value).resolve()
 if not path.is_relative_to(root) or not path.is_file(): raise HTTPException(404,detail='Файл недоступен')
 if hashlib.sha256(path.read_bytes()).hexdigest()!=checksum: raise HTTPException(409,detail='Нарушена целостность файла')
 return path
def safe_source(v): return safe_file(v.source_path,'template_sources',v.source_sha256)
def validate_docx(filename,mime,raw):
 if Path(filename).suffix.lower()!='.docx': raise HTTPException(415,detail='Поддерживается только DOCX')
 if mime and mime.lower() not in DOCX_MIMES: raise HTTPException(415,detail='Некорректный MIME-тип DOCX')
 if not raw: raise HTTPException(422,detail='Файл пустой')
 if len(raw)>settings.max_upload: raise HTTPException(413,detail='Файл превышает допустимый размер')
 try:
  with zipfile.ZipFile(BytesIO(raw)) as z:
   infos=z.infolist(); names={x.filename for x in infos}
   if len(infos)>5000 or sum(x.file_size for x in infos)>200*1024*1024: raise HTTPException(422,detail='DOCX содержит слишком много данных')
   if any(Path(x.filename).is_absolute() or '..' in Path(x.filename).parts for x in infos): raise HTTPException(422,detail='DOCX содержит небезопасные пути')
   if '[Content_Types].xml' not in names or 'word/document.xml' not in names or z.testzip() is not None: raise HTTPException(422,detail='Файл не имеет корректной структуры DOCX')
   ET.fromstring(z.read('word/document.xml'))
 except HTTPException: raise
 except (zipfile.BadZipFile,KeyError,ET.ParseError): raise HTTPException(422,detail='Файл не имеет корректной структуры DOCX')

class DocumentEditorAdapter(ABC):
 @abstractmethod
 def preview(self,path:Path)->dict[str,Any]: ...
 @abstractmethod
 def select_range(self,path:Path,locator:str,start:int,end:int,text:str)->dict[str,Any]: ...
 @abstractmethod
 def get_structural_locator(self,path:Path,selection:dict[str,Any])->dict[str,Any]: ...
class DocxXmlAdapter(DocumentEditorAdapter):
 def _parts(self,z): return [(x,ET.fromstring(z.read(x))) for x in ['word/document.xml']+sorted(n for n in z.namelist() if re.fullmatch(r'word/(?:header|footer)\d+\.xml',n))]
 def _runs(self,p): return [''.join(x.text or '' for x in r.findall('.//w:t',NS)) for r in p.findall('./w:r',NS)]
 def _node(self,part,locator,kind,text,runs,meta):
  result=[]; pos=0
  for i,value in enumerate(runs,1): result.append({'run':i,'start':pos,'end':pos+len(value)}); pos+=len(value)
  return {'part':part,'locator':locator,'kind':kind,'text':text,'runs':result,'metadata':meta}
 def preview(self,path):
  nodes=[]
  with zipfile.ZipFile(path) as z:
   for part,root in self._parts(z):
    body=root.find('w:body',NS) if part=='word/document.xml' else root
    if body is None: continue
    pno=tno=0
    for child in list(body):
     if child.tag==f'{{{WORD_NS}}}p':
      pno+=1; runs=self._runs(child); text=''.join(runs)
      if text.strip(): nodes.append(self._node(part,f'{part}#/p[{pno}]','paragraph',text,runs,{}))
     elif child.tag==f'{{{WORD_NS}}}tbl':
      tno+=1
      for rno,row in enumerate(child.findall('./w:tr',NS),1):
       for cno,cell in enumerate(row.findall('./w:tc',NS),1):
        props=cell.find('./w:tcPr',NS); grid=props.find('./w:gridSpan',NS) if props is not None else None; merge=props.find('./w:vMerge',NS) if props is not None else None
        meta={'table':tno,'row':rno,'cell':cno,'grid_span':int(grid.get(f'{{{WORD_NS}}}val','1')) if grid is not None else 1,'vertical_merge':merge.get(f'{{{WORD_NS}}}val','continue') if merge is not None else None}
        for cpno,p in enumerate(cell.findall('./w:p',NS),1):
         runs=self._runs(p); text=''.join(runs)
         # Empty table cells are selectable repeat-field targets. They are excluded
         # from the legacy fingerprint so existing uploaded versions remain valid.
         nodes.append(self._node(part,f'{part}#/tbl[{tno}]/tr[{rno}]/tc[{cno}]/p[{cpno}]','table_cell',text,runs,meta))
  fp=hashlib.sha256('\n'.join(x['locator']+'\0'+x['text'] for x in nodes if x['text'].strip()).encode()).hexdigest()
  return {'adapter':'DOCX_XML_READ_ONLY','locator_version':1,'structure_fingerprint':fp,'nodes':nodes,'limitations':['Разметка страницы отображается приблизительно.','Выделение поддерживается только внутри одного абзаца или ячейки.']}
 def select_range(self,path,locator,start,end,text):
  preview=self.preview(path); node=next((x for x in preview['nodes'] if x['locator']==locator),None)
  if not node: raise ValueError('Структурный фрагмент не найден в исходном DOCX.')
  empty_cell=node['kind']=='table_cell' and node['text']=='' and start==end==0 and text==''
  if not empty_cell and (start<0 or end<=start or end>len(node['text']) or node['text'][start:end]!=text): raise ValueError('Выделение не соответствует исходному DOCX.')
  return {'preview':preview,'node':node,'start':start,'end':end}
 def get_structural_locator(self,path,sel):
  preview,node,start,end=sel['preview'],sel['node'],sel['start'],sel['end']; touched=[x for x in node['runs'] if x['end']>start and x['start']<end]
  locator={'adapter':preview['adapter'],'locator_version':1,'structure_fingerprint':preview['structure_fingerprint'],'part':node['part'],'node':node['locator'],'text_start':start,'text_end':end}
  if touched: locator.update({'start_run':touched[0]['run'],'start_run_offset':start-touched[0]['start'],'end_run':touched[-1]['run'],'end_run_offset':end-touched[-1]['start']})
  else: locator['empty_cell']=True
  return locator
adapter=DocxXmlAdapter()

def template_out(x,s): return {'id':x.id,'code':x.code,'name':x.name,'category':x.category,'description':x.description,'status':x.status,'owner_id':x.owner_id,'current_version_id':x.current_version_id,'version_count':s.query(TemplateVersion).filter_by(template_id=x.id).count()}
def version_out(x): return {'id':x.id,'template_id':x.template_id,'version_number':x.version_number,'status':x.status,'original_filename':x.original_filename,'source_sha256':x.source_sha256,'structure_fingerprint':x.structure_fingerprint,'preview_metadata':x.preview_metadata,'created_at':x.created_at,'published_at':x.published_at}
def field_out(x): return {'id':x.id,'schema_id':x.schema_id,'field_key':x.field_key,'label':x.label,'field_type':x.field_type,'required':x.required,'position':x.position,'config':x.config}
def binding_out(x): return {'id':x.id,'field_id':x.field_id,'binding_key':x.binding_key,'occurrence_index':x.occurrence_index,'selected_text':x.selected_text,'locator':x.locator,'format_config':x.format_config or {},'status':x.status}
def region_out(x): return {'id':x.id,'region_key':x.region_key,'region_type':x.region_type,'locator':x.locator,'include_when':x.include_when,'original_text':x.original_text,'status':x.status}
def repeating_field_out(x): return {'id':x.id,'group_id':x.group_id,'field_key':x.field_key,'label':x.label,'field_type':x.field_type,'required':x.required,'position':x.position,'config':x.config,'locator':x.locator}
def repeating_group_out(x,s): return {'id':x.id,'group_key':x.group_key,'name':x.name,'description':x.description,'min_items':x.min_items,'max_items':x.max_items,'required':x.required,'position':x.position,'config':x.config,'locator':x.locator,'status':x.status,'fields':[repeating_field_out(f) for f in s.query(RepeatingGroupField).filter_by(group_id=x.id).order_by(RepeatingGroupField.position).all()]}
class CreateTemplate(BaseModel):
 code:str=Field(min_length=2,max_length=80,pattern=r'^[A-Za-z0-9_-]+$'); name:str=Field(min_length=2,max_length=240); category:str=Field(default='Другое',min_length=2,max_length=120); description:str=Field(default='',max_length=2000)
class FieldPayload(BaseModel):
 field_key:str=Field(min_length=2,max_length=120,pattern=r'^[A-Za-z][A-Za-z0-9_.-]*$'); label:str=Field(min_length=2,max_length=300); field_type:str; required:bool=False; position:int=Field(ge=1); description:str=Field(default='',max_length=1000); group:str=Field(default='Основные сведения',max_length=160); placeholder:str=Field(default='',max_length=300); validation:dict[str,Any]=Field(default_factory=dict); options:list[str]=Field(default_factory=list,max_length=100); visible_when:dict[str,Any]|None=None; required_when:dict[str,Any]|None=None; calculated:bool=False; expression:dict[str,Any]|None=None; text_expression:dict[str,Any]|None=None; precision:int=Field(default=2,ge=0,le=12); rounding:str='HALF_UP'; output_format:str='RAW'; item_config:dict[str,Any]=Field(default_factory=dict); components:list[dict[str,Any]]=Field(default_factory=list,max_length=30); separator:str=Field(default='; ',max_length=20)
class BindingPayload(BaseModel): field_id:str|None=None; binding_key:str|None=None; locator:str; start:int=Field(ge=0); end:int=Field(ge=0); selected_text:str=Field(default='',max_length=10000); output_format:str='RAW'; format_options:dict[str,Any]=Field(default_factory=dict)
class BindingFormatPayload(BaseModel): output_format:str='RAW'; format_options:dict[str,Any]=Field(default_factory=dict)
class RegionPayload(BaseModel):
 region_key:str=Field(min_length=2,max_length=120,pattern=r'^[A-Za-z][A-Za-z0-9_.-]*$'); region_type:str; start_locator:str; end_locator:str|None=None; start:int|None=Field(default=None,ge=0); end:int|None=Field(default=None,gt=0); selected_text:str|None=Field(default=None,max_length=10000); include_when:dict[str,Any]
class RegionUpdatePayload(BaseModel):
 region_key:str=Field(min_length=2,max_length=120,pattern=r'^[A-Za-z][A-Za-z0-9_.-]*$'); include_when:dict[str,Any]
class RepeatingGroupPayload(BaseModel):
 group_key:str=Field(min_length=2,max_length=120,pattern=r'^[A-Za-z][A-Za-z0-9_.-]*$'); name:str=Field(min_length=2,max_length=240); description:str=Field(default='',max_length=2000); min_items:int=Field(default=0,ge=0,le=50); max_items:int=Field(default=10,ge=1,le=50); required:bool=False; position:int=Field(ge=1); row_locator:str; visible_when:dict[str,Any]|None=None
class RepeatingGroupUpdatePayload(BaseModel):
 group_key:str=Field(min_length=2,max_length=120,pattern=r'^[A-Za-z][A-Za-z0-9_.-]*$'); name:str=Field(min_length=2,max_length=240); description:str=Field(default='',max_length=2000); min_items:int=Field(default=0,ge=0,le=50); max_items:int=Field(default=10,ge=1,le=50); required:bool=False; position:int=Field(ge=1); visible_when:dict[str,Any]|None=None
class RepeatingFieldPayload(BaseModel):
 field_key:str=Field(min_length=2,max_length=120,pattern=r'^[A-Za-z][A-Za-z0-9_.-]*$'); label:str=Field(min_length=2,max_length=300); field_type:str; required:bool=False; position:int=Field(ge=1); description:str=Field(default='',max_length=1000); placeholder:str=Field(default='',max_length=300); validation:dict[str,Any]=Field(default_factory=dict); options:list[str]=Field(default_factory=list,max_length=100); calculated:bool=False; expression:dict[str,Any]|None=None; precision:int=Field(default=2,ge=0,le=12); rounding:str='HALF_UP'; output_format:str='RAW'
class RepeatingBindingPayload(BaseModel): locator:str; start:int=Field(ge=0); end:int=Field(ge=0); selected_text:str=Field(default='',max_length=10000)
class AnswersPayload(BaseModel): answers:dict[str,Any]
class CandidateDecisionPayload(BaseModel):
 decision:str=Field(pattern=r'^(CONFIRMED|REJECTED)$')
class CandidateBatchPayload(BaseModel):
 candidate_ids:list[str]=Field(min_length=1,max_length=100)

def validate_field_payload(p):
 if p.field_type not in FIELD_TYPES: raise HTTPException(422,detail='Неподдерживаемый тип поля')
 if set(p.validation)-{'min_length','max_length','min','max','regex'}: raise HTTPException(422,detail='Параметры проверки содержат неподдерживаемые правила')
 if 'regex' in p.validation:
  pattern=str(p.validation['regex'])
  if len(pattern)>200 or '(?' in pattern or re.search(r'\\[1-9]',pattern): raise HTTPException(422,detail='Регулярное выражение слишком сложное')
  try: re.compile(pattern)
  except re.error: raise HTTPException(422,detail='Некорректное регулярное выражение')
 if p.field_type=='SINGLE_SELECT' and (not p.options or any(not str(x).strip() for x in p.options)): raise HTTPException(422,detail='Добавьте варианты выбора')
 if p.calculated and p.field_type not in {'INTEGER','DECIMAL','MONEY','TEXT'}: raise HTTPException(422,detail='Этот тип поля нельзя вычислять')
 if p.calculated and p.field_type=='TEXT' and not p.text_expression: raise HTTPException(422,detail='Текстовому вычисляемому полю нужно выражение')
 if p.calculated and p.field_type!='TEXT' and not p.expression: raise HTTPException(422,detail='Вычисляемому полю нужна числовая формула')
 if p.field_type=='LIST':
  if set(p.item_config)-{'min_items','max_items','min_length','max_length','regex'}:raise HTTPException(422,detail='Некорректные настройки списка')
  if not 0<=int(p.item_config.get('min_items',0))<=int(p.item_config.get('max_items',20))<=100:raise HTTPException(422,detail='Пределы списка некорректны')
 if p.field_type=='COMPOSITE':
  _validate_components(p.components)
 try: ValueFormatter.validate(p.output_format)
 except CalculationError as exc: raise HTTPException(422,detail=str(exc))
def calculation_config(p): return {'calculated':p.calculated,'expression':p.expression,'precision':p.precision,'rounding':p.rounding,'output_format':p.output_format}
def field_config(p): return {'description':p.description,'group':p.group,'placeholder':p.placeholder,'validation':p.validation,'options':p.options,'visible_when':p.visible_when,'required_when':p.required_when,'text_expression':p.text_expression,'item_config':p.item_config,'components':p.components,'separator':p.separator,**calculation_config(p)}
def repeating_field_config(p): return {'description':p.description,'placeholder':p.placeholder,'validation':p.validation,'options':p.options,**calculation_config(p)}
def _validate_components(components,depth=1):
 if depth>4:raise HTTPException(422,detail='Вложенность составного поля не может превышать 4 уровня')
 keys=[x.get('key') for x in components]
 if not keys or len(keys)!=len(set(keys)) or any(not re.fullmatch(r'[A-Za-z][A-Za-z0-9_.-]*',str(x)) for x in keys):raise HTTPException(422,detail='Компоненты должны иметь уникальные корректные ключи')
 allowed={'key','label','type','required','validation','components','separator'}
 if any(set(x)-allowed for x in components):raise HTTPException(422,detail='Компонент содержит неподдерживаемые настройки')
 for item in components:
  if item.get('type','TEXT') not in {'TEXT','DATE','INTEGER','DECIMAL','BOOLEAN','COMPOSITE'}:raise HTTPException(422,detail='Неподдерживаемый тип компонента')
  if item.get('type')=='COMPOSITE':_validate_components(item.get('components') or [],depth+1)
def _validate_composite(components,value):
 if not isinstance(value,dict):return 'Ожидалось составное значение'
 by_key={x['key']:x for x in components}
 if set(value)-set(by_key):return 'Составное значение содержит неизвестные компоненты'
 for key,item in by_key.items():
  current=value.get(key)
  if item.get('required') and current in (None,''):return f'Заполните «{item.get("label",key)}»'
  if current not in (None,'') and item.get('type')=='COMPOSITE':
   if error:=_validate_composite(item.get('components') or [],current):return f'«{item.get("label",key)}»: {error.lower()}'
 return None
def _render_composite(components,value,separator=', '):
 out=[]
 for item in components:
  current=value.get(item['key']) if isinstance(value,dict) else None
  if current in (None,''):continue
  out.append(_render_composite(item.get('components') or [],current,item.get('separator',', ')) if item.get('type')=='COMPOSITE' else ('Да' if current is True else 'Нет' if current is False else str(current)))
 return separator.join(out)
def validate_repeating_field_payload(p):
 if p.field_type not in FIELD_TYPES: raise HTTPException(422,detail='Неподдерживаемый тип поля строки')
 if set(p.validation)-{'min_length','max_length','min','max','regex'}: raise HTTPException(422,detail='Параметры проверки содержат неподдерживаемые правила')
 if p.field_type=='SINGLE_SELECT' and (not p.options or any(not str(x).strip() for x in p.options)): raise HTTPException(422,detail='Добавьте варианты выбора')
 if p.calculated and (p.field_type not in {'INTEGER','DECIMAL','MONEY'} or not p.expression): raise HTTPException(422,detail='Вычисляемому полю строки нужна числовая формула')
 try: ValueFormatter.validate(p.output_format)
 except CalculationError as exc: raise HTTPException(422,detail=str(exc))
 if 'regex' in p.validation:
  pattern=str(p.validation['regex'])
  if len(pattern)>200 or '(?' in pattern or re.search(r'\\[1-9]',pattern): raise HTTPException(422,detail='Регулярное выражение слишком сложное')
  try: re.compile(pattern)
  except re.error: raise HTTPException(422,detail='Некорректное регулярное выражение')
def validate_condition_schema(fields):
 try: ConditionEngine.validate_schema(fields)
 except ConditionConfigurationError as exc: raise HTTPException(422,detail=f'Некорректные условия анкеты: {exc}')
def validate_answer(field,value):
 if value in (None,''):
  return 'Обязательное поле' if field.required else None
 cfg=field.config or {}; rules=cfg.get('validation') or {}; kind=field.field_type
 try:
  if kind in {'TEXT','MULTILINE_TEXT'}:
   text=str(value)
   if rules.get('min_length') is not None and len(text)<int(rules['min_length']): return f'Минимум {rules["min_length"]} символов'
   if rules.get('max_length') is not None and len(text)>int(rules['max_length']): return f'Максимум {rules["max_length"]} символов'
   if rules.get('regex') and not re.fullmatch(str(rules['regex']),text): return 'Значение не соответствует требуемому формату'
  elif kind=='INTEGER': number=Decimal(int(value))
  elif kind in {'DECIMAL','MONEY'}: number=Decimal(str(value).replace(',','.'))
  elif kind=='DATE': date.fromisoformat(str(value))
  elif kind=='BOOLEAN' and not isinstance(value,bool): return 'Выберите да или нет'
  elif kind=='SINGLE_SELECT' and value not in cfg.get('options',[]): return 'Выберите значение из списка'
  elif kind=='LIST':
   if not isinstance(value,list) or any(not isinstance(x,str) for x in value):return 'Ожидался список строк'
   item=cfg.get('item_config') or {};minimum=int(item.get('min_items',0));maximum=int(item.get('max_items',20))
   if len(value)<minimum or len(value)>maximum:return f'Количество элементов должно быть от {minimum} до {maximum}'
   if any(not x.strip() or len(x)>int(item.get('max_length',500)) for x in value):return 'Проверьте элементы списка'
   if item.get('regex') and any(not re.fullmatch(str(item['regex']),x) for x in value):return 'Элемент списка не соответствует формату'
  elif kind=='COMPOSITE':
   if error:=_validate_composite(cfg.get('components',[]),value):return error
  if kind in {'INTEGER','DECIMAL','MONEY'}:
   if rules.get('min') is not None and number<Decimal(str(rules['min'])): return f'Значение не меньше {rules["min"]}'
   if rules.get('max') is not None and number>Decimal(str(rules['max'])): return f'Значение не больше {rules["max"]}'
 except (ValueError,TypeError,InvalidOperation): return 'Некорректный формат значения'
 return None
def render_value(field,value):
 if value in (None,''): return ''
 if field.field_type=='BOOLEAN': return 'Да' if value else 'Нет'
 if field.field_type=='LIST': return (field.config or {}).get('separator','; ').join(str(x) for x in value)
 if field.field_type=='COMPOSITE':
  cfg=field.config or {};return _render_composite(cfg.get('components',[]),value,cfg.get('separator',', '))
 return str(value)
def schema_for(version_id,s,create=False,u=None):
 x=s.query(QuestionnaireSchema).filter_by(template_version_id=version_id).order_by(QuestionnaireSchema.schema_version.desc()).first()
 if not x and create: x=QuestionnaireSchema(template_version_id=version_id,created_by_id=u.id); s.add(x); s.flush()
 return x
def schema_out(x,s): return {'id':x.id,'title':x.title,'status':x.status,'condition_schema_version':x.condition_schema_version,'calculation_schema_version':x.calculation_schema_version,'fields':[field_out(f) for f in s.query(QuestionnaireField).filter_by(schema_id=x.id).order_by(QuestionnaireField.position).all()],'repeating_groups':[repeating_group_out(g,s) for g in s.query(RepeatingGroup).filter_by(template_version_id=x.template_version_id,status='ACTIVE').order_by(RepeatingGroup.position).all()]}
def calculation_groups(version_id,s):
 groups=s.query(RepeatingGroup).filter_by(template_version_id=version_id,status='ACTIVE').order_by(RepeatingGroup.position).all()
 for group in groups: group.fields=s.query(RepeatingGroupField).filter_by(group_id=group.id).order_by(RepeatingGroupField.position).all()
 return groups
def validate_calculations(version_id,s):
 schema=schema_for(version_id,s); fields=s.query(QuestionnaireField).filter_by(schema_id=schema.id).all() if schema else []
 try: CalculationEngine.validate(fields,calculation_groups(version_id,s))
 except CalculationError as exc: raise HTTPException(422,detail=f'Некорректные вычисления: {exc}')
def calculated_document_values(version_id,s,fields,answers):
 groups=calculation_groups(version_id,s); active_keys={x.field_key for x in ConditionEngine.active_fields(fields,answers)}; active_groups={x.group_key for x in groups if group_is_active(x,fields,answers)}
 try:
  CalculationEngine.validate(fields,groups)
  return groups,*CalculationEngine.calculate(fields,groups,answers,active_keys,active_groups)
 except CalculationError as exc: raise HTTPException(422,detail=f'Ошибка вычисления: {exc}')
def formatted_bindings(bindings,values,fields=None):
 by_id={x.id:x for x in (fields or [])}
 result={}
 for binding in bindings:
  cfg=binding.format_config or {}; value=values.get(binding.binding_key,''); field=by_id.get(binding.field_id)
  result[binding.id]=render_value(field,value) if field and field.field_type=='BOOLEAN' and cfg.get('format','RAW')=='RAW' else ValueFormatter.format(value,cfg.get('format','RAW'),cfg.get('options') or {})
 return result
def validate_active_answers(fields,answers):
 errors={}
 for field in fields:
  if (field.config or {}).get('calculated'): continue
  state=ConditionEngine.field_state(field,answers)
  if not state['visible']: continue
  value=answers.get(field.field_key)
  if state['required'] and value in (None,''): errors[field.field_key]='Обязательное поле'
  elif value not in (None,'') and (error:=validate_answer(field,value)): errors[field.field_key]=error
 return errors
def condition_answers(fields,answers):
 """Hidden answers stay persisted, but are absent for legal-block decisions."""
 visible={x.field_key for x in ConditionEngine.active_fields(fields,answers)}
 return {key:value for key,value in answers.items() if key in visible}
def group_is_active(group,questionnaire_fields,answers):
 condition=(group.config or {}).get('visible_when')
 return ConditionEngine.evaluate(condition,condition_answers(questionnaire_fields,answers))
def validate_repeating_answers(groups,s,questionnaire_fields,answers,complete=True):
 errors={}
 for group in groups:
  items=answers.get(group.group_key,[]); active=group_is_active(group,questionnaire_fields,answers)
  if not isinstance(items,list): errors[group.group_key]=f'«{group.name}»: данные должны быть списком записей'; continue
  if len(items)>group.max_items: errors[group.group_key]=f'«{group.name}»: можно добавить не больше {group.max_items} записей'; continue
  minimum=max(group.min_items,1 if group.required else 0)
  if complete and active and len(items)<minimum: errors[group.group_key]=f'«{group.name}»: добавьте минимум {minimum} записей'; continue
  fields=s.query(RepeatingGroupField).filter_by(group_id=group.id).order_by(RepeatingGroupField.position).all(); known={x.field_key:x for x in fields}; ids=set()
  for index,item in enumerate(items,1):
   if not isinstance(item,dict): errors[group.group_key]=f'«{group.name}», запись № {index}: повреждённый формат'; break
   item_id=item.get('_id')
   if not isinstance(item_id,str) or not item_id or len(item_id)>100 or item_id in ids: errors[group.group_key]=f'«{group.name}», запись № {index}: не удалось определить запись'; break
   ids.add(item_id)
   if set(item)-set(known)-{'_id'}: errors[group.group_key]=f'«{group.name}», запись № {index}: найдены неизвестные поля'; break
   if any(k in item and (f.config or {}).get('calculated') for k,f in known.items()): errors[group.group_key]=f'«{group.name}», запись № {index}: вычисляемые поля нельзя изменять'; break
   if not active: continue
   for field in fields:
    if (field.config or {}).get('calculated'): continue
    value=item.get(field.field_key)
    if complete and field.required and value in (None,''): errors[group.group_key]=f'«{group.name}», запись № {index}: заполните поле «{field.label}»'; break
    if value not in (None,'') and (message:=validate_answer(field,value)): errors[group.group_key]=f'«{group.name}», запись № {index}, поле «{field.label}»: {message.lower()}'; break
   if group.group_key in errors: break
 return errors
def validate_bindings(v,s):
 preview=adapter.preview(safe_source(v))
 if preview['structure_fingerprint']!=v.structure_fingerprint: raise HTTPException(409,detail='DOCX изменился после создания версии')
 nodes={x['locator']:x for x in preview['nodes']}; occupied={}; bindings=s.query(DocumentBinding).filter_by(template_version_id=v.id,status='ACTIVE').all()
 for b in bindings:
  loc=b.locator or {}; node=nodes.get(loc.get('node')); start,end=loc.get('text_start'),loc.get('text_end')
  if not node or loc.get('part')!=node['part'] or loc.get('structure_fingerprint')!=v.structure_fingerprint: raise HTTPException(422,detail=f'Привязка «{b.binding_key}» устарела')
  if not isinstance(start,int) or not isinstance(end,int) or start<0 or end>len(node['text']) or node['text'][start:end]!=b.selected_text: raise HTTPException(422,detail=f'Диапазон привязки «{b.binding_key}» повреждён')
  ranges=occupied.setdefault(node['locator'],[])
  if any(start<e and end>s0 for s0,e in ranges): raise HTTPException(422,detail='Диапазоны привязок пересекаются')
  ranges.append((start,end))
 return bindings
def _element_hash(node): return hashlib.sha256(ET.tostring(node,encoding='utf-8')).hexdigest()
def _part_root(path,part):
 with zipfile.ZipFile(path) as z: return ET.fromstring(z.read(part))
def _body(root): return root.find('w:body',NS) if root.tag.endswith('document') else root
def _unsafe(node):
 names=('fldChar','instrText','drawing','pict','object','bookmarkStart','bookmarkEnd')
 return next((name for name in names if node.find(f'.//w:{name}',NS) is not None),None)
def _table_parts(locator):
 match=re.fullmatch(r'(.+)#/tbl\[(\d+)\]/tr\[(\d+)\]/tc\[(\d+)\]/p\[(\d+)\]',locator)
 return (match.group(1),*map(int,match.groups()[1:])) if match else None
def _top_paragraph(locator):
 match=re.fullmatch(r'(.+)#/p\[(\d+)\]',locator)
 return (match.group(1),int(match.group(2))) if match else None
def _paragraph_text(node): return ''.join(x.text or '' for x in node.findall('.//w:t',NS))
def _span(cell):
 node=cell.find('./w:tcPr/w:gridSpan',NS);return int(node.get(f'{{{WORD_NS}}}val','1')) if node is not None else 1
def _column_layout(table,column):
 grid=table.find('./w:tblGrid',NS);columns=grid.findall('./w:gridCol',NS) if grid is not None else []
 if not columns or column<1 or column>len(columns):raise ValueError('В таблице отсутствует корректный tblGrid')
 for row in table.findall('./w:tr',NS):
  cursor=1;hit=None
  for cell in row.findall('./w:tc',NS):
   span=_span(cell)
   if cursor<=column<cursor+span:hit=(cell,cursor,span);break
   cursor+=span
  if hit is None:raise ValueError('Строка не согласована с tblGrid')
  if hit[0].find('./w:tcPr/w:vMerge',NS) is not None:raise ValueError('Колонка пересекает вертикально объединённые ячейки')
 return columns
def _remove_table_column(table,column):
 columns=_column_layout(table,column);removed=columns[column-1];width=int(removed.get(f'{{{WORD_NS}}}w','0') or 0);neighbor=columns[column-2] if column>1 else (columns[1] if len(columns)>1 else None)
 if neighbor is not None and width:neighbor.set(f'{{{WORD_NS}}}w',str(int(neighbor.get(f'{{{WORD_NS}}}w','0') or 0)+width))
 table.find('./w:tblGrid',NS).remove(removed)
 for row in table.findall('./w:tr',NS):
  cursor=1
  for cell in row.findall('./w:tc',NS):
   span=_span(cell)
   if cursor<=column<cursor+span:
    if span==1:row.remove(cell)
    else:
     node=cell.find('./w:tcPr/w:gridSpan',NS)
     if span==2:cell.find('./w:tcPr',NS).remove(node)
     else:node.set(f'{{{WORD_NS}}}val',str(span-1))
    break
   cursor+=span
def _resolve_region(path,p:RegionPayload,structure_fingerprint):
 if p.region_type not in REGION_TYPES: raise ValueError('Неподдерживаемый тип условной области')
 preview=adapter.preview(path); nodes={x['locator']:x for x in preview['nodes']}; source=nodes.get(p.start_locator)
 if not source: raise ValueError('Начало области не найдено в исходном DOCX')
 part=source['part']; root=_part_root(path,part); body=_body(root); targets=[]; locator={'locator_version':2,'structure_fingerprint':structure_fingerprint,'part':part,'region_type':p.region_type}
 if p.region_type=='CONDITIONAL_RUN':
  if p.start is None or p.end is None or not p.selected_text: raise ValueError('Выделите текст условного фрагмента')
  adapter.get_structural_locator(path,adapter.select_range(path,p.start_locator,p.start,p.end,p.selected_text)); node=paragraph_for(root,p.start_locator); targets=[node]; locator.update({'node':p.start_locator,'nodes':[p.start_locator],'text_start':p.start,'text_end':p.end,'element_hash':_element_hash(node)})
 elif p.region_type=='CONDITIONAL_PARAGRAPH':
  node=paragraph_for(root,p.start_locator); targets=[node]; locator.update({'node':p.start_locator,'nodes':[p.start_locator],'element_hash':_element_hash(node)})
 elif p.region_type=='CONDITIONAL_PARAGRAPH_GROUP':
  first=_top_paragraph(p.start_locator); last=_top_paragraph(p.end_locator or '')
  if not first or not last or first[0]!=last[0] or first[0]!=part: raise ValueError('Блок можно создать только из соседних абзацев одной части документа')
  a,b=first[1],last[1]
  if a>b: raise ValueError('Конец блока должен находиться после начала')
  paragraphs=body.findall('./w:p',NS)
  if b>len(paragraphs): raise ValueError('Границы блока устарели')
  targets=paragraphs[a-1:b]; locators=[f'{part}#/p[{i}]' for i in range(a,b+1)]; locator.update({'nodes':locators,'start_node':locators[0],'end_node':locators[-1],'element_hashes':[_element_hash(x) for x in targets]})
 elif p.region_type in {'CONDITIONAL_TABLE_ROW','CONDITIONAL_TABLE','CONDITIONAL_TABLE_COLUMN'}:
  parsed=_table_parts(p.start_locator)
  if not parsed: raise ValueError('Выберите ячейку нужной таблицы')
  _,table_no,row_no,_,_=parsed; table=body.findall('./w:tbl',NS)[table_no-1]
  if p.region_type=='CONDITIONAL_TABLE_COLUMN':
   _,_,_,cell_no,_=parsed;selected_row=table.findall('./w:tr',NS)[row_no-1];cursor=1;column=None
   for index,cell in enumerate(selected_row.findall('./w:tc',NS),1):
    span=_span(cell)
    if index==cell_no:
     if span!=1:raise ValueError('Выберите ячейку, не объединённую по горизонтали')
     column=cursor;break
    cursor+=span
   if column is None:raise ValueError('Колонка не найдена')
   _column_layout(table,column);cell_nodes=[]
   for logical_row,row in enumerate(table.findall('./w:tr',NS),1):
    cursor=1
    for physical_cell,cell in enumerate(row.findall('./w:tc',NS),1):
     span=_span(cell)
     if cursor<=column<cursor+span:cell_nodes.append(f'{part}#/tbl[{table_no}]/tr[{logical_row}]/tc[{physical_cell}]');break
     cursor+=span
   targets=[table];locator.update({'table':table_no,'column':column,'nodes':[f'{part}#/tbl[{table_no}]/column[{column}]'],'cell_nodes':cell_nodes,'element_hash':_element_hash(table)})
  elif p.region_type=='CONDITIONAL_TABLE_ROW':
   row=table.findall('./w:tr',NS)[row_no-1]
   if row.find('.//w:vMerge',NS) is not None: raise ValueError('Нельзя условно удалить строку с вертикально объединёнными ячейками')
   targets=[row]; locator.update({'table':table_no,'row':row_no,'nodes':[f'{part}#/tbl[{table_no}]/tr[{row_no}]'],'element_hash':_element_hash(row)})
  else: targets=[table]; locator.update({'table':table_no,'nodes':[f'{part}#/tbl[{table_no}]'],'element_hash':_element_hash(table)})
 bad=next((_unsafe(x) for x in targets if _unsafe(x)),None)
 if bad: raise ValueError(f'Область содержит неподдерживаемую структуру Word ({bad})')
 original='\n'.join(_paragraph_text(x) for target in targets for x in ([target] if target.tag==f'{{{WORD_NS}}}p' else target.findall('.//w:p',NS)))
 if not original.strip(): raise ValueError('Условная область не содержит текста')
 return locator,original
def _binding_inside_region(binding,region):
 node=binding.locator['node']; loc=region.locator; kind=region.region_type
 if kind=='CONDITIONAL_RUN': return node==loc['node'] and binding.locator['text_start']>=loc['text_start'] and binding.locator['text_end']<=loc['text_end']
 if kind in {'CONDITIONAL_PARAGRAPH','CONDITIONAL_PARAGRAPH_GROUP'}: return node in loc.get('nodes',[])
 if kind=='CONDITIONAL_TABLE_COLUMN':
  return any(node.startswith(prefix+'/') for prefix in loc.get('cell_nodes',[]))
 prefix=f"{loc['part']}#/tbl[{loc['table']}]"
 return node.startswith(prefix+(f"/tr[{loc['row']}]" if kind=='CONDITIONAL_TABLE_ROW' else ''))
def validate_regions(v,s,bindings=None):
 fields=[]; schema=schema_for(v.id,s)
 if schema: fields=s.query(QuestionnaireField).filter_by(schema_id=schema.id).all()
 regions=s.query(ConditionalDocumentRegion).filter_by(template_version_id=v.id,status='ACTIVE').all(); seen=[]
 for region in regions:
  try: ConditionEngine.validate_condition(region.include_when,fields)
  except ConditionConfigurationError as exc: raise HTTPException(422,detail=f'Условная область «{region.region_key}»: {exc}')
  loc=region.locator or {}
  if loc.get('structure_fingerprint')!=v.structure_fingerprint: raise HTTPException(422,detail=f'Locator области «{region.region_key}» устарел')
  # Re-resolve stored boundaries and fingerprints against the immutable source.
  root=_part_root(safe_source(v),loc.get('part',''))
  try:
   if region.region_type in {'CONDITIONAL_RUN','CONDITIONAL_PARAGRAPH'}: actual=[paragraph_for(root,loc['node'])]
   elif region.region_type=='CONDITIONAL_PARAGRAPH_GROUP': actual=[paragraph_for(root,x) for x in loc['nodes']]
   elif region.region_type=='CONDITIONAL_TABLE_ROW': actual=[_body(root).findall('./w:tbl',NS)[loc['table']-1].findall('./w:tr',NS)[loc['row']-1]]
   else: actual=[_body(root).findall('./w:tbl',NS)[loc['table']-1]]
  except (KeyError,IndexError,ValueError): raise HTTPException(422,detail=f'Границы области «{region.region_key}» устарели')
  hashes=loc.get('element_hashes') or [loc.get('element_hash')]
  if [_element_hash(x) for x in actual]!=hashes: raise HTTPException(422,detail=f'Locator области «{region.region_key}» устарел')
  for old in seen:
   if _regions_overlap(old,region): raise HTTPException(422,detail=f'Условные области «{old.region_key}» и «{region.region_key}» пересекаются или вложены')
  seen.append(region)
 if bindings:
  for b in bindings:
   for r in regions:
    if r.region_type=='CONDITIONAL_RUN' and b.locator['node']==r.locator['node']:
     a,z=b.locator['text_start'],b.locator['text_end']; c,d=r.locator['text_start'],r.locator['text_end']
     if a<d and z>c and not (a>=c and z<=d): raise HTTPException(422,detail='Обычная привязка частично пересекает условный фрагмент')
 return regions
def _regions_overlap(a,b):
 la,lb=a.locator,b.locator
 if la.get('part')!=lb.get('part'): return False
 def paragraphs(region):
  loc=region.locator
  if region.region_type in {'CONDITIONAL_RUN','CONDITIONAL_PARAGRAPH'}: return {loc['node']}
  if region.region_type=='CONDITIONAL_PARAGRAPH_GROUP': return set(loc['nodes'])
  prefix=f"{loc['part']}#/tbl[{loc['table']}]"
  if region.region_type=='CONDITIONAL_TABLE_COLUMN': return {prefix+f"/column[{loc['column']}]"}
  return {prefix+(f"/tr[{loc['row']}]" if region.region_type=='CONDITIONAL_TABLE_ROW' else '')}
 pa,pb=paragraphs(a),paragraphs(b)
 for x in pa:
  for y in pb:
   if x==y or x.startswith(y+'/') or y.startswith(x+'/'):
    if a.region_type==b.region_type=='CONDITIONAL_RUN': return la['text_start']<lb['text_end'] and la['text_end']>lb['text_start']
    return True
 return False
def _resolve_repeating_row(path,row_locator,structure_fingerprint):
 preview=adapter.preview(path); node=next((x for x in preview['nodes'] if x['locator']==row_locator),None)
 parsed=_table_parts(row_locator)
 if not node or not parsed: raise ValueError('Выберите ячейку строки таблицы')
 part,table_no,row_no,_,_=parsed; root=_part_root(path,part); body=_body(root)
 try: row=body.findall('./w:tbl',NS)[table_no-1].findall('./w:tr',NS)[row_no-1]
 except IndexError: raise ValueError('Строка-образец не найдена')
 if row.find('.//w:vMerge',NS) is not None: raise ValueError('Строки с вертикально объединёнными ячейками пока нельзя повторять')
 bad=_unsafe(row)
 if bad: raise ValueError(f'Строка содержит неподдерживаемую структуру Word ({bad})')
 cells=row.findall('./w:tc',NS)
 if not cells: raise ValueError('Строка-образец не содержит ячеек')
 return {'locator_version':3,'structure_fingerprint':structure_fingerprint,'part':part,'table':table_no,'row':row_no,'row_node':f'{part}#/tbl[{table_no}]/tr[{row_no}]','row_hash':_element_hash(row),'cell_count':len(cells),'has_grid_span':row.find('.//w:gridSpan',NS) is not None,'prototype_text':' | '.join(_paragraph_text(x) for x in cells)}
def _group_row(root,loc): return _body(root).findall('./w:tbl',NS)[loc['table']-1].findall('./w:tr',NS)[loc['row']-1]
def _field_in_group_row(locator,group):
 parsed=_table_parts(locator)
 return bool(parsed and parsed[0]==group.locator['part'] and parsed[1]==group.locator['table'] and parsed[2]==group.locator['row'])
def validate_repeating_groups(v,s,bindings=None,regions=None):
 schema=schema_for(v.id,s); questionnaire=s.query(QuestionnaireField).filter_by(schema_id=schema.id).all() if schema else []
 groups=s.query(RepeatingGroup).filter_by(template_version_id=v.id,status='ACTIVE').order_by(RepeatingGroup.position).all(); occupied=set()
 for group in groups:
  if group.min_items<0 or group.max_items<max(group.min_items,1) or group.max_items>50: raise HTTPException(422,detail=f'Группа «{group.name}»: проверьте минимальное и максимальное количество')
  condition=(group.config or {}).get('visible_when')
  if condition is not None:
   try: ConditionEngine.validate_condition(condition,questionnaire)
   except ConditionConfigurationError as exc: raise HTTPException(422,detail=f'Группа «{group.name}»: {exc}')
  loc=group.locator or {}
  if loc.get('structure_fingerprint')!=v.structure_fingerprint: raise HTTPException(422,detail=f'Строка-образец группы «{group.name}» устарела')
  try: root=_part_root(safe_source(v),loc['part']); row=_group_row(root,loc)
  except (KeyError,IndexError): raise HTTPException(422,detail=f'Строка-образец группы «{group.name}» не найдена')
  if _element_hash(row)!=loc.get('row_hash'): raise HTTPException(422,detail=f'Строка-образец группы «{group.name}» устарела')
  key=(loc['part'],loc['table'],loc['row'])
  if key in occupied: raise HTTPException(422,detail='Несколько повторяемых групп используют одну строку-образец')
  occupied.add(key)
  fields=s.query(RepeatingGroupField).filter_by(group_id=group.id).order_by(RepeatingGroupField.position).all()
  if not fields: raise HTTPException(422,detail=f'Добавьте поля в группу «{group.name}»')
  for field in fields:
   if not field.locator: raise HTTPException(422,detail=f'Свяжите поле «{field.label}» с текстом в строке-образце')
   floc=field.locator
   if not _field_in_group_row(floc.get('node',''),group) or floc.get('structure_fingerprint')!=v.structure_fingerprint: raise HTTPException(422,detail=f'Привязка поля «{field.label}» устарела')
   try: p=paragraph_for(root,floc['node']); text=_paragraph_text(p)
   except (KeyError,IndexError,ValueError): raise HTTPException(422,detail=f'Привязка поля «{field.label}» не найдена')
   if text[floc['text_start']:floc['text_end']]!=floc.get('selected_text'): raise HTTPException(422,detail=f'Привязка поля «{field.label}» устарела')
  ranges={}
  for field in fields:
   locf=field.locator; ranges.setdefault(locf['node'],[]).append((locf['text_start'],locf['text_end']))
  if any(any(a<d and b>c for i,(a,b) in enumerate(items) for c,d in items[i+1:]) for items in ranges.values()): raise HTTPException(422,detail=f'Привязки полей группы «{group.name}» пересекаются')
  for binding in bindings or []:
   if _field_in_group_row(binding.locator.get('node',''),group): raise HTTPException(422,detail=f'Обычная привязка пересекает строку-образец группы «{group.name}»')
  for region in regions or []:
   fake=type('RowRegion',(),{'region_type':'CONDITIONAL_TABLE_ROW','locator':loc})()
   if _regions_overlap(fake,region): raise HTTPException(422,detail=f'Условная область пересекает строку-образец группы «{group.name}»')
 return groups
def paragraph_for(root,locator):
 path=locator.split('#/',1)[1]; body=root.find('w:body',NS) if root.tag.endswith('document') else root
 m=re.fullmatch(r'p\[(\d+)\]',path)
 if m: return body.findall('./w:p',NS)[int(m.group(1))-1]
 m=re.fullmatch(r'tbl\[(\d+)\]/tr\[(\d+)\]/tc\[(\d+)\]/p\[(\d+)\]',path)
 if not m: raise ValueError('Неподдерживаемый структурный locator')
 t,r,c,p=map(int,m.groups()); return body.findall('./w:tbl',NS)[t-1].findall('./w:tr',NS)[r-1].findall('./w:tc',NS)[c-1].findall('./w:p',NS)[p-1]
def replace_range(p,start,end,value):
 if p.find('.//w:fldChar',NS) is not None or p.find('.//w:instrText',NS) is not None: raise ValueError('Привязка внутри поля Word не поддерживается')
 spans=[]; pos=0
 for node in p.findall('.//w:t',NS): text=node.text or ''; spans.append((node,pos,pos+len(text),text)); pos+=len(text)
 if start==end==0 and pos==0:
  run=p.find('./w:r',NS)
  if run is None: run=ET.SubElement(p,f'{{{WORD_NS}}}r')
  text_node=run.find('./w:t',NS)
  if text_node is None: text_node=ET.SubElement(run,f'{{{WORD_NS}}}t')
  text_node.text=str(value)
  return
 touched=[x for x in spans if x[2]>start and x[1]<end]
 if end>pos or not touched: raise ValueError('Диапазон привязки выходит за границы текста')
 first,last=touched[0],touched[-1]
 for node,s,e,text in touched:
  prefix=text[:max(0,start-s)] if node is first[0] else ''; suffix=text[max(0,end-s):] if node is last[0] else ''; node.text=(prefix+str(value)+suffix) if node is first[0] else suffix
  if node.text.startswith(' ') or node.text.endswith(' '): node.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
def _parent_of(root,node):
 for parent in root.iter():
  if node in list(parent): return parent
 return None
def render_docx(v,bindings,values,target,regions=None,included_region_ids=None,repeating_groups=None,repeating_values=None,repeating_fields=None):
 regions=regions or []; included_region_ids=set(included_region_ids or [])
 repeating_groups=repeating_groups or []; repeating_values=repeating_values or {}; repeating_fields=repeating_fields or {}
 excluded=[x for x in regions if x.id not in included_region_ids]
 bindings=[b for b in bindings if not any(_binding_inside_region(b,r) for r in excluded)]
 grouped={}
 for b in bindings: grouped.setdefault(b.locator['part'],{'bindings':[],'regions':[]})['bindings'].append(b)
 for r in excluded: grouped.setdefault(r.locator['part'],{'bindings':[],'regions':[]})['regions'].append(r)
 for g in repeating_groups: grouped.setdefault(g.locator['part'],{'bindings':[],'regions':[],'groups':[]}).setdefault('groups',[]).append(g)
 with zipfile.ZipFile(safe_source(v),'r') as src,zipfile.ZipFile(target,'w') as dst:
  for info in src.infolist():
   raw=src.read(info.filename)
   if info.filename in grouped:
    root=ET.fromstring(raw); by_node={}; jobs=[]
    for group in grouped[info.filename].get('groups',[]): jobs.append((group,_group_row(root,group.locator)))
    for b in grouped[info.filename]['bindings']: by_node.setdefault(b.locator['node'],[]).append((b.locator['text_start'],b.locator['text_end'],values[b.id] if b.id in values else values[b.binding_key]))
    for region in grouped[info.filename]['regions']:
     if region.region_type=='CONDITIONAL_RUN': by_node.setdefault(region.locator['node'],[]).append((region.locator['text_start'],region.locator['text_end'],''))
    for locator,items in by_node.items():
     p=paragraph_for(root,locator)
     for start,end,value in sorted(items,key=lambda x:x[0],reverse=True): replace_range(p,start,end,value)
    for group,row in jobs:
     table=_parent_of(root,row)
     if table is None: raise ValueError(f'Строка-образец группы «{group.name}» недоступна')
     insert_at=list(table).index(row); items=repeating_values.get(group.group_key,[]); fields=repeating_fields.get(group.id,[])
     for offset,item in enumerate(items):
      clone=copy.deepcopy(row); actions={}
      for field in fields:
       loc=field.locator; parsed=_table_parts(loc['node'])
       if not parsed: raise ValueError(f'Привязка поля «{field.label}» повреждена')
       _,_,_,cell_no,paragraph_no=parsed
       try: paragraph=clone.findall('./w:tc',NS)[cell_no-1].findall('./w:p',NS)[paragraph_no-1]
       except IndexError: raise ValueError(f'Привязка поля «{field.label}» устарела')
       cfg=field.config or {}; formatted=render_value(field,item.get(field.field_key,'')) if field.field_type=='BOOLEAN' and cfg.get('output_format','RAW')=='RAW' else ValueFormatter.format(item.get(field.field_key,''),cfg.get('output_format','RAW'),{'precision':cfg.get('precision',2),'rounding':cfg.get('rounding','HALF_UP')})
       actions.setdefault(paragraph,[]).append((loc['text_start'],loc['text_end'],formatted))
      for paragraph,ranges in actions.items():
       for start,end,value in sorted(ranges,key=lambda x:x[0],reverse=True): replace_range(paragraph,start,end,value)
      table.insert(insert_at+offset,clone)
     table.remove(row)
    removals=[]
    for region in grouped[info.filename]['regions']:
     loc=region.locator
     if region.region_type=='CONDITIONAL_RUN': continue
     elif region.region_type in {'CONDITIONAL_PARAGRAPH','CONDITIONAL_PARAGRAPH_GROUP'}: removals.extend(paragraph_for(root,x) for x in loc['nodes'])
     elif region.region_type=='CONDITIONAL_TABLE_ROW': removals.append(_body(root).findall('./w:tbl',NS)[loc['table']-1].findall('./w:tr',NS)[loc['row']-1])
     elif region.region_type=='CONDITIONAL_TABLE_COLUMN': _remove_table_column(_body(root).findall('./w:tbl',NS)[loc['table']-1],loc['column'])
     else: removals.append(_body(root).findall('./w:tbl',NS)[loc['table']-1])
    for node in removals:
     parent=_parent_of(root,node)
     if parent is None: raise ValueError('Не удалось безопасно удалить OOXML-область')
     parent.remove(node)
    raw=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   dst.writestr(info,raw)
 validate_docx(target.name,DOCX_MIME,target.read_bytes())

def build_contract_platform_router(current_dependency:Callable)->APIRouter:
 r=APIRouter(tags=['contract-platform']); prefix='/api/v1/document-templates'
 @r.get('/api/v1/contract-catalog')
 def catalog(q:str='',u=Depends(current_dependency),s:DBS=Depends(db)):
  query=s.query(DocumentTemplate).filter_by(status='ACTIVE').filter(DocumentTemplate.current_version_id.is_not(None))
  product_versions=s.query(ContractProductVersion.id).filter_by(status='PUBLISHED')
  product_template_versions=s.query(ContractProductVariant.template_version_id).filter(ContractProductVariant.product_version_id.in_(product_versions))
  query=query.filter(~DocumentTemplate.current_version_id.in_(product_template_versions))
  if q: query=query.filter(DocumentTemplate.name.ilike(f'%{q}%'))
  return [template_out(x,s) for x in query.order_by(DocumentTemplate.name).all()]
 @r.post('/api/v1/contract-drafts',status_code=201)
 def create_draft(payload:dict,u=Depends(current_dependency),s:DBS=Depends(db)):
  t=s.get(DocumentTemplate,payload.get('template_id'))
  if not t or t.status!='ACTIVE' or not t.current_version_id: raise HTTPException(404,detail='Опубликованный шаблон не найден')
  v=s.get(TemplateVersion,t.current_version_id); schema=schema_for(v.id,s)
  if not schema or schema.status!='PUBLISHED': raise HTTPException(409,detail='Анкета шаблона не опубликована')
  d=DocumentDraft(owner_id=u.id,template_version_id=v.id,questionnaire_schema_id=schema.id); s.add(d); s.commit(); return {'id':d.id,'template':template_out(t,s),'status':d.status,'answers':d.answers,'schema':schema_out(schema,s)}
 @r.get('/api/v1/contract-drafts')
 def drafts(u=Depends(current_dependency),s:DBS=Depends(db)):
  out=[]
  for d in s.query(DocumentDraft).filter_by(owner_id=u.id).order_by(DocumentDraft.updated_at.desc()).all():
   v=s.get(TemplateVersion,d.template_version_id); t=s.get(DocumentTemplate,v.template_id); g=s.query(GeneratedDocument).filter_by(draft_id=d.id).order_by(GeneratedDocument.generation_number.desc()).first(); out.append({'id':d.id,'template_name':t.name,'template_version':v.version_number,'status':d.status,'created_at':d.created_at,'updated_at':d.updated_at,'generated':g.id if g else None,'generated_stale':bool(g and g.created_at<d.updated_at)})
  return out
 @r.get('/api/v1/contract-drafts/{id}')
 def draft(id:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  d=own_draft(id,u,s); v=s.get(TemplateVersion,d.template_version_id); t=s.get(DocumentTemplate,v.template_id); schema=s.get(QuestionnaireSchema,d.questionnaire_schema_id); return {'id':d.id,'status':d.status,'answers':d.answers,'template':template_out(t,s),'schema':schema_out(schema,s)}
 @r.patch('/api/v1/contract-drafts/{id}')
 def save(id:str,payload:AnswersPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  d=own_draft(id,u,s); fields=s.query(QuestionnaireField).filter_by(schema_id=d.questionnaire_schema_id).all(); groups=s.query(RepeatingGroup).filter_by(template_version_id=d.template_version_id,status='ACTIVE').all(); known={x.field_key:x for x in fields}; known.update({x.group_key:x for x in groups})
  validate_condition_schema(fields)
  if set(payload.answers)-set(known): raise HTTPException(422,detail='Анкета содержит неизвестные поля')
  if any(k in payload.answers and (f.config or {}).get('calculated') for k,f in known.items() if isinstance(f,QuestionnaireField)): raise HTTPException(422,detail='Вычисляемые поля нельзя изменять')
  answers={**(d.answers or {}),**payload.answers}; errors=validate_active_answers(fields,answers); errors.update(validate_repeating_answers(groups,s,fields,answers,True))
  if errors: raise HTTPException(422,detail={'message':'Исправьте ответы','fields':errors})
  d.answers=answers; d.status='DRAFT'; d.updated_at=now(); s.commit(); return {'id':d.id,'status':d.status,'answers':d.answers}
 @r.patch('/api/v1/contract-drafts/{id}/autosave')
 def autosave(id:str,payload:AnswersPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  d=own_draft(id,u,s); fields=s.query(QuestionnaireField).filter_by(schema_id=d.questionnaire_schema_id).all(); groups=s.query(RepeatingGroup).filter_by(template_version_id=d.template_version_id,status='ACTIVE').all(); known={x.field_key for x in fields}|{x.group_key for x in groups}
  if set(payload.answers)-known: raise HTTPException(422,detail='Анкета содержит неизвестные поля')
  if any(x.field_key in payload.answers and (x.config or {}).get('calculated') for x in fields): raise HTTPException(422,detail='Вычисляемые поля нельзя изменять')
  answers={**(d.answers or {}),**payload.answers}; errors=validate_repeating_answers(groups,s,fields,answers,False)
  if errors: raise HTTPException(422,detail={'message':'Не удалось сохранить черновик','fields':errors})
  for field in fields:
   value=answers.get(field.field_key)
   if not (field.config or {}).get('calculated') and value not in (None,'') and (message:=validate_answer(field,value)): errors[field.field_key]=message
  if errors: raise HTTPException(422,detail={'message':'Не удалось сохранить черновик','fields':errors})
  d.answers=answers; d.status='DRAFT'; d.updated_at=now(); s.commit(); return {'id':d.id,'status':d.status,'answers':d.answers}
 @r.post('/api/v1/contract-drafts/{id}/generate',status_code=201)
 def generate(id:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  d=own_draft(id,u,s); v=s.get(TemplateVersion,d.template_version_id); fields=s.query(QuestionnaireField).filter_by(schema_id=d.questionnaire_schema_id).all(); validate_condition_schema(fields); answers=d.answers or {}; errors=validate_active_answers(fields,answers)
  groups=s.query(RepeatingGroup).filter_by(template_version_id=v.id,status='ACTIVE').all(); errors.update(validate_repeating_answers(groups,s,fields,answers,True))
  if errors: raise HTTPException(422,detail={'message':'Заполните обязательные поля','fields':errors})
  bindings=validate_bindings(v,s); regions=validate_regions(v,s,bindings); groups=validate_repeating_groups(v,s,bindings,regions); active={x.field_key:x for x in ConditionEngine.active_fields(fields,answers)}; bindings=[x for x in bindings if x.binding_key in active]; legal_answers=condition_answers(fields,answers); included={x.id for x in regions if ConditionEngine.evaluate(x.include_when,legal_answers)}; groups,calculated_values,repeated=calculated_document_values(v.id,s,fields,answers); values=formatted_bindings(bindings,calculated_values,fields); group_fields={x.id:x.fields for x in groups}; number=s.query(GeneratedDocument).filter_by(draft_id=d.id).count()+1; root=(Path(settings.storage)/'generated_contracts'/d.id).resolve(); root.mkdir(parents=True,exist_ok=True); target=root/f'{number}.docx'
  try: render_docx(v,bindings,values,target,regions,included,groups,repeated,group_fields)
  except (ValueError,IndexError,KeyError,CalculationError) as exc: raise HTTPException(422,detail=f'Не удалось безопасно сформировать DOCX: {exc}')
  raw=target.read_bytes(); g=GeneratedDocument(draft_id=d.id,template_version_id=v.id,generation_number=number,path=str(target),sha256=hashlib.sha256(raw).hexdigest(),mime_type=DOCX_MIME,size=len(raw),created_by_id=u.id); s.add(g); d.status='GENERATED'; s.commit(); return {'id':g.id,'generation_number':number,'filename':f'Договор_{number}.docx','download_url':f'/api/v1/generated-documents/{g.id}/download'}
 @r.get('/api/v1/generated-documents/{id}/download')
 def generated(id:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  g=s.get(GeneratedDocument,id)
  if not g: raise HTTPException(404,detail='Документ не найден')
  d=own_draft(g.draft_id,u,s); t=s.get(DocumentTemplate,s.get(TemplateVersion,d.template_version_id).template_id); return FileResponse(safe_file(g.path,'generated_contracts',g.sha256),filename=f'{t.name}.docx',media_type=DOCX_MIME)
 @r.get(prefix)
 def templates(u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); query=s.query(DocumentTemplate)
  if 'administrator' not in roles(u): query=query.filter_by(owner_id=u.id)
  return [template_out(x,s) for x in query.order_by(DocumentTemplate.updated_at.desc()).all()]
 @r.post(prefix,status_code=201)
 def create_template(payload:CreateTemplate,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); code=payload.code.upper()
  if s.query(DocumentTemplate).filter_by(owner_id=u.id,code=code).first(): raise HTTPException(409,detail='Шаблон с таким кодом уже существует')
  x=DocumentTemplate(owner_id=u.id,code=code,name=payload.name.strip(),category=payload.category.strip(),description=payload.description.strip()); s.add(x); s.flush(); s.add(TemplateAuditEvent(template_id=x.id,actor_id=u.id,event_type='TEMPLATE_CREATED')); s.commit(); return template_out(x,s)
 @r.get(prefix+'/{tid}')
 def get_template(tid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); return template_out(owned_template(tid,u,s),s)
 @r.get(prefix+'/{tid}/versions')
 def versions(tid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); owned_template(tid,u,s); return [version_out(x) for x in s.query(TemplateVersion).filter_by(template_id=tid).order_by(TemplateVersion.version_number.desc()).all()]
 @r.post(prefix+'/{tid}/versions',status_code=201)
 async def upload(tid:str,file:UploadFile=File(...),u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); t=owned_template(tid,u,s); raw=await file.read(); filename=Path(file.filename or '').name; validate_docx(filename,file.content_type or '',raw); last=s.query(TemplateVersion).filter_by(template_id=t.id).order_by(TemplateVersion.version_number.desc()).first(); v=TemplateVersion(template_id=t.id,version_number=last.version_number+1 if last else 1,original_filename=filename,source_path='',source_sha256=hashlib.sha256(raw).hexdigest(),source_size=len(raw),mime_type=DOCX_MIME,structure_fingerprint='',created_by_id=u.id); s.add(v); s.flush(); root=(Path(settings.storage)/'template_sources'/t.id/v.id).resolve(); root.mkdir(parents=True); path=root/'source.docx'; path.write_bytes(raw); path.chmod(0o444); preview=adapter.preview(path); v.source_path=str(path); v.structure_fingerprint=preview['structure_fingerprint']; v.preview_metadata={'adapter':preview['adapter'],'node_count':len(preview['nodes']),'limitations':preview['limitations']}; schema_for(v.id,s,True,u); s.add(TemplateAuditEvent(template_id=t.id,template_version_id=v.id,actor_id=u.id,event_type='VERSION_UPLOADED')); s.commit(); return version_out(v)
 @r.get(prefix+'/{tid}/versions/{vid}/preview')
 def preview(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  result=adapter.preview(safe_source(v))
  if result['structure_fingerprint']!=v.structure_fingerprint: raise HTTPException(409,detail='Структура исходного файла изменилась')
  return result
 @r.get(prefix+'/{tid}/versions/{vid}/source')
 def source(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  return FileResponse(safe_source(v),filename=v.original_filename,media_type=DOCX_MIME)
 @r.get(prefix+'/{tid}/versions/{vid}/binding-candidates')
 def binding_candidates(tid:str,vid:str,rf_id:str|None=None,variant:str|None=None,status:str|None=None,operation:str|None=None,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  items=(v.preview_metadata or {}).get('binding_candidates',[])
  filters={'rf_id':rf_id,'variant':variant,'status':status,'operation':operation}
  return [item for item in items if all(value is None or item.get(key)==value for key,value in filters.items())]
 @r.patch(prefix+'/{tid}/versions/{vid}/binding-candidates/{candidate_id}')
 def decide_binding_candidate(tid:str,vid:str,candidate_id:str,p:CandidateDecisionPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  metadata=copy.deepcopy(v.preview_metadata or {}); candidates=metadata.get('binding_candidates',[])
  candidate=next((x for x in candidates if x.get('candidate_id')==candidate_id),None)
  if candidate is None: raise HTTPException(404,detail='Кандидат разметки не найден')
  candidate['review_decision']=p.decision
  candidate['status']='REVIEW_REQUIRED' if p.decision=='REJECTED' else candidate.get('status','CANDIDATE_FOUND')
  metadata['binding_candidates']=candidates; v.preview_metadata=metadata
  s.add(TemplateAuditEvent(template_id=t.id,template_version_id=v.id,actor_id=u.id,event_type='BINDING_CANDIDATE_'+p.decision,metadata_json={'candidate_id':candidate_id}))
  s.commit(); return candidate
 @r.post(prefix+'/{tid}/versions/{vid}/binding-candidates/{candidate_id}/materialize')
 def materialize_candidate(tid:str,vid:str,candidate_id:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  from app.binding_materialization import BindingMaterializationService,MaterializationError
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  try:return BindingMaterializationService(s,u).materialize(v,candidate_id,manual=True)
  except MaterializationError as exc:s.rollback();raise HTTPException(422,detail=str(exc))
 @r.post(prefix+'/{tid}/versions/{vid}/binding-candidates/auto-materialize')
 def auto_materialize_candidates(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  from app.binding_materialization import BindingMaterializationService,MaterializationError
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  service=BindingMaterializationService(s,u);done=[];blocked=[]
  for item in list((v.preview_metadata or {}).get('binding_candidates',[])):
   try:done.append(service.materialize(v,item['candidate_id'],manual=False)['candidate_id'])
   except MaterializationError as exc:s.rollback();v=s.get(TemplateVersion,vid);blocked.append({'candidate_id':item['candidate_id'],'reason':str(exc)})
  return {'materialized':done,'blocked':blocked}
 @r.post(prefix+'/{tid}/versions/{vid}/binding-candidates/batch-materialize')
 def batch_materialize_candidates(tid:str,vid:str,p:CandidateBatchPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  from app.binding_materialization import BindingMaterializationService,MaterializationError
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  if len(set(p.candidate_ids))!=len(p.candidate_ids):raise HTTPException(422,detail='Список содержит дубликаты')
  service=BindingMaterializationService(s,u);done=[]
  try:
   for candidate_id in p.candidate_ids:done.append(service.materialize(v,candidate_id,manual=True)['candidate_id']);v=s.get(TemplateVersion,vid)
  except MaterializationError as exc:s.rollback();raise HTTPException(422,detail={'message':str(exc),'materialized_before_error':done})
  return {'materialized':done}
 @r.delete(prefix+'/{tid}/versions/{vid}/binding-candidates/{candidate_id}/materialization')
 def rollback_candidate(tid:str,vid:str,candidate_id:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  from app.binding_materialization import BindingMaterializationService,MaterializationError
  require_manager(u);v,t=owned_version(vid,u,s)
  if t.id!=tid:raise HTTPException(404,detail='Версия не найдена')
  try:return BindingMaterializationService(s,u).rollback(v,candidate_id)
  except MaterializationError as exc:s.rollback();raise HTTPException(422,detail=str(exc))
 @r.post(prefix+'/{tid}/versions/{vid}/test-drive',status_code=201)
 def draft_test_drive(tid:str,vid:str,p:AnswersPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u);v,t=owned_version(vid,u,s);require_draft(v)
  if t.id!=tid:raise HTTPException(404,detail='Версия не найдена')
  candidates=(v.preview_metadata or {}).get('binding_candidates',[]);unresolved=[x['candidate_id'] for x in candidates if x.get('status') not in {'CONFIGURED','GENERATED_AND_VERIFIED'}]
  if unresolved:raise HTTPException(409,detail={'message':'Сначала завершите структурную разметку','unresolved_count':len(unresolved),'candidate_ids':unresolved[:50]})
  schema=schema_for(v.id,s);fields=s.query(QuestionnaireField).filter_by(schema_id=schema.id).all() if schema else []
  known={x.field_key for x in fields}
  if set(p.answers)-known:raise HTTPException(422,detail='Анкета содержит неизвестные поля')
  validate_condition_schema(fields);errors=validate_active_answers(fields,p.answers)
  groups=s.query(RepeatingGroup).filter_by(template_version_id=v.id,status='ACTIVE').all();errors.update(validate_repeating_answers(groups,s,fields,p.answers,True))
  if errors:raise HTTPException(422,detail={'message':'Исправьте тестовые ответы','fields':errors})
  bindings=validate_bindings(v,s);regions=validate_regions(v,s,bindings);groups=validate_repeating_groups(v,s,bindings,regions)
  active={x.field_key:x for x in ConditionEngine.active_fields(fields,p.answers)};bindings=[x for x in bindings if x.binding_key in active];legal=condition_answers(fields,p.answers);included={x.id for x in regions if ConditionEngine.evaluate(x.include_when,legal)}
  groups,calculated,repeated=calculated_document_values(v.id,s,fields,p.answers);values=formatted_bindings(bindings,calculated,fields);group_fields={x.id:x.fields for x in groups}
  document_id=str(uuid.uuid4());root=(Path(settings.storage)/'test_contracts'/v.id).resolve();root.mkdir(parents=True,exist_ok=True);target=root/f'{document_id}.docx'
  try:
   render_docx(v,bindings,values,target,regions,included,groups,repeated,group_fields)
   marked=Document(target);marked.core_properties.title='ТЕСТОВЫЙ ДОКУМЕНТ';marked.core_properties.comments='Создано в draft test-drive. Не является официальным договором.'
   for section in marked.sections:
    paragraph=section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph();paragraph.text='ТЕСТОВЫЙ ДОКУМЕНТ · НЕ ДЛЯ ПОДПИСАНИЯ'
   marked.save(target)
  except (ValueError,IndexError,KeyError,CalculationError) as exc:raise HTTPException(422,detail=f'Не удалось сформировать тестовый DOCX: {exc}')
  raw=target.read_bytes();checksum=hashlib.sha256(raw).hexdigest();metadata=copy.deepcopy(v.preview_metadata or {});documents=metadata.get('test_documents',[]);documents.append({'id':document_id,'path':str(target),'sha256':checksum,'created_by_id':u.id,'created_at':now().isoformat()});metadata['test_documents']=documents[-20:];v.preview_metadata=metadata
  s.add(TemplateAuditEvent(template_id=t.id,template_version_id=v.id,actor_id=u.id,event_type='DRAFT_TEST_DOCUMENT_GENERATED',metadata_json={'document_id':document_id,'sha256':checksum}));s.commit()
  return {'id':document_id,'test_only':True,'sha256':checksum,'download_url':f'{prefix}/{tid}/versions/{vid}/test-documents/{document_id}/download'}
 @r.get(prefix+'/{tid}/versions/{vid}/test-documents/{document_id}/download')
 def download_test_document(tid:str,vid:str,document_id:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u);v,t=owned_version(vid,u,s)
  if t.id!=tid:raise HTTPException(404,detail='Версия не найдена')
  item=next((x for x in (v.preview_metadata or {}).get('test_documents',[]) if x.get('id')==document_id and (x.get('created_by_id')==u.id or 'administrator' in roles(u))),None)
  if not item:raise HTTPException(404,detail='Тестовый документ не найден')
  return FileResponse(safe_file(item['path'],'test_contracts',item['sha256']),filename=f'ТЕСТОВЫЙ_{t.name}.docx',media_type=DOCX_MIME)
 @r.get(prefix+'/{tid}/versions/{vid}/questionnaire')
 def questionnaire(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  return schema_out(schema_for(v.id,s,True,u),s)
 @r.post(prefix+'/{tid}/versions/{vid}/fields',status_code=201)
 def create_field(tid:str,vid:str,p:FieldPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); validate_field_payload(p); schema=schema_for(v.id,s,True,u)
  if s.query(QuestionnaireField).filter_by(schema_id=schema.id,field_key=p.field_key).first(): raise HTTPException(409,detail='Поле с таким ключом уже существует')
  if s.query(QuestionnaireField).filter_by(schema_id=schema.id,position=p.position).first(): raise HTTPException(409,detail='Позиция уже занята')
  x=QuestionnaireField(schema_id=schema.id,field_key=p.field_key,label=p.label,field_type=p.field_type,required=p.required,position=p.position,config=field_config(p)); s.add(x); s.flush()
  try:
   validate_condition_schema(s.query(QuestionnaireField).filter_by(schema_id=schema.id).all())
   validate_regions(v,s)
   validate_calculations(v.id,s)
  except HTTPException: s.rollback(); raise
  s.commit(); return field_out(x)
 @r.patch(prefix+'/{tid}/versions/{vid}/fields/{fid}')
 def update_field(tid:str,vid:str,fid:str,p:FieldPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); validate_field_payload(p); schema=schema_for(v.id,s); x=s.get(QuestionnaireField,fid)
  if not x or not schema or x.schema_id!=schema.id: raise HTTPException(404,detail='Поле не найдено')
  duplicate=s.query(QuestionnaireField).filter(QuestionnaireField.schema_id==schema.id,QuestionnaireField.field_key==p.field_key,QuestionnaireField.id!=x.id).first()
  occupied=s.query(QuestionnaireField).filter(QuestionnaireField.schema_id==schema.id,QuestionnaireField.position==p.position,QuestionnaireField.id!=x.id).first()
  if duplicate: raise HTTPException(409,detail='Поле с таким ключом уже существует')
  if occupied: raise HTTPException(409,detail='Позиция уже занята')
  if x.field_key!=p.field_key:
   for binding in s.query(DocumentBinding).filter_by(field_id=x.id): binding.binding_key=p.field_key
  x.field_key=p.field_key; x.label=p.label; x.field_type=p.field_type; x.required=p.required; x.position=p.position; x.config=field_config(p); s.flush()
  try:
   validate_condition_schema(s.query(QuestionnaireField).filter_by(schema_id=schema.id).all())
   validate_regions(v,s)
   validate_calculations(v.id,s)
  except HTTPException: s.rollback(); raise
  s.commit(); return field_out(x)
 @r.delete(prefix+'/{tid}/versions/{vid}/fields/{fid}',status_code=204)
 def delete_field(tid:str,vid:str,fid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); schema=schema_for(v.id,s); x=s.get(QuestionnaireField,fid)
  if not x or not schema or x.schema_id!=schema.id: raise HTTPException(404,detail='Поле не найдено')
  if s.query(DocumentBinding).filter_by(field_id=x.id).first(): raise HTTPException(409,detail='Сначала удалите привязки поля')
  s.delete(x); s.flush()
  try:
   validate_condition_schema(s.query(QuestionnaireField).filter_by(schema_id=schema.id).all())
   validate_regions(v,s)
   validate_calculations(v.id,s)
  except HTTPException: s.rollback(); raise
  s.commit()
 @r.get(prefix+'/{tid}/versions/{vid}/bindings')
 def bindings(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); return [binding_out(x) for x in s.query(DocumentBinding).filter_by(template_version_id=v.id).all()]
 @r.post(prefix+'/{tid}/versions/{vid}/bindings',status_code=201)
 def bind(tid:str,vid:str,p:BindingPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); schema=schema_for(v.id,s); field=s.get(QuestionnaireField,p.field_id) if p.field_id else None
  if p.field_id and (not field or not schema or field.schema_id!=schema.id): raise HTTPException(404,detail='Поле анкеты не найдено')
  key=field.field_key if field else p.binding_key
  if not key: raise HTTPException(422,detail='Выберите поле анкеты')
  try: ValueFormatter.validate(p.output_format); ValueFormatter.validate_options(p.format_options); path=safe_source(v); loc=adapter.get_structural_locator(path,adapter.select_range(path,p.locator,p.start,p.end,p.selected_text))
  except (ValueError,CalculationError) as exc: raise HTTPException(422,detail=str(exc))
  for old in s.query(DocumentBinding).filter_by(template_version_id=v.id,status='ACTIVE'):
   o=old.locator
   if o.get('node')==loc['node'] and loc['text_start']<o['text_end'] and loc['text_end']>o['text_start']: raise HTTPException(409,detail='Выделенный диапазон пересекается с существующей привязкой')
  n=s.query(DocumentBinding).filter_by(template_version_id=v.id,binding_key=key).count()+1; x=DocumentBinding(template_version_id=v.id,field_id=field.id if field else None,binding_key=key,occurrence_index=n,locator=loc,selected_text=p.selected_text,format_config={'format':p.output_format,'options':p.format_options},created_by_id=u.id); s.add(x); s.flush(); s.add(TemplateAuditEvent(template_id=t.id,template_version_id=v.id,actor_id=u.id,event_type='BINDING_CREATED',metadata_json={'binding_id':x.id})); s.commit(); return binding_out(x)
 @r.patch(prefix+'/{tid}/versions/{vid}/bindings/{bid}/format')
 def update_binding_format(tid:str,vid:str,bid:str,p:BindingFormatPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); x=s.get(DocumentBinding,bid)
  if not x or x.template_version_id!=v.id: raise HTTPException(404,detail='Привязка не найдена')
  try: ValueFormatter.validate(p.output_format); ValueFormatter.validate_options(p.format_options)
  except CalculationError as exc: raise HTTPException(422,detail=str(exc))
  x.format_config={'format':p.output_format,'options':p.format_options}; s.commit(); return binding_out(x)
 @r.delete(prefix+'/{tid}/versions/{vid}/bindings/{bid}',status_code=204)
 def delete_binding(tid:str,vid:str,bid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); x=s.get(DocumentBinding,bid)
  if not x or x.template_version_id!=v.id: raise HTTPException(404,detail='Привязка не найдена')
  s.delete(x); s.commit()
 @r.get(prefix+'/{tid}/versions/{vid}/conditional-regions')
 def conditional_regions(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  return [region_out(x) for x in s.query(ConditionalDocumentRegion).filter_by(template_version_id=v.id).order_by(ConditionalDocumentRegion.created_at).all()]
 @r.post(prefix+'/{tid}/versions/{vid}/conditional-regions',status_code=201)
 def create_region(tid:str,vid:str,p:RegionPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); schema=schema_for(v.id,s); fields=s.query(QuestionnaireField).filter_by(schema_id=schema.id).all() if schema else []
  if s.query(ConditionalDocumentRegion).filter_by(template_version_id=v.id,region_key=p.region_key).first(): raise HTTPException(409,detail='Область с таким ключом уже существует')
  try:
   ConditionEngine.validate_condition(p.include_when,fields); locator,original=_resolve_region(safe_source(v),p,v.structure_fingerprint)
  except (ConditionConfigurationError,ValueError,IndexError) as exc: raise HTTPException(422,detail=str(exc))
  x=ConditionalDocumentRegion(template_version_id=v.id,region_key=p.region_key,region_type=p.region_type,locator=locator,include_when=p.include_when,original_text=original,created_by_id=u.id); s.add(x); s.flush()
  try: validate_regions(v,s,validate_bindings(v,s))
  except HTTPException: s.rollback(); raise
  s.add(TemplateAuditEvent(template_id=t.id,template_version_id=v.id,actor_id=u.id,event_type='CONDITIONAL_REGION_CREATED',metadata_json={'region_id':x.id,'region_type':x.region_type})); s.commit(); return region_out(x)
 @r.patch(prefix+'/{tid}/versions/{vid}/conditional-regions/{rid}')
 def update_region(tid:str,vid:str,rid:str,p:RegionUpdatePayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); x=s.get(ConditionalDocumentRegion,rid)
  if not x or x.template_version_id!=v.id: raise HTTPException(404,detail='Условная область не найдена')
  duplicate=s.query(ConditionalDocumentRegion).filter(ConditionalDocumentRegion.template_version_id==v.id,ConditionalDocumentRegion.region_key==p.region_key,ConditionalDocumentRegion.id!=x.id).first()
  if duplicate: raise HTTPException(409,detail='Область с таким ключом уже существует')
  fields=s.query(QuestionnaireField).filter_by(schema_id=schema_for(v.id,s).id).all()
  try: ConditionEngine.validate_condition(p.include_when,fields)
  except ConditionConfigurationError as exc: raise HTTPException(422,detail=str(exc))
  x.region_key=p.region_key; x.include_when=p.include_when; s.flush()
  try: validate_regions(v,s,validate_bindings(v,s))
  except HTTPException: s.rollback(); raise
  s.commit(); return region_out(x)
 @r.delete(prefix+'/{tid}/versions/{vid}/conditional-regions/{rid}',status_code=204)
 def delete_region(tid:str,vid:str,rid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); x=s.get(ConditionalDocumentRegion,rid)
  if not x or x.template_version_id!=v.id: raise HTTPException(404,detail='Условная область не найдена')
  s.delete(x); s.commit()
 @r.get(prefix+'/{tid}/versions/{vid}/repeating-groups')
 def repeating_groups(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s)
  if t.id!=tid: raise HTTPException(404,detail='Версия не найдена')
  return [repeating_group_out(x,s) for x in s.query(RepeatingGroup).filter_by(template_version_id=v.id).order_by(RepeatingGroup.position).all()]
 @r.post(prefix+'/{tid}/versions/{vid}/repeating-groups',status_code=201)
 def create_repeating_group(tid:str,vid:str,p:RepeatingGroupPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v)
  if p.max_items<p.min_items or (p.required and p.max_items<1): raise HTTPException(422,detail='Максимальное количество должно быть не меньше минимального')
  if s.query(RepeatingGroup).filter_by(template_version_id=v.id,group_key=p.group_key).first(): raise HTTPException(409,detail='Группа с таким ключом уже существует')
  if s.query(RepeatingGroup).filter_by(template_version_id=v.id,position=p.position).first(): raise HTTPException(409,detail='Позиция группы уже занята')
  fields=s.query(QuestionnaireField).filter_by(schema_id=schema_for(v.id,s).id).all()
  try:
   if p.visible_when is not None: ConditionEngine.validate_condition(p.visible_when,fields)
   locator=_resolve_repeating_row(safe_source(v),p.row_locator,v.structure_fingerprint)
  except (ConditionConfigurationError,ValueError) as exc: raise HTTPException(422,detail=str(exc))
  x=RepeatingGroup(template_version_id=v.id,group_key=p.group_key,name=p.name.strip(),description=p.description.strip(),min_items=p.min_items,max_items=p.max_items,required=p.required,position=p.position,config={'visible_when':p.visible_when},locator=locator,created_by_id=u.id); s.add(x); s.flush(); s.add(TemplateAuditEvent(template_id=t.id,template_version_id=v.id,actor_id=u.id,event_type='REPEATING_GROUP_CREATED',metadata_json={'group_id':x.id})); s.commit(); return repeating_group_out(x,s)
 @r.patch(prefix+'/{tid}/versions/{vid}/repeating-groups/{gid}')
 def update_repeating_group(tid:str,vid:str,gid:str,p:RepeatingGroupUpdatePayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); x=s.get(RepeatingGroup,gid)
  if not x or x.template_version_id!=v.id: raise HTTPException(404,detail='Повторяемая группа не найдена')
  if p.max_items<p.min_items: raise HTTPException(422,detail='Максимальное количество должно быть не меньше минимального')
  duplicate=s.query(RepeatingGroup).filter(RepeatingGroup.template_version_id==v.id,RepeatingGroup.group_key==p.group_key,RepeatingGroup.id!=x.id).first(); occupied=s.query(RepeatingGroup).filter(RepeatingGroup.template_version_id==v.id,RepeatingGroup.position==p.position,RepeatingGroup.id!=x.id).first()
  if duplicate: raise HTTPException(409,detail='Группа с таким ключом уже существует')
  if occupied: raise HTTPException(409,detail='Позиция группы уже занята')
  fields=s.query(QuestionnaireField).filter_by(schema_id=schema_for(v.id,s).id).all()
  try:
   if p.visible_when is not None: ConditionEngine.validate_condition(p.visible_when,fields)
  except ConditionConfigurationError as exc: raise HTTPException(422,detail=str(exc))
  x.group_key=p.group_key; x.name=p.name.strip(); x.description=p.description.strip(); x.min_items=p.min_items; x.max_items=p.max_items; x.required=p.required; x.position=p.position; x.config={'visible_when':p.visible_when}; s.commit(); return repeating_group_out(x,s)
 @r.delete(prefix+'/{tid}/versions/{vid}/repeating-groups/{gid}',status_code=204)
 def delete_repeating_group(tid:str,vid:str,gid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); x=s.get(RepeatingGroup,gid)
  if not x or x.template_version_id!=v.id: raise HTTPException(404,detail='Повторяемая группа не найдена')
  s.query(RepeatingGroupField).filter_by(group_id=x.id).delete(); s.delete(x); s.commit()
 @r.post(prefix+'/{tid}/versions/{vid}/repeating-groups/{gid}/fields',status_code=201)
 def create_repeating_field(tid:str,vid:str,gid:str,p:RepeatingFieldPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); group=s.get(RepeatingGroup,gid); validate_repeating_field_payload(p)
  if not group or group.template_version_id!=v.id: raise HTTPException(404,detail='Повторяемая группа не найдена')
  if s.query(RepeatingGroupField).filter_by(group_id=gid,field_key=p.field_key).first(): raise HTTPException(409,detail='Поле с таким ключом уже существует')
  if s.query(RepeatingGroupField).filter_by(group_id=gid,position=p.position).first(): raise HTTPException(409,detail='Позиция поля уже занята')
  x=RepeatingGroupField(group_id=gid,field_key=p.field_key,label=p.label.strip(),field_type=p.field_type,required=p.required,position=p.position,config=repeating_field_config(p)); s.add(x); s.flush()
  try: validate_calculations(v.id,s)
  except HTTPException: s.rollback(); raise
  s.commit(); return repeating_field_out(x)
 @r.patch(prefix+'/{tid}/versions/{vid}/repeating-groups/{gid}/fields/{fid}')
 def update_repeating_field(tid:str,vid:str,gid:str,fid:str,p:RepeatingFieldPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); group=s.get(RepeatingGroup,gid); x=s.get(RepeatingGroupField,fid); validate_repeating_field_payload(p)
  if not group or group.template_version_id!=v.id or not x or x.group_id!=gid: raise HTTPException(404,detail='Поле группы не найдено')
  duplicate=s.query(RepeatingGroupField).filter(RepeatingGroupField.group_id==gid,RepeatingGroupField.field_key==p.field_key,RepeatingGroupField.id!=fid).first(); occupied=s.query(RepeatingGroupField).filter(RepeatingGroupField.group_id==gid,RepeatingGroupField.position==p.position,RepeatingGroupField.id!=fid).first()
  if duplicate: raise HTTPException(409,detail='Поле с таким ключом уже существует')
  if occupied: raise HTTPException(409,detail='Позиция поля уже занята')
  x.field_key=p.field_key; x.label=p.label.strip(); x.field_type=p.field_type; x.required=p.required; x.position=p.position; x.config=repeating_field_config(p); s.flush()
  try: validate_calculations(v.id,s)
  except HTTPException: s.rollback(); raise
  s.commit(); return repeating_field_out(x)
 @r.delete(prefix+'/{tid}/versions/{vid}/repeating-groups/{gid}/fields/{fid}',status_code=204)
 def delete_repeating_field(tid:str,vid:str,gid:str,fid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); group=s.get(RepeatingGroup,gid); x=s.get(RepeatingGroupField,fid)
  if not group or group.template_version_id!=v.id or not x or x.group_id!=gid: raise HTTPException(404,detail='Поле группы не найдено')
  s.delete(x); s.flush()
  try: validate_calculations(v.id,s)
  except HTTPException: s.rollback(); raise
  s.commit()
 @r.put(prefix+'/{tid}/versions/{vid}/repeating-groups/{gid}/fields/{fid}/binding')
 def bind_repeating_field(tid:str,vid:str,gid:str,fid:str,p:RepeatingBindingPayload,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); v,t=owned_version(vid,u,s); require_draft(v); group=s.get(RepeatingGroup,gid); field=s.get(RepeatingGroupField,fid)
  if not group or group.template_version_id!=v.id or not field or field.group_id!=gid: raise HTTPException(404,detail='Поле группы не найдено')
  try: loc=adapter.get_structural_locator(safe_source(v),adapter.select_range(safe_source(v),p.locator,p.start,p.end,p.selected_text))
  except ValueError as exc: raise HTTPException(422,detail=str(exc))
  if not _field_in_group_row(loc['node'],group): raise HTTPException(422,detail='Выделение должно находиться в строке-образце этой группы')
  loc['selected_text']=p.selected_text
  for other in s.query(RepeatingGroupField).filter(RepeatingGroupField.group_id==gid,RepeatingGroupField.id!=fid):
   old=other.locator or {}
   if old.get('node')==loc['node'] and loc['text_start']<old.get('text_end',0) and loc['text_end']>old.get('text_start',0): raise HTTPException(409,detail='Выделение пересекается с другой привязкой строки')
  field.locator=loc; s.commit(); return repeating_field_out(field)
 @r.post(prefix+'/{tid}/versions/{vid}/publish')
 def publish(tid:str,vid:str,u=Depends(current_dependency),s:DBS=Depends(db)):
  require_manager(u); require_publisher(u); v,t=owned_version(vid,u,s); require_draft(v); schema=schema_for(v.id,s)
  if not schema: raise HTTPException(422,detail='Создайте анкету')
  fields=s.query(QuestionnaireField).filter_by(schema_id=schema.id).all(); bindings=validate_bindings(v,s); regions=validate_regions(v,s,bindings); groups=validate_repeating_groups(v,s,bindings,regions)
  if not fields: raise HTTPException(422,detail='Добавьте хотя бы одно поле анкеты')
  validate_condition_schema(fields)
  validate_calculations(v.id,s)
  bound={x.field_id for x in bindings}
  if any((x.required or (x.config or {}).get('required_when') is not None) and x.id not in bound for x in fields): raise HTTPException(422,detail='Все обязательные поля должны иметь привязку в DOCX')
  if any(not x.field_id for x in bindings): raise HTTPException(422,detail='Найдены привязки без поля анкеты')
  probe={x.field_key:('Да' if x.field_type=='BOOLEAN' else '2026-01-01' if x.field_type=='DATE' else '1') for x in fields if not (x.config or {}).get('calculated')}; temp=Path(v.source_path).parent/'publish-check.docx'
  group_fields={x.id:s.query(RepeatingGroupField).filter_by(group_id=x.id).order_by(RepeatingGroupField.position).all() for x in groups}
  def samples(count):
   def sample_value(field):
    if field.field_type=='DATE': return '2026-01-01'
    if field.field_type=='BOOLEAN': return True
    if field.field_type in {'INTEGER','DECIMAL','MONEY'}: return 1
    if field.field_type=='SINGLE_SELECT': return (field.config or {}).get('options',['Тест'])[0]
    return 'Тест'
   return {g.group_key:[{'_id':f'preflight-{i}',**{f.field_key:sample_value(f) for f in group_fields[g.id] if not (f.config or {}).get('calculated')}} for i in range(count)] for g in groups}
  def checked_values(repeated):
   source={**probe,**repeated}; _,values,rows=calculated_document_values(v.id,s,fields,source); return formatted_bindings(bindings,values),rows
  try:
   first=samples(1); values,rows=checked_values(first); render_docx(v,bindings,values,temp,regions,{x.id for x in regions},groups,rows,group_fields); temp.unlink(missing_ok=True)
   zero_or_min={g.group_key:([] if g.min_items==0 and not g.required else samples(max(g.min_items,1))[g.group_key]) for g in groups}
   values,rows=checked_values(zero_or_min); render_docx(v,bindings,values,temp,regions,set(),groups,rows,group_fields); temp.unlink(missing_ok=True)
   three=samples(3); values,rows=checked_values(three); render_docx(v,bindings,values,temp,regions,{x.id for x in regions},groups,rows,group_fields); temp.unlink(missing_ok=True)
   if groups:
    maximum={g.group_key:samples(g.max_items)[g.group_key] for g in groups}
    values,rows=checked_values(maximum); render_docx(v,bindings,values,temp,regions,{x.id for x in regions},groups,rows,group_fields); temp.unlink(missing_ok=True)
  except (ValueError,IndexError,KeyError,CalculationError) as exc: raise HTTPException(422,detail=f'Проверка генерации не пройдена: {exc}')
  for old in s.query(TemplateVersion).filter(TemplateVersion.template_id==t.id,TemplateVersion.status=='PUBLISHED'): old.status='SUPERSEDED'
  v.status='PUBLISHED'; v.published_by_id=u.id; v.published_at=now(); schema.status='PUBLISHED'; t.status='ACTIVE'; t.current_version_id=v.id; s.add(TemplateAuditEvent(template_id=t.id,template_version_id=v.id,actor_id=u.id,event_type='VERSION_PUBLISHED')); s.commit(); return version_out(v)
 return r
