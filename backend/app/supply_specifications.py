"""Schema-driven supply specification dialogue and approved-template renderer."""
from __future__ import annotations

import re
import shutil
from copy import deepcopy
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable

from docx import Document


NUMBER_WORDS = {
    'один': 1, 'одна': 1, 'одно': 1, 'два': 2, 'две': 2, 'три': 3,
    'четыре': 4, 'пять': 5, 'шесть': 6, 'семь': 7, 'восемь': 8,
    'девять': 9, 'десять': 10, 'одиннадцать': 11, 'двенадцать': 12,
}
CONFIRMATIONS = {'да', 'верно', 'все верно', 'всё верно', 'подтверждаю'}
REJECTIONS = {'нет', 'не', 'неверно', 'не все верно', 'не всё верно', 'есть ошибки', 'хочу исправить', 'нужно изменить'}


def confirmation_intent(text: str) -> str:
    normalized = re.sub(r'[.!?]+$', '', re.sub(r'\s+', ' ', text.strip().lower().replace('ё', 'е')))
    if normalized in {value.replace('ё', 'е') for value in CONFIRMATIONS}: return 'CONFIRM'
    if normalized in {value.replace('ё', 'е') for value in REJECTIONS}: return 'REJECT_CONFIRMATION'
    if re.match(r'^(?:нет|неверно|не все верно|есть ошибки|хочу исправить|нужно изменить)\s*[,;:]\s*\S+', normalized): return 'EDIT_COMMAND'
    if any(word in normalized for word in ('поменяй', 'измени', 'исправь', 'убери', 'удали', 'добавь')): return 'EDIT_COMMAND'
    return 'UNKNOWN'


class SpecificationTemplateRegistry:
    templates = {'SUPPLY': 'supply/supply_specification.docx'}

    @classmethod
    def resolve(cls, root: Path, code: str = 'SUPPLY') -> Path:
        relative = cls.templates.get(code)
        if not relative:
            raise ValueError(f'Неизвестный шаблон ТЗ: {code}')
        path = (root / relative).resolve()
        if not path.is_relative_to(root.resolve()) or not path.is_file():
            raise FileNotFoundError(path)
        return path


class SpecificationRouter:
    def route(self, _: str = '') -> str:
        return 'SUPPLY'


def empty_supply_specification() -> dict[str, Any]:
    return {
        'kind': 'SUPPLY', 'template_code': 'SUPPLY', 'equipment_category': '',
        'subject': '', 'subject_locked': False, 'acceptance_criteria': [],
        'delivery_address': '', 'currency': '', 'vat_mode': '', 'vat_rate': None,
        'vat_basis': '', 'delivery_term': '', 'support_required': None,
        'support_services': '', 'support_term': '', 'support_start_condition': '',
        'warranty_service_description': '', 'warranty_terms_url': '', 'items': [],
        'total_amount': '0', 'vat_amount': '0', 'confirmed': False,
        'artifact_status': 'NONE', 'provenance': {}, 'previous_state': None,
    }


def parse_positive_integer(text: str) -> int | None:
    match = re.search(r'\b\d+\b', text.replace(' ', ' '))
    if match:
        value = int(match.group())
        return value if value > 0 else None
    for word, value in NUMBER_WORDS.items():
        if re.search(rf'\b{word}\b', text.lower()):
            return value
    return None


def parse_money(text: str) -> Decimal | None:
    value = text.lower().replace(' ', ' ')
    match = re.search(r'(\d[\d ]*(?:[,.]\d{1,2})?)', value)
    if not match:
        return None
    try:
        result = Decimal(match.group(1).replace(' ', '').replace(',', '.'))
    except InvalidOperation:
        return None
    if re.search(r'\bтыс(?:\.|яч[аи]?)?', value):
        result *= 1000
    return result if result > 0 else None


def money(value: Any) -> str:
    amount = Decimal(str(value or 0))
    return f'{amount:,.2f}'.replace(',', ' ').replace('.00', '').replace('.', ',')


def _decimal_string(value: Decimal) -> str:
    return format(value.quantize(Decimal('0.01')), 'f')


def recalculate(data: dict[str, Any]) -> dict[str, Any]:
    total = Decimal('0')
    for index, item in enumerate(data.get('items', []), 1):
        item['line_number'] = index
        price = Decimal(str(item.get('unit_price') or 0))
        quantity = Decimal(str(item.get('quantity') or 0))
        item['total_price'] = _decimal_string(price * quantity)
        total += price * quantity
    data['total_amount'] = _decimal_string(total)
    rate = Decimal(str(data.get('vat_rate') or 0))
    data['vat_amount'] = _decimal_string(total * rate / (Decimal('100') + rate)) if data.get('vat_mode') == 'VAT_PAYER' else '0'
    return data


def _clean(text: str) -> str:
    return re.sub(r'\s+', ' ', text).strip(' .')


def extract_items(text: str) -> list[dict[str, Any]]:
    """Extract one or more product clauses without owning dialogue state."""
    chunks = [part.strip() for part in re.split(r'\s+(?:и|а также)\s+(?=\d+\s)', text, flags=re.I) if part.strip()]
    items = []
    for chunk in chunks:
        quantity = parse_positive_integer(chunk)
        product = re.search(r'\b\d+\s*(?:шт\.?\s*)?(.+?)(?=,|\s+парт(?:-|\s)?номер|\s+гарант|\s+по\s+[\d ]|$)', chunk, re.I)
        if not product:
            continue
        description = _clean(product.group(1))
        part = re.search(r'парт(?:-|\s)?номер\s*[:№]?\s*([\w-]+)', chunk, re.I)
        warranty = re.search(r'гаранти[яию]?\s*(\d+\s*(?:месяц(?:ев|а)?|мес\.?|год(?:а|ов)?))', chunk, re.I)
        price = re.search(r'(?:по|цена(?:\s+за\s+единицу)?)[^\d]*(\d[\d ]*(?:[,.]\d{1,2})?\s*(?:тыс\.?)?(?:\s*руб(?:лей|ля|\.)?)?)', chunk, re.I)
        manufacturer = ''
        known = re.search(r'\b(Lenovo|Samsung|Apple|Huawei|Asus|Acer|Dell|HP|Xiaomi)\b', chunk, re.I)
        if known:
            manufacturer = known.group(1)
        item = {'line_number': len(items) + 1, 'part_number': part.group(1) if part else '',
                'description': description, 'manufacturer': manufacturer,
                'warranty_term': warranty.group(1) if warranty else '',
                'unit_price': _decimal_string(parse_money(price.group(1))) if price and parse_money(price.group(1)) else '',
                'quantity': quantity, 'total_price': '0'}
        items.append(item)
    return items


@dataclass(frozen=True)
class Question:
    question_id: str
    target: str
    parser: str
    text: str
    dependency: Callable[[dict[str, Any]], bool] = lambda _: True


BASE_QUESTIONS = (
    Question('SUPPLY_CATEGORY', 'equipment_category', 'TEXT', 'Что закупаем? Укажите категорию оборудования.'),
    Question('SUPPLY_ITEM_COUNT', 'item_count', 'INTEGER', 'Сколько разных видов товаров будет в ТЗ? Например, ноутбуки и мониторы — это 2 позиции.'),
    Question('SUPPLY_ADDRESS', 'delivery_address', 'TEXT', 'Укажите полный адрес доставки.'),
    Question('SUPPLY_CURRENCY', 'currency', 'CURRENCY', 'В какой валюте указаны цены? Например: рубли.'),
    Question('SUPPLY_VAT_MODE', 'vat_mode', 'VAT_MODE', 'Цены указаны с НДС или без НДС?'),
    Question('SUPPLY_VAT_RATE', 'vat_rate', 'RATE', 'Укажите ставку НДС, например 22%.', lambda d: d.get('vat_mode') == 'VAT_PAYER'),
    Question('SUPPLY_VAT_BASIS', 'vat_basis', 'TEXT', 'Укажите основание, по которому НДС не облагается.', lambda d: d.get('vat_mode') == 'NO_VAT'),
    Question('SUPPLY_DELIVERY_TERM', 'delivery_term', 'TEXT', 'Укажите срок поставки товара.'),
    Question('SUPPLY_SUPPORT_REQUIRED', 'support_required', 'BOOLEAN', 'Нужна ли техническая поддержка производителя?'),
    Question('SUPPLY_SUPPORT_SERVICES', 'support_services', 'TEXT', 'Перечислите услуги технической поддержки.', lambda d: d.get('support_required') is True),
    Question('SUPPLY_SUPPORT_TERM', 'support_term', 'TEXT', 'Укажите срок технической поддержки.', lambda d: d.get('support_required') is True),
    Question('SUPPLY_SUPPORT_START', 'support_start_condition', 'TEXT', 'С какого события начинается срок технической поддержки?', lambda d: d.get('support_required') is True),
    Question('SUPPLY_WARRANTY_SERVICE', 'warranty_service_description', 'TEXT', 'Кратко опишите гарантийное обслуживание.'),
    Question('SUPPLY_WARRANTY_URL', 'warranty_terms_url', 'URL', 'Укажите ссылку на условия гарантийного обслуживания производителя.'),
)


def _item_question(index: int, field: str) -> Question:
    labels = {
        'description': ('TEXT', 'описание товара'),
        'manufacturer': ('TEXT', 'производителя'), 'warranty_term': ('WARRANTY', 'срок гарантии'),
        'unit_price': ('MONEY', 'цену за единицу'), 'quantity': ('INTEGER', 'количество'),
    }
    parser, label = labels[field]
    return Question(f'SUPPLY_ITEM_{index + 1}_{field.upper()}', f'items[{index}].{field}', parser,
                    f'Позиция {index + 1}: укажите {label}.')


def _get_target(data: dict[str, Any], target: str) -> Any:
    match = re.fullmatch(r'items\[(\d+)]\.(\w+)', target)
    if match:
        return data['items'][int(match.group(1))].get(match.group(2))
    return data.get(target)


def _set_target(data: dict[str, Any], target: str, value: Any) -> None:
    match = re.fullmatch(r'items\[(\d+)]\.(\w+)', target)
    if match:
        data['items'][int(match.group(1))][match.group(2)] = value
    else:
        data[target] = value


def _parse(parser: str, text: str) -> tuple[Any, str | None]:
    raw = text.strip()
    low = raw.lower()
    if parser == 'INTEGER':
        value = parse_positive_integer(raw)
        return (value, None) if value else (None, 'Укажите целое число больше нуля.')
    if parser == 'MONEY':
        value = parse_money(raw)
        return (_decimal_string(value), None) if value else (None, 'Укажите положительную цену, например 100 000 руб.')
    if parser == 'BOOLEAN':
        if low in {'да', 'нужна', 'нужно', 'требуется'}: return True, None
        if low in {'нет', 'не нужна', 'не нужно', 'не требуется'}: return False, None
        return None, 'Ответьте «Да» или «Нет».'
    if parser == 'VAT_MODE':
        if 'без ндс' in low: return 'NO_VAT', None
        if 'с ндс' in low or low == 'ндс': return 'VAT_PAYER', None
        return None, 'Укажите «С НДС» или «Без НДС».'
    if parser == 'RATE':
        match = re.search(r'(\d+(?:[,.]\d+)?)\s*%', raw)
        return (float(match.group(1).replace(',', '.')), None) if match else (None, 'Укажите ставку в процентах, например 22%.')
    if parser == 'CURRENCY':
        if any(x in low for x in ('руб', 'rub', '₽')): return 'рублей', None
        if any(x in low for x in ('доллар', 'usd', '$')): return 'долларов США', None
        if any(x in low for x in ('евро', 'eur', '€')): return 'евро', None
        return None, 'Укажите валюту: рубли, доллары США или евро.'
    if parser == 'URL':
        match = re.search(r'https?://\S+', raw)
        return (match.group().rstrip('.,'), None) if match else (None, 'Укажите полную ссылку, начинающуюся с http:// или https://.')
    if parser == 'WARRANTY':
        match = re.search(r'\b\d+\s*(?:месяц(?:ев|а)?|мес\.?|год(?:а|ов)?)\b', low)
        return (match.group(), None) if match else (None, 'Укажите срок гарантии, например 12 месяцев или 3 года.')
    value = _clean(raw)
    return (value, None) if value else (None, 'Значение не должно быть пустым.')


def _ensure_items(data: dict[str, Any], count: int) -> None:
    while len(data['items']) < count:
        data['items'].append({'line_number': len(data['items']) + 1, 'part_number': '', 'description': '',
                              'manufacturer': '', 'warranty_term': '', 'unit_price': '',
                              'quantity': None, 'total_price': '0'})
    del data['items'][count:]


def next_question(data: dict[str, Any]) -> Question | None:
    if not data.get('equipment_category'):
        return BASE_QUESTIONS[0]
    if not data.get('items'):
        return BASE_QUESTIONS[1]
    for index, item in enumerate(data['items']):
        for field in ('description', 'manufacturer', 'warranty_term', 'unit_price', 'quantity'):
            if item.get(field) in (None, ''):
                return _item_question(index, field)
    for question in BASE_QUESTIONS[2:]:
        if question.dependency(data) and _get_target(data, question.target) in (None, ''):
            return question
    return None


def summary(data: dict[str, Any]) -> str:
    vat = f"НДС {data.get('vat_rate')}%" if data.get('vat_mode') == 'VAT_PAYER' else f"Без НДС: {data.get('vat_basis')}"
    support = 'требуется' if data.get('support_required') else 'не требуется'
    items = '\n'.join(f"{index}. {item.get('description','')}\n   Количество: {item.get('quantity','—')}\n   Цена: {money(item.get('unit_price'))} {data.get('currency','')}\n   Гарантия: {item.get('warranty_term') or 'не указана'}" for index,item in enumerate(data['items'],1))
    criteria = '; '.join(data.get('acceptance_criteria') or []) or 'не указаны'
    return (f"Проверьте ТЗ:\n\nПредмет закупки: {data.get('subject') or 'не указан'}\n\nПозиции:\n{items}\n\n"
            f"Адрес доставки: {data['delivery_address']}\nСрок поставки: {data['delivery_term']}\n"
            f"НДС: {vat}\nТехническая поддержка: {support}\nКритерии приёмки: {criteria}\n"
            f"Общая стоимость: {money(data['total_amount'])} {data['currency']}\n\nВсё верно?")


def supply_from_import(generic: dict[str, Any], text: str = '') -> dict[str, Any]:
    """Convert imported legacy extraction into the canonical supply state."""
    data = empty_supply_specification()
    generic_items = generic.get('items') or []
    data['items'] = [{
        'line_number': index, 'part_number': item.get('part_number',''),
        'description': item.get('description') or item.get('name') or '',
        'manufacturer': item.get('manufacturer',''),
        'warranty_term': item.get('warranty_term') or item.get('warranty') or '',
        'unit_price': str(item.get('unit_price') or ''), 'quantity': item.get('quantity'),
        'total_price': str(item.get('total_price') or '0'),
    } for index,item in enumerate(generic_items,1)]
    subject = generic.get('subject') or (data['items'][0]['description'] if data['items'] else '')
    data.update(subject=subject,subject_locked=bool(subject),equipment_category=subject,
                delivery_term=generic.get('delivery_term') or generic.get('delivery_terms') or generic.get('deadline') or '',
                acceptance_criteria=generic.get('acceptance_criteria') if isinstance(generic.get('acceptance_criteria'),list) else ([generic['acceptance_criteria']] if generic.get('acceptance_criteria') else []))
    address = re.search(r'(?:адрес\s+доставки|поставка\s+по\s+адресу)\s*[:—-]?\s*([^\n.;]+)', text, re.I)
    if address: data['delivery_address'] = address.group(1).strip()
    vat = re.search(r'ндс\s*(\d+(?:[,.]\d+)?)\s*%', text, re.I)
    if vat: data.update(vat_mode='VAT_PAYER',vat_rate=float(vat.group(1).replace(',','.')))
    elif re.search(r'без\s+ндс',text,re.I): data.update(vat_mode='NO_VAT',vat_basis='Указано в исходном ТЗ')
    if re.search(r'\b(?:руб(?:ль|ля|лей|\.)?|₽)\b',text,re.I): data['currency']='рублей'
    data['provenance']={key:'EXTRACTED' for key,value in data.items() if value not in ('',None,[],False)}
    return recalculate(data)


def _snapshot(data: dict[str, Any]) -> dict[str, Any]:
    snapshot = deepcopy(data)
    snapshot['previous_state'] = None
    return snapshot


class SupplySpecificationUpdateService:
    """The single sparse-update path used by chat and HTTP editing."""
    editable_fields = {
        'equipment_category', 'subject', 'delivery_address', 'currency', 'vat_mode',
        'vat_rate', 'vat_basis', 'delivery_term', 'support_required', 'support_services',
        'support_term', 'support_start_condition', 'warranty_service_description',
        'warranty_terms_url', 'acceptance_criteria', 'items',
    }

    def apply(self, data: dict[str, Any], changes: list[dict[str, Any]], source: str = 'USER_EXPLICIT') -> dict[str, Any]:
        updated = deepcopy(data)
        updated.setdefault('provenance', {})
        updated.setdefault('acceptance_criteria', [])
        updated.setdefault('artifact_status', 'NONE')
        before = _snapshot(data)
        changed = False
        for change in changes:
            field = change.get('field', '')
            operation = change.get('operation', 'replace')
            value = change.get('value')
            item_match = re.fullmatch(r'items\[(\d+)]\.(description|part_number|manufacturer|warranty_term|unit_price|quantity)', field)
            if item_match:
                index, item_field = int(item_match.group(1)), item_match.group(2)
                if not 0 <= index < len(updated['items']): raise ValueError('Позиция не найдена')
                if item_field == 'quantity':
                    value = parse_positive_integer(str(value))
                    if value is None: raise ValueError('Количество должно быть целым числом больше нуля')
                elif item_field == 'unit_price':
                    parsed = parse_money(str(value))
                    if parsed is None: raise ValueError('Цена должна быть положительным числом')
                    value = _decimal_string(parsed)
                updated['items'][index][item_field] = '' if operation == 'delete' else value
                updated['provenance'][field] = source; changed = True; continue
            if field == 'items.add':
                item = deepcopy(value or {})
                item.setdefault('part_number', '')
                item.setdefault('total_price', '0')
                updated['items'].append(item); updated['provenance'][f'items[{len(updated["items"])-1}]'] = source; changed = True; continue
            if field == 'items.remove':
                index = int(value)
                if not 0 <= index < len(updated['items']): raise ValueError('Позиция не найдена')
                updated['items'].pop(index); changed = True; continue
            if field not in self.editable_fields: raise ValueError(f'Недопустимое поле: {field}')
            if field == 'acceptance_criteria':
                if isinstance(value, str): value = [part.strip(' -•') for part in re.split(r'[;\n]+', value) if part.strip(' -•')]
                if not isinstance(value, list): raise ValueError('Критерии приёмки должны быть текстом или списком')
            if field == 'subject' and updated.get('subject_locked') and source != 'USER_EXPLICIT':
                continue
            updated[field] = ([] if field in {'items', 'acceptance_criteria'} else '') if operation == 'delete' else value
            updated['provenance'][field] = source
            if field == 'subject' and source == 'USER_EXPLICIT': updated['subject_locked'] = True
            changed = True
        if changed:
            updated['previous_state'] = before
            updated['confirmed'] = False
            updated['artifact_status'] = 'STALE' if data.get('artifact_status') == 'CURRENT' else data.get('artifact_status', 'NONE')
        return recalculate(updated)

    def undo(self, data: dict[str, Any]) -> dict[str, Any]:
        previous = data.get('previous_state')
        if not previous: raise ValueError('Нет изменения для отмены')
        restored = deepcopy(previous); restored['previous_state'] = None
        restored['confirmed'] = False
        restored['artifact_status'] = 'STALE' if data.get('artifact_status') == 'CURRENT' else restored.get('artifact_status', 'NONE')
        return recalculate(restored)


class SupplySpecificationDialogueEngine:
    def __init__(self):
        self.updates = SupplySpecificationUpdateService()

    def process(self, data: dict[str, Any] | None, text: str, pending: dict[str, Any] | None) -> tuple[dict[str, Any], dict[str, Any] | None, str, bool]:
        data = deepcopy(data or empty_supply_specification())
        # Keep older/in-progress JSON documents compatible now that part number is optional.
        for item in data.get('items', []):
            item.setdefault('part_number', '')
        normalized = re.sub(r'[.!?]+$', '', text.strip().lower())
        intent = confirmation_intent(text)
        if next_question(data) is None and intent == 'REJECT_CONFIRMATION':
            data['confirmed'] = False
            return data, None, 'Хорошо. Что нужно изменить в ТЗ?', False
        if next_question(data) is None and intent == 'EDIT_COMMAND' and re.match(r'^\s*(?:нет|неверно|не\s+вс[её]\s+верно|есть ошибки)\s*[,;:]', text, re.I):
            text = re.sub(r'^\s*(?:нет|неверно|не\s+вс[её]\s+верно|есть ошибки)\s*[,;:]\s*', '', text, flags=re.I)
            text = 'Исправь ' + text
        if normalized in {'отмени последнее изменение', 'верни как было', 'откатить последнее изменение'}:
            try: data = self.updates.undo(data)
            except ValueError as exc: return data, pending, str(exc) + '.', False
            question = next_question(data)
            return data, self.pending(question), 'Последнее изменение отменено.' + (f'\n\n{question.text}' if question else '\n\n' + summary(data)), False
        if normalized in CONFIRMATIONS and next_question(data) is None:
            data['confirmed'] = True
            return recalculate(data), None, 'Техническое задание готово. Сформируйте и скачайте DOCX.', True

        edit = self._parse_edit(data, text)
        if edit:
            changes, edit_message = edit
            try: data = self.updates.apply(data, changes)
            except ValueError as exc: return data, pending, str(exc) + '.', False
            question = next_question(data)
            tail = f'\n\n{question.text}' if question else '\n\nМожете внести ещё правки или подтвердить ТЗ.'
            return data, self.pending(question), edit_message + '\nОстальные данные ТЗ оставил без изменений.' + tail, False

        extracted = extract_items(text)
        if extracted and (not data.get('items') or not pending or pending.get('target') in {'equipment_category', 'item_count'}):
            data['items'] = extracted
            if not data.get('equipment_category'):
                data['equipment_category'] = 'компьютерное оборудование' if any(re.search(r'ноутбук|монитор|компьютер', x['description'], re.I) for x in extracted) else extracted[0]['description']
                data['subject'] = 'Поставка ' + data['equipment_category']
                data['subject_locked'] = True
                data['provenance']['subject'] = 'EXTRACTED'

        if pending:
            target, parser = pending['target'], pending['parser']
            if target == 'item_count':
                count = parse_positive_integer(text)
                if not count:
                    return data, pending, 'Укажите целое количество позиций больше нуля.', False
                _ensure_items(data, count)
            elif not (extracted and target in {'equipment_category', 'item_count'}):
                value, error = _parse(parser, text)
                if error:
                    return data, pending, error, False
                _set_target(data, target, value)
                data['provenance'][target] = 'USER_EXPLICIT'
        elif not data.get('equipment_category') and not extracted:
            data['equipment_category'] = _clean(text)
            data['subject'] = 'Поставка ' + data['equipment_category']
            data['subject_locked'] = True
            data['provenance'].update(equipment_category='USER_EXPLICIT', subject='DERIVED')

        recalculate(data)
        question = next_question(data)
        if question:
            return data, self.pending(question), question.text, False
        return data, None, summary(data), False

    @staticmethod
    def pending(question: Question | None) -> dict[str, Any] | None:
        return None if not question else {'question_id': question.question_id, 'target': question.target,
                                          'parser': question.parser, 'question': question.text}

    def _parse_edit(self, data: dict[str, Any], text: str) -> tuple[list[dict[str, Any]], str] | None:
        low = text.lower()
        if not any(word in low for word in ('поменяй', 'измени', 'исправь', 'убери', 'удали', 'добавь')):
            return None
        remove = re.search(r'(?:убери|удали)\s+(?:позицию\s*)?(\d+|перв\w+|втор\w+|треть\w+)', low)
        if remove:
            token = remove.group(1); index = int(token) - 1 if token.isdigit() else {'перв': 0, 'втор': 1, 'трет': 2}[token[:4]]
            if 0 <= index < len(data['items']):
                return [{'field':'items.remove','value':index}], 'Готово. Позиция удалена.'
        subject = re.search(r'(?:теперь закупаем|предмет закупки\s*[:—-]?|замени предмет на)\s*(.+)', text, re.I)
        if subject:
            return [{'field':'subject','value':_clean(subject.group(1))}], 'Готово. Предмет закупки изменён.'
        address = re.search(r'(?:адрес(?:\s+доставки)?)(?:\s+на|\s*[:—-]|\s+должен\s+быть)\s*(.+)', text, re.I)
        if address:
            return [{'field':'delivery_address','value':address.group(1).strip()}], 'Готово. Адрес доставки изменён.'
        delivery = re.search(r'(?:срок поставки)(?:\s+на|\s*[:—-]|\s+должен\s+быть)\s*(.+)', text, re.I)
        if delivery:
            return [{'field':'delivery_term','value':delivery.group(1).strip(' .')}], 'Готово. Срок поставки изменён.'
        criterion = re.search(r'(?:критери[йяи]\s+при[её]мки)(?:\s*[:—-])?\s*(.+)', text, re.I)
        if criterion:
            existing = list(data.get('acceptance_criteria') or []); existing.append(criterion.group(1).strip(' .'))
            return [{'field':'acceptance_criteria','value':existing}], 'Готово. Критерии приёмки обновлены.'
        warranty = re.search(r'(?:гаранти[юя]|срок гарантии)(?:\s+на|\s*[:—-])\s*(\d+\s*(?:месяц(?:ев|а)?|мес\.?|год(?:а|ов)?))', text, re.I)
        if warranty and data['items']:
            index_match=re.search(r'(перв\w+|втор\w+|треть\w+|\d+)\s+позици',low)
            index=0 if not index_match else (int(index_match.group(1))-1 if index_match.group(1).isdigit() else {'перв':0,'втор':1,'трет':2}[index_match.group(1)[:4]])
            return [{'field':f'items[{index}].warranty_term','value':warranty.group(1)}], 'Готово. Срок гарантии изменён.'
        quantity = re.search(r'(?:количеств\w*(?:\s+(?:перв\w+|втор\w+|треть\w+|[а-яё-]+)\s*(?:позици\w*)?)?(?:\s+должно\s+быть)?|на)\s+(\d+)\s*(?:шт\.?)?', low)
        if quantity and data['items']:
            ordinal=re.search(r'(перв\w+|втор\w+|треть\w+)\s+позици',low)
            descriptor = re.search(r'(?:количеств\w*)\s+(.+?)\s+на\s+\d+', low)
            index = ({'перв':0,'втор':1,'трет':2}[ordinal.group(1)[:4]] if ordinal else next((i for i, item in enumerate(data['items']) if descriptor and descriptor.group(1) in item.get('description', '').lower()), 0))
            return [{'field':f'items[{index}].quantity','value':int(quantity.group(1))}], 'Готово. Количество позиции изменено.'
        if 'добавь' in low:
            additions = extract_items(text)
            if additions:
                return ([{'field':'items.add','value':item} for item in additions], 'Готово. Новая позиция добавлена.')
        return None


def _replace_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]: run.text = ''
    else:
        paragraph.add_run(text)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


ONES = ['', 'один', 'два', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
TEENS = {10:'десять',11:'одиннадцать',12:'двенадцать',13:'тринадцать',14:'четырнадцать',15:'пятнадцать',16:'шестнадцать',17:'семнадцать',18:'восемнадцать',19:'девятнадцать'}
TENS = ['', '', 'двадцать', 'тридцать', 'сорок', 'пятьдесят', 'шестьдесят', 'семьдесят', 'восемьдесят', 'девяносто']
HUNDREDS = ['', 'сто', 'двести', 'триста', 'четыреста', 'пятьсот', 'шестьсот', 'семьсот', 'восемьсот', 'девятьсот']


def _under_thousand(number: int, feminine: bool = False) -> str:
    words = [HUNDREDS[number // 100]]; remainder = number % 100
    if 10 <= remainder <= 19: words.append(TEENS[remainder])
    else:
        words.extend((TENS[remainder // 10], ONES[remainder % 10]))
    if feminine and remainder % 10 in (1, 2) and not 10 <= remainder <= 19:
        words[-1] = 'одна' if remainder % 10 == 1 else 'две'
    return ' '.join(word for word in words if word)


def amount_in_words(value: Any) -> str:
    number = int(Decimal(str(value or 0)))
    if number == 0: return 'ноль'
    groups = ((1_000_000, 'миллион', 'миллиона', 'миллионов', False), (1000, 'тысяча', 'тысячи', 'тысяч', True), (1, '', '', '', False))
    result = []
    for divisor, one, few, many, feminine in groups:
        part, number = divmod(number, divisor)
        if not part: continue
        result.append(_under_thousand(part, feminine))
        if divisor > 1:
            ending = one if part % 10 == 1 and part % 100 != 11 else few if part % 10 in (2,3,4) and not 12 <= part % 100 <= 14 else many
            result.append(ending)
    return ' '.join(result)


class SupplySpecificationRenderer:
    def __init__(self, template_root: Path):
        self.template_root = template_root

    def render(self, data: dict[str, Any], output: Path) -> Path:
        source = SpecificationTemplateRegistry.resolve(self.template_root)
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, output)
        document = Document(output)
        table = document.tables[0]
        while len(data['items']) > len(table.rows) - 3:
            total_row = table.rows[-2]._tr
            total_row.addprevious(deepcopy(table.rows[-3]._tr))
        item_rows = table.rows[1:1 + len(data['items'])]
        vat_text = f"с НДС {data['vat_rate']}%" if data['vat_mode'] == 'VAT_PAYER' else f"без НДС ({data['vat_basis']})"
        currency = data['currency']
        for row, item in zip(item_rows, data['items']):
            values = (item['line_number'], item['part_number'], item['description'], item['manufacturer'], item['warranty_term'],
                      f"{money(item['unit_price'])} {currency}, {vat_text}", item['quantity'],
                      f"{money(item['total_price'])} {currency}, {vat_text}")
            for cell, value in zip(row.cells, values): cell.text = str(value)
        for row in table.rows[1 + len(data['items']):-2]:
            for cell in row.cells: cell.text = ''
        total_row = table.rows[-2]
        total_row.cells[0].text = f'ИТОГО {vat_text}, {currency}:'
        total_row.cells[-1].text = f"{money(data['total_amount'])} {currency}"
        if data['vat_mode'] == 'VAT_PAYER':
            vat_row = table.rows[-1]
            vat_row.cells[0].text = f"В том числе НДС {data['vat_rate']}% {currency}:"
            vat_row.cells[-1].text = f"{money(data['vat_amount'])} {currency}"
        else:
            table._tbl.remove(table.rows[-1]._tr)

        table.rows[0].cells[5].text = f'Цена за ед., {vat_text}, {currency}'
        table.rows[0].cells[7].text = f'Стоимость, {vat_text}, {currency}'
        paragraphs = list(document.paragraphs)
        for paragraph in paragraphs:
            text = paragraph.text
            if text.startswith('Поставщик обязуется поставить'):
                suffix = ' и Сертификат(-ы) на техническую поддержку' if data['support_required'] else ''
                _replace_paragraph(paragraph, f"Поставщик обязуется поставить Покупателю следующее оборудование категории «{data['equipment_category']}»{suffix}, далее — «Товар», по адресу: {data['delivery_address']}:")
            elif text.startswith('Общая стоимость Товара составляет'):
                vat_clause = f"в том числе НДС ({data['vat_rate']}%) — {money(data['vat_amount'])} {currency}" if data['vat_mode'] == 'VAT_PAYER' else f"НДС не облагается: {data['vat_basis']}"
                support_clause = ', а также техническую поддержку Производителя' if data['support_required'] else ''
                _replace_paragraph(paragraph, f"Общая стоимость Товара составляет {money(data['total_amount'])} ({amount_in_words(data['total_amount'])}) {currency}, {vat_clause}, включая доставку, погрузку/разгрузку и гарантийное обслуживание{support_clause}.")
            elif text.startswith('[указать общий перечень услуг'):
                if data['support_required']: _replace_paragraph(paragraph, f"Техническая поддержка включает: {data['support_services']}.")
                else: _remove_paragraph(paragraph)
            elif text.startswith('Срок поставки Товара:'):
                _replace_paragraph(paragraph, f"Срок поставки Товара: {data['delivery_term']}.")
            elif text.startswith('Срок действия Сертификата'):
                if data['support_required']: _replace_paragraph(paragraph, f"Срок технической поддержки: {data['support_term']}; начало срока: {data['support_start_condition']}.")
                else: _remove_paragraph(paragraph)
            elif text.startswith('Гарантийное обслуживание Товара в течение'):
                _replace_paragraph(paragraph, f"Гарантийное обслуживание Товара в течение срока, указанного в Таблице № 1, включает: {data['warranty_service_description']}. Условия размещены на сайте Производителя: {data['warranty_terms_url']}.")
            elif text.startswith('Поставщик обязуется предоставить/обеспечить') and not data['support_required']:
                _replace_paragraph(paragraph, 'Поставщик обязуется обеспечить гарантийное обслуживание поставленного Товара на условиях Договора и настоящей Спецификации в течение всего гарантийного срока.')
        document.save(output)
        self._validate(document)
        return output

    @staticmethod
    def _validate(document) -> None:
        text = '\n'.join([*(p.text for p in document.paragraphs), *(cell.text for table in document.tables for row in table.rows for cell in row.cells)])
        unresolved = [token for token in re.findall(r'\[[^\]]+\]', text) if token not in {'[__]', '[__________]'}]
        if unresolved:
            raise ValueError(f'В документе остались placeholders: {sorted(set(unresolved))}')
