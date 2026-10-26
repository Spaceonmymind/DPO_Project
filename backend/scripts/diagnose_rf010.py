"""Generate isolated, non-official RF010 diagnostics from the approved source DOCX."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document

from app.contract_platform import calculated_document_values, formatted_bindings, render_docx, safe_source
from app.db import DocumentBinding, QuestionnaireField, QuestionnaireSchema, SessionLocal, TemplateVersion


parser = argparse.ArgumentParser()
parser.add_argument("--output-dir", type=Path, required=True)
args = parser.parse_args()
args.output_dir.mkdir(parents=True, exist_ok=True)
session = SessionLocal()
try:
    version = session.get(TemplateVersion, "4eff6663-f03e-4858-bb2a-6dc39b4b1b95")
    schema = session.query(QuestionnaireSchema).filter_by(template_version_id=version.id).first()
    fields = session.query(QuestionnaireField).filter_by(schema_id=schema.id).all()
    bindings = session.query(DocumentBinding).filter(DocumentBinding.template_version_id == version.id, DocumentBinding.binding_key.in_(["ip.invoice_clause", "ip.ogrnip"]), DocumentBinding.status == "ACTIVE").all()
    assert {binding.binding_key for binding in bindings} == {"ip.invoice_clause", "ip.ogrnip"}
    original_sha = hashlib.sha256(safe_source(version).read_bytes()).hexdigest()
    expected = {
        "VAT_PAYER": "Счета-фактуры выставляются Исполнителем в соответствии с требованиями законодательства Российской Федерации.",
        "NO_VAT": "Если Исполнитель не является плательщиком НДС, счет-фактура не выставляется.",
    }
    for status, phrase in expected.items():
        _, calculated, _ = calculated_document_values(version.id, session, fields, {"tax.vat_status": status, "ip.ogrnip": "123456789012345"})
        assert calculated["ip.invoice_clause"] == phrase
        output = args.output_dir / f"rf010-{status.lower()}.docx"
        render_docx(version, bindings, formatted_bindings(bindings, calculated, fields), output)
        document = Document(output)
        found = [paragraph.text for paragraph in document.paragraphs if phrase in paragraph.text]
        assert len(found) == 1, (status, found)
        assert any("ОГРНИП 123456789012345" in paragraph.text for table in document.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs)
        assert "Если Исполнитель не является плательщиком НДС, счет-фактура не выставляется." not in found[0] if status == "VAT_PAYER" else "Счета-фактуры выставляются Исполнителем" not in found[0]
        print(json.dumps({"status": status, "calculated": calculated["ip.invoice_clause"], "paragraph": found[0], "path": str(output), "sha256": hashlib.sha256(output.read_bytes()).hexdigest()}, ensure_ascii=False))
    assert hashlib.sha256(safe_source(version).read_bytes()).hexdigest() == original_sha
finally:
    session.close()
