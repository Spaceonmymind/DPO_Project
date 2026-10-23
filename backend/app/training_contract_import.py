"""Idempotent, fail-closed R3.6.1 import of the training-product draft."""
from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import load_workbook

from app.db import (
    ContractProduct, ContractProductAuditEvent, ContractProductField,
    ContractProductVersion, DocumentTemplate, SessionLocal, TemplateVersion,
    QuestionnaireSchema,QuestionnaireField,settings,
)
from app.training_contract_audit import EXPECTED, audit_training_sources
from app.training_mapping import discover_candidates

PRODUCT_CODE = "TRAINING_CONTRACT"
PRODUCT_NAME = "Договор на проведение тренировок"
UNSUPPORTED_FACT_TYPES = {"ARRAY_STRING", "ARRAY_OBJECT"}
TYPE_MAP = {"ENUM": "SINGLE_SELECT", "BOOLEAN": "BOOLEAN", "STRING": "TEXT", "DECIMAL": "DECIMAL", "INTEGER": "INTEGER", "DATE": "DATE"}
RF_TYPE_MAP={"SELECT":"SINGLE_SELECT","ARRAY":"LIST","OBJECT":"COMPOSITE","BOOLEAN_BLOCK":"BOOLEAN","INTEGER":"INTEGER","STRING":"TEXT","TEXT":"TEXT","DERIVED":"TEXT","BLOCK":"BOOLEAN"}


def _options(raw):
    values = [x.strip() for x in str(raw or "").split("|")]
    return [x for x in values if x and x not in {"true", "false", "null"}]


def build_plan(root: Path) -> dict:
    audit = audit_training_sources(root)
    workbook = load_workbook(root / "Матрица_применимости_LLM_v2.xlsx", data_only=True)
    fields = []
    blocked = []
    for position, row in enumerate(workbook["FACT_SCHEMA"].iter_rows(min_row=2, values_only=True), 1):
        key, source_type, requirement, allowed, meaning, question, default = row
        if source_type in UNSUPPORTED_FACT_TYPES:
            blocked.append({"key": key, "reason": f"Тип {source_type} не поддержан универсальным полем"})
            continue
        kind = TYPE_MAP[source_type]
        config = {"group": "Применимость" if position <= 12 else "Условия договора", "description": meaning or "", "source_requirement": requirement}
        if kind == "SINGLE_SELECT": config["options"] = _options(allowed)
        if key == "ip.ogrnip": config["validation"] = {"regex": "^[0-9]{15}$"}
        fields.append({"field_key": key, "label": question if question and question != "—" else meaning or key, "field_type": kind, "required": requirement in {"CRITICAL", "REQUIRED_TO_RENDER"}, "position": position, "config": config})
    blockers = [
        {"code": "DRAFT_TEMPLATE_VERSIONS", "severity": "P0", "detail": "Оба реальных TemplateVersion не опубликованы; варианты продукта создавать нельзя"},
        {"code": "CANDIDATE_REVIEW", "severity": "P0", "detail": "Структурные кандидаты требуют решения специалиста; неразрешённые места нельзя материализовать автоматически"},
        {"code": "REAL_BINDINGS_MISSING", "severity": "P0", "detail": "140 физических вхождений ещё не имеют проверенных runtime bindings/regions/groups"},
        {"code": "CONTROL_GENERATION", "severity": "P0", "detail": "Оба варианта ещё не прошли draft generation и OOXML/Word acceptance"},
    ]
    signature = hashlib.sha256((audit["files"]["Матрица_применимости_LLM_v2.xlsx"]["sha256"] + PRODUCT_CODE).encode()).hexdigest()
    candidates=discover_candidates(root)
    return {"signature": signature, "audit": audit, "fields": fields, "candidates": candidates, "unsupported_facts": blocked, "blockers": blockers, "source_counts": {"facts": 40, "rf_fields": 25, "occurrences": 74, "physical_occurrences":len(candidates), "scenarios": 120}}


def apply_plan(plan: dict, owner, session=None) -> dict:
    own_session = session is None
    session = session or SessionLocal()
    try:
        product = session.query(ContractProduct).filter_by(owner_id=owner.id, code=PRODUCT_CODE).first()
        created_product = product is None
        if product is None:
            product = ContractProduct(owner_id=owner.id, code=PRODUCT_CODE, name=PRODUCT_NAME, category="Усуги", description="Черновой пилот по матрице R3.6.1")
            session.add(product); session.flush()
            version = ContractProductVersion(product_id=product.id, version_number=1, created_by_id=owner.id, rules={"import_signature": plan["signature"], "blocking_rules": [], "manual_review_rules": [], "test_scenarios": []})
            session.add(version); session.flush()
        else:
            version = session.query(ContractProductVersion).filter_by(product_id=product.id, status="DRAFT").order_by(ContractProductVersion.version_number.desc()).first()
            if version is None:
                raise ValueError("У продукта нет черновой версии; опубликованные версии не изменяются")
        existing = {x.field_key for x in session.query(ContractProductField).filter_by(product_version_id=version.id)}
        created = 0
        for item in plan["fields"]:
            if item["field_key"] in existing: continue
            session.add(ContractProductField(product_version_id=version.id, **item)); created += 1
        if created_product:
            session.add(ContractProductAuditEvent(product_id=product.id, product_version_id=version.id, actor_id=owner.id, event_type="R361_DRAFT_IMPORTED", metadata_json={"signature": plan["signature"]}))
        candidate_counts={};template_field_counts={}
        workbook=load_workbook(Path(settings.contract_templates_path)/"training"/"Матрица_применимости_LLM_v2.xlsx",data_only=True)
        rf_rows=list(workbook["RED_FIELD_SCHEMA"].iter_rows(min_row=2,values_only=True))
        for variant,filename in (("PHYSICAL_PERSON","Шаблон_тренировки_физлицо.docx"),("INDIVIDUAL_ENTREPRENEUR","Шаблон_тренировки_ИП.docx")):
            sha=EXPECTED[filename];template_version=session.query(TemplateVersion).filter_by(source_sha256=sha,status="DRAFT").first()
            if template_version:
                metadata=dict(template_version.preview_metadata or {});existing={x.get("candidate_id"):x for x in metadata.get("binding_candidates",[])}
                for candidate in (x for x in plan["candidates"] if x["variant"]==variant):
                    previous=existing.get(candidate["candidate_id"],{});preserved={key:value for key,value in previous.items() if key in {"status","review_decision","materialized_id","materialized_type","materialization_source","materialized_at","review_group","review_reason","required_specialist_action"}}
                    existing[candidate["candidate_id"]]={**candidate,**preserved}
                    if preserved.get("materialized_id"):existing[candidate["candidate_id"]]["status"]="CONFIGURED"
                metadata["binding_candidates"]=list(existing.values());metadata["r363_import_signature"]=plan["signature"];template_version.preview_metadata=metadata;candidate_counts[variant]=len(existing)
                schema=session.query(QuestionnaireSchema).filter_by(template_version_id=template_version.id).first()
                known={x.field_key for x in session.query(QuestionnaireField).filter_by(schema_id=schema.id)} if schema else set();added=0
                for position,row in enumerate(rf_rows,1):
                    rf_id,key,kind,scope,condition,required,allowed,render,question,operation,missing=row
                    if scope not in {"BOTH",variant}:continue
                    if key in known:continue
                    config={"group":"Поля договора","description":render or "","rf_id":rf_id,"source_operation":operation}
                    if kind=="SELECT":config["options"]=_options(allowed)
                    if kind=="ARRAY":config["item_config"]={"min_items":0,"max_items":50,"max_length":500}
                    if kind=="OBJECT":config["components"]=[{"key":"passport","label":"Паспортные данные","field_type":"TEXT","required":True},{"key":"snils","label":"СНИЛС","field_type":"TEXT","required":True}]
                    session.add(QuestionnaireField(schema_id=schema.id,field_key=key,label=question if question and question!='—' else render or key,field_type=RF_TYPE_MAP.get(kind,"TEXT"),required=required=="YES",position=position,config=config));added+=1
                template_field_counts[variant]=added
        session.commit()
        return {"product_id": product.id, "version_id": version.id, "created_product": created_product, "created_fields": created, "total_fields": session.query(ContractProductField).filter_by(product_version_id=version.id).count(), "candidate_counts":candidate_counts,"template_fields_created":template_field_counts,"variants_created": 0, "published": False}
    except Exception:
        session.rollback(); raise
    finally:
        if own_session: session.close()
