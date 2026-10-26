"""Isolated, non-official combined OOXML diagnostics for both training drafts."""
from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document

from app.conditions import ConditionEngine
from app.contract_platform import (
    calculated_document_values, condition_answers, formatted_bindings, render_docx,
    safe_source, validate_bindings, validate_regions, validate_repeating_groups,
)
from app.db import QuestionnaireField, QuestionnaireSchema, SessionLocal, TemplateVersion


VERSIONS = {
    "ip": "4eff6663-f03e-4858-bb2a-6dc39b4b1b95",
    "physical": "3dbdc61e-953b-49ce-a5c7-57df2613e454",
}
OUTPUT = Path("/app/storage/training-combined-diagnostics")
OUTPUT.mkdir(parents=True, exist_ok=True)
session = SessionLocal()
try:
    for variant, version_id in VERSIONS.items():
        version = session.get(TemplateVersion, version_id)
        schema = session.query(QuestionnaireSchema).filter_by(template_version_id=version.id).first()
        fields = session.query(QuestionnaireField).filter_by(schema_id=schema.id).all()
        bindings = validate_bindings(version, session)
        regions = validate_regions(version, session, bindings)
        groups = validate_repeating_groups(version, session, bindings, regions)
        original_sha = hashlib.sha256(safe_source(version).read_bytes()).hexdigest()
        for guarantee, vat_status, extra in ((False, "NO_VAT", False), (True, "VAT_PAYER", True)):
            answers = {
                "payment.model": "PREPAY", "payment.bank_guarantee_clause": guarantee,
                "schedule.place_column": True, "extra_option.enabled": extra,
                "schedule.entries": [
                    {"_id": "one", "month": "Октябрь", "date": "2026-10-26", "weekday": "Понедельник", "time": "10:00", "quantity": 1, "price": 5000, "place": "Зал 1"},
                    {"_id": "two", "month": "Ноябрь", "date": "2026-11-02", "weekday": "Понедельник", "time": "11:00", "quantity": 2, "price": 6000, "place": "Зал 2"},
                ],
            }
            if variant == "ip":
                answers.update({"tax.vat_status": vat_status, "ip.ogrnip": "123456789012345"})
            included = {region.id for region in regions if ConditionEngine.evaluate(region.include_when, condition_answers(fields, answers))}
            calculated_groups, calculated, repeated = calculated_document_values(version.id, session, fields, answers)
            output = OUTPUT / f"{variant}-guarantee-{int(guarantee)}.docx"
            render_docx(version, bindings, formatted_bindings(bindings, calculated, fields), output,
                regions, included, calculated_groups, repeated, {group.id: group.fields for group in calculated_groups})
            with zipfile.ZipFile(output) as archive:
                assert archive.testzip() is None
            doc = Document(output)
            full_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            assert ("предоставления Исполнителем банковской гарантии" in full_text) == guarantee
            assert len(doc.tables[1].rows) == 6 + int(extra) and len(doc.tables[6].rows) == 6 + int(extra)
            assert doc.tables[1].rows[2].cells[6].text == "Зал 1"
            assert doc.tables[6].rows[4].cells[6].text == "Зал 2"
            if variant == "ip":
                assert ("В случае нарушения Исполнителем требований ст. 168 НК РФ" in full_text) == (vat_status == "VAT_PAYER")
                assert ("Счета-фактуры выставляются Исполнителем" in full_text) == (vat_status == "VAT_PAYER")
                assert "ОГРНИП 123456789012345" in " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            assert hashlib.sha256(safe_source(version).read_bytes()).hexdigest() == original_sha
            print(json.dumps({"variant": variant, "guarantee": guarantee, "vat_status": vat_status if variant == "ip" else None,
                "extra_option": extra, "path": str(output), "size": output.stat().st_size,
                "sha256": hashlib.sha256(output.read_bytes()).hexdigest(), "tables": len(doc.tables)}, ensure_ascii=False))
finally:
    session.close()
