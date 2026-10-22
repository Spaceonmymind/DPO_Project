"""Read-only R3.6 inventory audit for the immutable training-contract fixtures."""
from __future__ import annotations

import hashlib
from collections import Counter
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

EXPECTED = {
    "Матрица_применимости_LLM_v2.xlsx": "5ee3ae75b9aa82f54ecf5c1daa0aad215e69fbc9fb8ca3e6d6aa18e291e7d97e",
    "Шаблон_тренировки_ИП.docx": "9560a30d32c08b39f5c0dde1671829302a6331ec96ab0146b03489503ef7d8a5",
    "Шаблон_тренировки_физлицо.docx": "1e39f28d118567ef22e411ab60d5d836fea5c5d71bbaf7917298651879ac94f8",
}


def audit_training_sources(root: Path) -> dict:
    files = {}
    for name, expected_sha in EXPECTED.items():
        path = root / name
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        files[name] = {"sha256": actual, "immutable_sha_matches": actual == expected_sha}
    workbook = load_workbook(root / "Матрица_применимости_LLM_v2.xlsx", data_only=True)
    facts = list(workbook["FACT_SCHEMA"].iter_rows(min_row=2, values_only=True))
    fields = list(workbook["RED_FIELD_SCHEMA"].iter_rows(min_row=2, values_only=True))
    occurrences = list(workbook["RED_OCCURRENCES"].iter_rows(min_row=2, values_only=True))
    tests = list(workbook["PROMPT_TESTS"].iter_rows(min_row=2, values_only=True))
    field_ids = [row[0] for row in fields]
    occurrence_field_ids = [row[3] for row in occurrences]
    docs = {}
    for name in EXPECTED:
        if not name.endswith(".docx"):
            continue
        document = Document(root / name)
        docs[name] = {"tables": len(document.tables), "paragraphs": len(document.paragraphs)}
    return {
        "files": files,
        "inventory": {
            "facts": len(facts), "red_fields": len(fields), "red_occurrences": len(occurrences),
            "decision_rules": workbook["DECISION_RULES"].max_row - 1,
            "prompt_tests": len(tests), "docx": docs,
        },
        "integrity": {
            "unique_field_ids": len(set(field_ids)) == 25,
            "all_occurrences_reference_fields": set(occurrence_field_ids) <= set(field_ids),
            "occurrences_per_field": dict(sorted(Counter(occurrence_field_ids).items())),
            "all_files_match": all(item["immutable_sha_matches"] for item in files.values()),
        },
    }
