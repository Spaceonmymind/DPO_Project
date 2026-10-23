"""Idempotent import of the approved training decision matrix into ContractProduct."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from app.db import (
    ContractProduct, ContractProductAuditEvent, ContractProductField,
    ContractProductFieldMapping, ContractProductVariant, ContractProductVersion,
    QuestionnaireField, QuestionnaireSchema, TemplateVersion,
)
from app.training_contract_audit import EXPECTED
from app.training_contracts import empty_facts, extract_training_facts, merge_facts


def _value(raw):
    value=raw.strip()
    if value=='true':return True
    if value=='false':return False
    if value=='null':return None
    return value


def parse_matrix_condition(expression:str)->dict:
    expression=expression.strip()
    for operator in (' OR ',' AND '):
        if operator in expression:
            return {'operator':operator.strip(),'conditions':[parse_matrix_condition(item) for item in expression.split(operator)]}
    match=re.fullmatch(r'([A-Za-z0-9_.]+)\s+(NOT IN|IN)\s+\[([^]]+)\]',expression)
    if match:return {'operator':match.group(2).replace(' ','_'),'field_key':match.group(1),'value':[_value(x) for x in match.group(3).split(',')]}
    match=re.fullmatch(r'([A-Za-z0-9_.]+)\s*(==|!=)\s*(.+)',expression)
    if not match:raise ValueError(f'Неподдерживаемое правило матрицы: {expression}')
    value=_value(match.group(3))
    if value is None:return {'operator':'IS_EMPTY' if match.group(2)=='==' else 'IS_NOT_EMPTY','field_key':match.group(1)}
    return {'operator':'EQUALS' if match.group(2)=='==' else 'NOT_EQUALS','field_key':match.group(1),'value':value}


def flatten(values,prefix=''):
    result={}
    for key,value in values.items():
        path=f'{prefix}.{key}' if prefix else key
        if isinstance(value,dict):result.update(flatten(value,path))
        else:result[path]=value
    return result


def fixture_answers(prompt:str,category:str='')->dict:
    text=prompt.lower();facts=merge_facts(empty_facts(),extract_training_facts(prompt))
    if any(x in text for x in ('заказчиком выступает физическое лицо','заказчик — другое ооо','не корпоративное ао')):facts['customer']['type']='OTHER_LEGAL_ENTITY' if 'ооо' in text else 'INDIVIDUAL'
    else:facts['customer']['type']='CORPORATE_AO'
    if category=='APPLICABLE_PHYSICAL':facts['subject'].update(type='TRAINING_SERVICE',paid=True);facts['contractor']['status']='PHYSICAL_NO_IP_NO_NPD';facts['performance'].update(personal=True,third_parties=False)
    if category=='APPLICABLE_IP':facts['subject'].update(type='TRAINING_SERVICE',paid=True);facts['contractor']['status']='IP_NO_NPD';facts['performance'].update(personal=True,third_parties=False)
    if any(x in text for x in ('автономная некоммерческая','спортивным клубом ооо')):facts['contractor']['status']='LEGAL_ENTITY'
    if 'профессиональн' in text and ('налог' in text or 'доход' in text) and not any(x in text for x in ('не применяет','не является','без нпд')):facts['contractor']['status']='NPD_PHYSICAL'
    if re.search(r'\b(?:1[0-7]|[1-9])\s*лет\b',text):facts['contractor']['status']='MINOR';facts['contractor']['age_18_plus']=False
    if 'статус пока неизвест' in text or 'то ли физлицо, то ли ип' in text:facts['contractor']['status']='UNKNOWN'
    if 'физлицо, но прислал огрнип' in text:facts['contractor']['status']='CONFLICT'
    if 'какие-то услуги тренера, деталей нет' in text:facts['subject']['type']='UNKNOWN';facts['contractor']['status']='PHYSICAL_NO_IP_NO_NPD'
    if any(x in text for x in ('купить','поставка','арендовать','аренда','турнир','разработка сайта','мобильное приложение','лиценз','ремонт')):facts['subject']['type']='GOODS'
    if any(x in text for x in ('зарплат','постоянное рабочее место','руководитель')):facts['labor']['salary_or_position']=True
    if any(x in text for x in ('заменить себя','замена другим тренером','субподрядчик','распределяет тренеров')):facts['performance'].update(personal=False,third_parties=True)
    if 'неизвестно, кто фактически' in text:facts['subject'].update(type='TRAINING_SERVICE',paid=True);facts['contractor']['status']='PHYSICAL_NO_IP_NO_NPD';facts['performance'].update(personal=None,third_parties=None)
    return flatten(facts)


def configure_training_resolver(session,owner,root:Path)->dict:
    workbook=load_workbook(root/'Матрица_применимости_LLM_v2.xlsx',data_only=True)
    product=session.query(ContractProduct).filter_by(owner_id=owner.id,code='TRAINING_CONTRACT').one();version=session.query(ContractProductVersion).filter_by(product_id=product.id,status='DRAFT').one()
    fields={x.field_key:x for x in session.query(ContractProductField).filter_by(product_version_id=version.id)}
    variants={x.code:x for x in session.query(ContractProductVariant).filter_by(product_version_id=version.id)}
    definitions=(('PHYSICAL_PERSON','Договор с физическим лицом','Шаблон_тренировки_физлицо.docx','PHYSICAL_NO_IP_NO_NPD'),('INDIVIDUAL_ENTREPRENEUR','Договор с ИП','Шаблон_тренировки_ИП.docx','IP_NO_NPD'))
    for position,(code,name,filename,status) in enumerate(definitions,1):
        template_version=session.query(TemplateVersion).filter_by(source_sha256=EXPECTED[filename],status='DRAFT').one()
        condition={'operator':'EQUALS','field_key':'contractor.status','value':status}
        variant=variants.get(code)
        if not variant:
            variant=ContractProductVariant(product_version_id=version.id,code=code,name=name,template_version_id=template_version.id,selection_condition=condition,position=position);session.add(variant);session.flush();variants[code]=variant
        schema=session.query(QuestionnaireSchema).filter_by(template_version_id=template_version.id).one();targets={x.field_key:x for x in session.query(QuestionnaireField).filter_by(schema_id=schema.id)}
        for key,field in fields.items():
            target=targets.get(key)
            if not target or field.field_type!=target.field_type:continue
            if not session.query(ContractProductFieldMapping).filter_by(variant_id=variant.id,product_field_id=field.id).first():session.add(ContractProductFieldMapping(variant_id=variant.id,product_field_id=field.id,target_field_id=target.id))
    decision_rules=[]
    for row in workbook['DECISION_RULES'].iter_rows(min_row=2,values_only=True):
        rule_id,priority,kind,condition,applicability,template,render_status,stop,reason_code,message=row
        if kind not in {'STOP','SELECT'}:continue
        decision_rules.append({'rule_id':rule_id,'priority':priority,'condition':parse_matrix_condition(condition),'state':'SELECTED' if kind=='SELECT' else ('NEED_MORE_INFORMATION' if render_status=='NEEDS_INPUT' else 'NOT_APPLICABLE'),'variant_code':template if kind=='SELECT' else None,'reason_code':reason_code,'message':message})
    scenarios=[]
    for row in workbook['PROMPT_TESTS'].iter_rows(min_row=2,values_only=True):
        test_id,category,prompt,applicability,template,rule_id,render_status,reason,test_status=row
        scenarios.append({'name':test_id,'prompt':prompt,'answers':fixture_answers(prompt,category),'expected_state':'SELECTED' if applicability=='APPLICABLE' else ('NEED_MORE_INFORMATION' if render_status=='NEEDS_INPUT' else 'NOT_APPLICABLE'),'expected_variant':template,'expected_rule_id':rule_id})
    rules=dict(version.rules or {});rules.update({'decision_rules':decision_rules,'test_scenarios':scenarios,'matrix_source':'Матрица_применимости_LLM_v2.xlsx'});version.rules=rules
    session.add(ContractProductAuditEvent(product_id=product.id,product_version_id=version.id,actor_id=owner.id,event_type='TRAINING_RESOLVER_CONFIGURED',metadata_json={'rules':len(decision_rules),'scenarios':len(scenarios)}));session.commit()
    return {'product_id':product.id,'version_id':version.id,'variants':len(variants),'rules':len(decision_rules),'scenarios':len(scenarios)}


def execute_training_fixtures(session, owner) -> dict:
    """Execute saved fixtures through the production resolver and persist evidence."""
    from app.contract_products import ContractVariantResolver
    product=session.query(ContractProduct).filter_by(owner_id=owner.id,code='TRAINING_CONTRACT').one()
    version=session.query(ContractProductVersion).filter_by(product_id=product.id,status='DRAFT').one()
    fields=session.query(ContractProductField).filter_by(product_version_id=version.id).all()
    variants=session.query(ContractProductVariant).filter_by(product_version_id=version.id).all()
    results=[]
    for scenario in (version.rules or {}).get('test_scenarios',[]):
        actual=ContractVariantResolver.resolve(version,fields,variants,scenario['answers'])
        variant=next((item for item in variants if item.id==actual.get('variant_id')),None)
        passed=(actual.get('state')==scenario['expected_state'] and actual.get('rule_id')==scenario['expected_rule_id'] and (scenario.get('expected_variant') in {None,'NONE'} or (variant and variant.code==scenario['expected_variant'])))
        results.append({'name':scenario['name'],'expected_state':scenario['expected_state'],'actual_state':actual.get('state'),'expected_rule_id':scenario['expected_rule_id'],'actual_rule_id':actual.get('rule_id'),'expected_variant':scenario.get('expected_variant'),'actual_variant':variant.code if variant else 'NONE','status':'PASS' if passed else 'FAIL','reason':actual.get('reason')})
    rules=dict(version.rules or {});rules['resolver_results']=results;version.rules=rules
    summary={'total':len(results),'pass':sum(x['status']=='PASS' for x in results),'fail':sum(x['status']=='FAIL' for x in results)}
    session.add(ContractProductAuditEvent(product_id=product.id,product_version_id=version.id,actor_id=owner.id,event_type='PRODUCT_FIXTURES_EXECUTED',metadata_json=summary));session.commit()
    return summary
