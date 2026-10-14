import hashlib
from copy import deepcopy
from pathlib import Path

import openpyxl
from docx import Document

from app.documents.service import generate_xlsx
from app.training_contracts import decide_training, empty_facts, extract_training_facts, load_occurrences, render_training_contract


_repository_templates=Path('/app/contract_templates') if Path('/app/contract_templates/training').is_dir() else Path(__file__).resolve().parents[2]/'contract_templates'
CONTRACT_ROOT=_repository_templates
TRAINING=CONTRACT_ROOT/'training'


def complete_facts(status='PHYSICAL_NO_IP_NO_NPD'):
    facts=empty_facts()
    facts.update(subject={'type':'TRAINING_SERVICE','paid':True},contractor={'status':status,'citizenship':'RU','age_18_plus':True},performance={'personal':True,'third_parties':False})
    facts['contract']['form_label']='TYPICAL_FORM'; facts['service']['name']='Проведение тренировок'
    facts['price']['total']=100000; facts['term'].update(start='2026-11-01',end='2027-01-31',more_than_one_year=False)
    facts['payment'].update(model='POSTPAY',postpay_percent=100,postpay_days=10,bank_guarantee=False)
    facts['schedule'].update(months=['2026-11'],entries=[{'date':'2026-11-10','weekday':'вторник','time':'10:00','count':1,'unit_price':100000,'place':'Москва'}])
    facts['places']['multiple_or_not_in_contract']=True; facts['extra_option']['enabled']=False
    if status=='PHYSICAL_NO_IP_NO_NPD':
        facts['tax'].update(ndfl_rate=13,ndfl_amount=13000); facts['physical'].update(passport='45 01 123456',snils='123-456-789 00')
    else:
        facts['tax'].update(vat_status='NO_VAT',no_vat_basis='УСН'); facts['ip']['ogrnip']='123456789012345'
    return facts


def test_decision_priority_and_readiness():
    assert extract_training_facts('Обычное физлицо без ИП')['contractor']['status']=='PHYSICAL_NO_IP_NO_NPD'
    scenarios=[
        ('LEGAL_ENTITY','D004'),('NPD_PHYSICAL','D005'),('IP_NPD','D005'),('FOREIGN_PHYSICAL','D006'),('MINOR','D007'),
    ]
    for status,rule in scenarios:
        facts=complete_facts(); facts['contractor']['status']=status
        assert decide_training(facts).matched_rule_id==rule
    facts=complete_facts(); facts['performance']['third_parties']=True; assert decide_training(facts).matched_rule_id=='D008'
    facts=complete_facts(); facts['labor']['fixed_workday']=True; assert decide_training(facts).matched_rule_id=='D009'
    for status,template,rule in [('PHYSICAL_NO_IP_NO_NPD','PHYSICAL_PERSON','D030'),('IP_NO_NPD','INDIVIDUAL_ENTREPRENEUR','D031')]:
        decision=decide_training(complete_facts(status))
        assert (decision.template,decision.matched_rule_id,decision.render_status)==(template,rule,'READY')
        assert len(decision.red_fields)==25 and not decision.missing_field_ids
    incomplete=complete_facts(); incomplete['schedule']['entries']=None
    assert decide_training(incomplete).render_status=='NEEDS_INPUT' and 'RF021' in decide_training(incomplete).missing_field_ids


def test_matrix_and_both_renderers_preserve_sources(tmp_path):
    assert len(load_occurrences(CONTRACT_ROOT))==74
    for status in ('PHYSICAL_NO_IP_NO_NPD','IP_NO_NPD'):
        decision=decide_training(complete_facts(status)); source_name='Шаблон_тренировки_физлицо.docx' if status.startswith('PHYSICAL') else 'Шаблон_тренировки_ИП.docx'
        source=TRAINING/source_name; digest=hashlib.sha256(source.read_bytes()).hexdigest(); output=tmp_path/f'{status}.docx'
        result=render_training_contract(CONTRACT_ROOT,complete_facts(status),decision,output)
        assert result['occurrences']==74 and not result['mapping_errors'] and output.stat().st_size>1000
        assert hashlib.sha256(source.read_bytes()).hexdigest()==digest
        document=Document(output); text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
        assert 'Проведение тренировок' in text and 'красным' not in text.lower()


def test_xlsx_export_contains_summary_items_and_requirements(tmp_path):
    path=tmp_path/'spec.xlsx'
    data={'subject':'Поставка ноутбуков','purpose':'Оснащение','scope':'2 штуки','deadline':'2026-12-01','acceptance_criteria':'По количеству и характеристикам','items':[{'name':'MacBook Pro','characteristics':'M4, 24 ГБ','quantity':2,'unit':'шт.','warranty':'12 месяцев','delivery_address':'Екатеринбург'}],'blocking_requirements':['Только новое оборудование']}
    generate_xlsx(data,path); sheet=openpyxl.load_workbook(path).active
    assert sheet['B3'].value=='Поставка ноутбуков' and sheet['B6'].value=='2026-12-01' and sheet['B7'].value=='По количеству и характеристикам'
    assert sheet['B11'].value=='MacBook Pro' and sheet['C11'].value=='M4, 24 ГБ' and sheet['D11'].value==2
    assert sheet['G11'].value=='12 месяцев' and sheet['I11'].value=='Екатеринбург' and 'Только новое' in sheet['B12'].value
