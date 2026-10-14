import hashlib
from copy import deepcopy
from pathlib import Path

from docx import Document

from app.supply_specifications import (
    SupplySpecificationDialogueEngine,
    SupplySpecificationRenderer,
    empty_supply_specification,
    extract_items,
    next_question,
    parse_money,
    parse_positive_integer,
    recalculate,
)


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_ROOT = ROOT / 'specification_templates'
SOURCE = TEMPLATE_ROOT / 'supply' / 'supply_specification.docx'


def complete_data(item_count=1, vat_mode='VAT_PAYER', support=True):
    items=[]
    for index in range(item_count):
        items.append({'line_number':index+1,'part_number':f'PN-{index+1}','description':f'Ноутбук модель {index+1}',
                      'manufacturer':'Lenovo','warranty_term':'3 года','unit_price':str(100000+index*1000),
                      'quantity':index+1,'total_price':'0'})
    data=empty_supply_specification()
    data.update(equipment_category='компьютерное оборудование',delivery_address='Москва, ул. Большая Татарская, 11',
                currency='рублей',vat_mode=vat_mode,vat_rate=20 if vat_mode=='VAT_PAYER' else None,
                vat_basis='' if vat_mode=='VAT_PAYER' else 'НДС не облагается на основании применения УСН',
                delivery_term='в течение 10 рабочих дней',support_required=support,
                support_services='консультации и устранение неисправностей' if support else '',
                support_term='12 месяцев' if support else '',support_start_condition='с даты поставки' if support else '',
                warranty_service_description='ремонт и замена неисправных компонентов',
                warranty_terms_url='https://example.test/warranty',items=items,confirmed=True)
    return recalculate(data)


def all_text(document):
    return '\n'.join([*(p.text for p in document.paragraphs),*(cell.text for table in document.tables for row in table.rows for cell in row.cells)])


def test_template_structure_matches_approved_supply_form():
    document=Document(SOURCE)
    assert len(document.tables[0].columns)==8
    assert len(document.tables[0].rows)==6
    assert document.tables[0].rows[-2].cells[0].text.startswith('ИТОГО')
    assert document.tables[0].rows[-1].cells[0].text.startswith('В том числе НДС')
    text='\n'.join(p.text for p in document.paragraphs)
    assert 'Таблица № 1' in text and 'Срок поставки Товара' in text and 'Гарантийное обслуживание' in text


def test_natural_item_and_numeric_normalization():
    items=extract_items('Нужно 10 ноутбуков Lenovo ThinkPad, партномер 21ABC, гарантия 3 года, по 100 тыс. руб.')
    assert len(items)==1
    assert items[0]['quantity']==10 and items[0]['part_number']=='21ABC'
    assert items[0]['manufacturer'].lower()=='lenovo' and items[0]['unit_price']=='100000.00'
    assert parse_positive_integer('десять шт.')==10 and parse_money('100 000 руб.')==100000


def test_renderer_single_item_vat_totals_and_source_immutable(tmp_path):
    digest=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    data=complete_data(); output=tmp_path/'single.docx'
    SupplySpecificationRenderer(TEMPLATE_ROOT).render(data,output)
    document=Document(output); table=document.tables[0]; text=all_text(document)
    assert table.rows[1].cells[1].text=='PN-1' and table.rows[1].cells[7].text.startswith('100 000')
    assert 'ИТОГО' in table.rows[-2].cells[0].text and 'НДС 20%' in table.rows[-1].cells[0].text
    assert '100 000' in text and 'сто тысяч' in text
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==digest


def test_renderer_adds_fourth_item_preserves_totals_and_removes_no_vat_support(tmp_path):
    data=complete_data(4,vat_mode='NO_VAT',support=False); output=tmp_path/'four.docx'
    SupplySpecificationRenderer(TEMPLATE_ROOT).render(data,output)
    document=Document(output); table=document.tables[0]; text=all_text(document)
    assert len(table.rows)==6  # header + four items + total; VAT row removed
    assert [table.rows[index].cells[1].text for index in range(1,5)]==['PN-1','PN-2','PN-3','PN-4']
    assert table.rows[-1].cells[0].text.startswith('ИТОГО')
    assert 'В том числе НДС' not in text and 'техническую поддержку (сервисный' not in text
    assert 'НДС не облагается на основании применения УСН' in text
    assert not any(token in text for token in ('[указать ', '[добавить ', '[выбрать '))


def test_total_calculation_two_items():
    data=complete_data(2)
    data['items'][0].update(unit_price='100000',quantity=10)
    data['items'][1].update(unit_price='20000',quantity=5)
    recalculate(data)
    assert data['items'][0]['total_price']=='1000000.00'
    assert data['items'][1]['total_price']=='100000.00'
    assert data['total_amount']=='1100000.00' and data['vat_amount']=='183333.33'


def test_schema_dialogue_never_repeats_valid_question_and_confirms():
    engine=SupplySpecificationDialogueEngine(); data=empty_supply_specification(); pending=engine.pending(next_question(data)); seen=[]
    answers={
      'equipment_category':'компьютерное оборудование','item_count':'1','items[0].description':'Ноутбук Lenovo ThinkPad',
      'items[0].part_number':'21ABC','items[0].manufacturer':'Lenovo','items[0].warranty_term':'3 года',
      'items[0].unit_price':'100 тыс. руб.','items[0].quantity':'10 шт.','delivery_address':'Москва, ул. Большая Татарская, 11',
      'currency':'рубли','vat_mode':'С НДС','vat_rate':'20%','delivery_term':'10 рабочих дней',
      'support_required':'Нет','warranty_service_description':'ремонт и замена','warranty_terms_url':'https://example.test/warranty',
    }
    while pending:
        seen.append(pending['question_id']); data,pending,_,ready=engine.process(data,answers[pending['target']],pending); assert not ready
    assert len(seen)==len(set(seen))
    data,pending,message,ready=engine.process(data,'Подтверждаю',None)
    assert ready and data['confirmed'] and pending is None and 'готово' in message.lower()


def test_editing_via_chat_is_sparse():
    engine=SupplySpecificationDialogueEngine(); data=complete_data(2); original=deepcopy(data['items'][1])
    changed,_,_,_=engine.process(data,'Поменяй количество ноутбуков на 12.',None)
    assert changed['items'][0]['quantity']==12 and changed['items'][1]['part_number']==original['part_number']
    changed,_,_,_=engine.process(changed,'Убери вторую позицию.',None)
    assert len(changed['items'])==1
