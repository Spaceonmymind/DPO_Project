"""Structural candidate discovery for the immutable training-contract matrix."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from app.contract_platform import adapter

DOCS={
 "PHYSICAL_PERSON":"Шаблон_тренировки_физлицо.docx",
 "INDIVIDUAL_ENTREPRENEUR":"Шаблон_тренировки_ИП.docx",
}

def matrix_locator(value:str)->str|None:
 if match:=re.fullmatch(r'P(\d+)',value):return f'word/document.xml#/p[{int(match.group(1))+1}]'
 if match:=re.fullmatch(r'T(\d+)R(\d+)C(\d+)P(\d+)',value):
  t,r,c,p=map(int,match.groups());return f'word/document.xml#/tbl[{t+1}]/tr[{r+1}]/tc[{c+1}]/p[{p+1}]'
 return None

def discover_candidates(root:Path)->list[dict]:
 workbook=load_workbook(root/'Матрица_применимости_LLM_v2.xlsx',data_only=True)
 fields={row[0]:{'field_path':row[1],'operation':row[9]} for row in workbook['RED_FIELD_SCHEMA'].iter_rows(min_row=2,values_only=True)}
 sheet=workbook['RED_OCCURRENCES'];headers=[x.value for x in sheet[1]]
 previews={key:adapter.preview(root/name) for key,name in DOCS.items()};result=[]
 for row in sheet.iter_rows(min_row=2,values_only=True):
  item=dict(zip(headers,row));targets=list(DOCS) if item['template_scope'] in {'BOTH','NONE'} else [item['template_scope']]
  for variant in targets:
   locator=matrix_locator(str(item['docx_location']));nodes={x['locator']:x for x in previews[variant]['nodes']};node=nodes.get(locator);red=str(item['red_text'] or '')
   if not node:status='NOT_FOUND';start=end=None
   elif red and red in node['text']:status='CANDIDATE_FOUND';start=node['text'].index(red);end=start+len(red)
   else:status='REVIEW_REQUIRED';start=0;end=len(node['text'])
   result.append({'candidate_id':f"{item['occurrence_id']}:{variant}",'occurrence_id':item['occurrence_id'],'rf_id':item['canonical_field_id'],'field_path':fields[item['canonical_field_id']]['field_path'],'variant':variant,'source_filename':DOCS[variant],'matrix_location':item['docx_location'],'locator':locator,'part':node['part'] if node else None,'structure_fingerprint':previews[variant]['structure_fingerprint'],'start':start,'end':end,'selected_text':node['text'][start:end] if node and start is not None else None,'context':node['text'][:500] if node else None,'operation':fields[item['canonical_field_id']]['operation'],'status':status})
 return result
