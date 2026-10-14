"""Deterministic training-contract workflow based on matrix specification v2.0."""
from __future__ import annotations

import json
import logging
import re
from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

log = logging.getLogger('dpo.training')

TEMPLATES = {
    'PHYSICAL_PERSON': ('physical_person.docx', 'Шаблон_тренировки_физлицо.docx'),
    'INDIVIDUAL_ENTREPRENEUR': ('individual_entrepreneur.docx', 'Шаблон_тренировки_ИП.docx'),
}
MATRIX_NAMES = ('applicability_matrix.xlsx', 'Матрица_применимости_LLM_v2.xlsx')


@dataclass(frozen=True)
class TrainingDecision:
    applicability: str
    template: str
    matched_rule_id: str
    reason_code: str
    reason: str
    render_status: str
    missing_field_ids: list[str]
    red_fields: list[dict[str, Any]]


def _get(data: dict, path: str, default=None):
    node: Any = data
    for part in path.split('.'):
        if not isinstance(node, dict): return default
        node = node.get(part, default)
    return node


def _set(data: dict, path: str, value: Any):
    parts=path.split('.'); node=data
    for part in parts[:-1]: node=node.setdefault(part,{})
    node[parts[-1]]=value


def empty_facts() -> dict:
    return {
        'subject': {'type':'UNKNOWN','paid':None},
        'customer': {'type':'CORPORATE_AO'},
        'contractor': {'status':'UNKNOWN','citizenship':'UNKNOWN','age_18_plus':None},
        'performance': {'personal':None,'third_parties':None},
        'labor': {'salary_or_position':None,'internal_rules':None,'fixed_workday':None,'leave_benefits':None},
        'service': {'name':None}, 'contract': {'form_label':None}, 'appendices': {'other':[]},
        'price': {'total':None}, 'tax': {'ndfl_rate':None,'ndfl_amount':None,'vat_status':None,'vat_rate':None,'vat_amount':None,'no_vat_basis':None},
        'payment': {'model':None,'advance_percent':None,'advance_days':None,'postpay_percent':None,'postpay_days':None,'bank_guarantee':None},
        'term': {'start':None,'end':None,'more_than_one_year':None},
        'schedule': {'months':None,'entries':None}, 'places': {'multiple_or_not_in_contract':None},
        'extra_option': {'enabled':None,'limit':None}, 'physical': {'passport':None,'snils':None},
        'ip': {'ogrnip':None,'long_term_resources_clause':None},
    }


def merge_facts(current: dict | None, patch: dict) -> dict:
    result=deepcopy(current or empty_facts())
    def merge(target, source):
        for key,value in source.items():
            if isinstance(value,dict) and isinstance(target.get(key),dict): merge(target[key],value)
            elif value is not None: target[key]=value
    merge(result,patch); return result


def combine_patches(*patches: dict) -> dict:
    """Combine sparse patches without materialising UNKNOWN defaults."""
    result: dict={}
    def merge(target,source):
        for key,value in source.items():
            if isinstance(value,dict): merge(target.setdefault(key,{}),value)
            elif value is not None: target[key]=value
    for patch in patches: merge(result,patch)
    return result


def extract_training_facts(text: str, current: dict | None=None) -> dict:
    """Mock-compatible fact extraction only; it never decides applicability."""
    t=' '.join(text.lower().replace('ё','е').split()); patch: dict[str,Any]={}
    def put(path,value): _set(patch,path,value)
    if any(x in t for x in ('трениров','тренер','заняти','спортивн')): put('subject.type','TRAINING_SERVICE')
    elif any(x in t for x in ('товар','купить','поставка')): put('subject.type','GOODS')
    elif any(x in t for x in ('аренд','лизинг')): put('subject.type','LEASE')
    elif any(x in t for x in ('разработка','сайт','программ')): put('subject.type','IT_SERVICE')
    if any(x in t for x in ('бесплат','безвозмезд')): put('subject.paid',False)
    elif any(x in t for x in ('платн','оплат','стоимост','рубл','аванс','постоплат')) or _get(patch,'subject.type')=='TRAINING_SERVICE': put('subject.paid',True)
    statuses=[]
    explicitly_physical=any(x in t for x in ('обычное физлицо','обычный физ','физлицо без ип','физическое лицо без ип','не ип и не самозанят'))
    if explicitly_physical: statuses.append('PHYSICAL_NO_IP_NO_NPD')
    elif any(x in t for x in ('ип на нпд','ип применяет нпд','ип-самозанят')): statuses.append('IP_NPD')
    elif any(x in t for x in ('самозанят','плательщик нпд','налог на профессиональный доход')): statuses.append('NPD_PHYSICAL')
    if any(x in t for x in ('юридическое лицо','юрлицо','ооо ')): statuses.append('LEGAL_ENTITY')
    if any(x in t for x in ('иностран','не гражданин рф')): statuses.append('FOREIGN_PHYSICAL'); put('contractor.citizenship','NON_RU')
    if any(x in t for x in ('несовершеннолет','младше 18')): statuses.append('MINOR'); put('contractor.age_18_plus',False)
    if not statuses and re.search(r'\bип\b|индивидуальн\w* предпринимател',t): statuses.append('IP_NO_NPD')
    if len(set(statuses))>1: put('contractor.status','CONFLICT')
    elif statuses: put('contractor.status',statuses[0])
    if any(x in t for x in ('гражданин рф','гражданка рф','российское гражданство')): put('contractor.citizenship','RU')
    if any(x in t for x in ('старше 18','совершеннолет')): put('contractor.age_18_plus',True)
    if any(x in t for x in ('лично','сам проводит','сама проводит','самостоятельно проводит')): put('performance.personal',True)
    if any(x in t for x in ('не лично','сотрудник','субподряд','другие тренер','замещающ','помощник')): put('performance.personal',False); put('performance.third_parties',True)
    if any(x in t for x in ('без третьих лиц','без привлечения других лиц','других тренеров не','замены не','помощники не','сотрудники не привлека','никого не привлека')): put('performance.third_parties',False)
    if any(x in t for x in ('будут еще другие тренеры','будут ещё другие тренеры','привлекать других исполнителей','привлекать третьих лиц')): put('performance.third_parties',True)
    labor={'salary_or_position':('оклад','должност','в штат','штатная'), 'internal_rules':('внутренн','трудов распоряд'), 'fixed_workday':('рабочий день','с 9 до 18'), 'leave_benefits':('отпуск','больничн')}
    for field,words in labor.items():
        if any(word in t for word in words): put('labor.'+field,True)
        elif any(x in t for x in ('без признаков трудовых','гпх, трудовых признаков нет')): put('labor.'+field,False)
    patterns={
      'price.total':r'(?:стоимость|цена)\s*(?:(?:договора|услуг)\s*)?[:—-]?\s*([\d ]+(?:[,.]\d+)?)',
      'tax.ndfl_rate':r'ндфл\s*[:—-]?\s*(\d+(?:[,.]\d+)?)\s*%', 'tax.ndfl_amount':r'(?:сумма ндфл|ндфл составляет)\s*[:—-]?\s*([\d ]+(?:[,.]\d+)?)',
      'tax.vat_rate':r'ндс\s*[:—-]?\s*(\d+(?:[,.]\d+)?)\s*%', 'tax.vat_amount':r'(?:сумма ндс|ндс составляет)\s*[:—-]?\s*([\d ]+(?:[,.]\d+)?)',
      'payment.advance_percent':r'аванс\s*[:—-]?\s*(\d+(?:[,.]\d+)?)\s*%', 'payment.advance_days':r'аванс[^.]{0,50}?(\d+)\s*(?:календарн\w*\s*)?дн',
      'payment.postpay_percent':r'(?:постоплат|остаток)[^\d]{0,15}(\d+(?:[,.]\d+)?)\s*%', 'payment.postpay_days':r'(?:постоплат|окончательн\w* оплат|оплат\w* после оказания услуг)[^.]{0,70}?(\d+)\s*(?:календарн\w*\s*)?дн',
      'extra_option.limit':r'(?:не более|лимит)\s*(\d+)', 'ip.ogrnip':r'огрнип\D*(\d{15})',
    }
    for path,pattern in patterns.items():
        match=re.search(pattern,t)
        if match:
            raw=match.group(1).replace(' ','').replace(',','.')
            value=float(raw) if path not in {'payment.advance_days','payment.postpay_days','extra_option.limit','ip.ogrnip'} else (match.group(1) if path=='ip.ogrnip' else int(match.group(1)))
            if path=='price.total' and re.search(r'\bтыс(?:яч[аи]?)?\b',t): value=float(value)*1000
            put(path,value)
    advance_before=re.search(r'(\d+(?:[,.]\d+)?)\s*%\s*аванс',t)
    advance_days=re.search(r'аванс\w*[^.]{0,40}?(?:за|в течение)\s*(\d+)\s*(?:календарн\w*\s*)?дн',t)
    postpay_days=re.search(r'(?:остаток|оплат\w* после оказания услуг)[^.]{0,50}?(?:за|в течение)\s*(\d+)\s*(?:календарн\w*\s*)?дн',t)
    if advance_before: put('payment.advance_percent',float(advance_before.group(1).replace(',','.')))
    if advance_days: put('payment.advance_days',int(advance_days.group(1)))
    if postpay_days: put('payment.postpay_days',int(postpay_days.group(1)))
    for path,label in (('physical.passport','паспорт'),('physical.snils','снилс'),('tax.no_vat_basis','без ндс на основании'),('service.name','услуги:')):
        match=re.search(re.escape(label)+r'\s*[:—-]?\s*([^;,\n]+)',text,re.I)
        if match: put(path,match.group(1).strip(' .'))
    if any(x in t for x in ('типовая форма','типовой договор','стандартная форма','обычная типовая')): put('contract.form_label','TYPICAL_FORM')
    elif 'иная форма' in t or 'собственная форма' in t: put('contract.form_label','CUSTOM_FORM')
    if 'без ндс' in t or 'ндс не облагается' in t: put('tax.vat_status','NO_VAT')
    elif 'плательщик ндс' in t or re.search(r'ндс\s*\d+\s*%',t): put('tax.vat_status','VAT_PAYER')
    if 'смешанн' in t or ('аванс' in t and ('остаток' in t or 'постоплат' in t)): put('payment.model','MIXED')
    elif 'аванс' in t or 'предоплат' in t: put('payment.model','PREPAY')
    elif 'постоплат' in t or 'после акт' in t or 'оплата после оказания услуг' in t: put('payment.model','POSTPAY')
    if 'банковская гарантия не' in t or 'без банковской гарантии' in t: put('payment.bank_guarantee',False)
    elif 'банковская гарантия' in t: put('payment.bank_guarantee',True)
    dates=re.findall(r'\b(20\d{2}-\d{2}-\d{2})\b',t)
    if dates: put('term.start',dates[0])
    if len(dates)>1: put('term.end',dates[1])
    months=re.findall(r'\b(20\d{2}-(?:0[1-9]|1[0-2]))\b',t)
    month_names={'январ':1,'феврал':2,'март':3,'апрел':4,'ма':5,'июн':6,'июл':7,'август':8,'сентябр':9,'октябр':10,'ноябр':11,'декабр':12}
    for stem,number in month_names.items():
        match=re.search(stem+r'\w*\s+(20\d{2})',t)
        if match: months.append(f'{match.group(1)}-{number:02d}')
    if months: put('schedule.months',list(dict.fromkeys(months)))
    if 'несколько мест' in t or 'места в договоре не указаны' in t: put('places.multiple_or_not_in_contract',True)
    elif 'одно место' in t: put('places.multiple_or_not_in_contract',False)
    if 'без дополнительных тренировок' in t or 'дополнительная опция не нужна' in t: put('extra_option.enabled',False)
    elif 'дополнительн' in t and ('трениров' in t or 'сопровожд' in t): put('extra_option.enabled',True)
    if 'свыше года' in t or 'более года' in t: put('term.more_than_one_year',True)
    elif 'не более года' in t: put('term.more_than_one_year',False)
    if 'запрос документов о ресурсах' in t: put('ip.long_term_resources_clause',True)
    schedule=[]
    for match in re.finditer(r'(20\d{2}-\d{2}-\d{2})\s*,\s*([^,;]+),\s*([^,;]+),\s*(\d+)\s*,\s*([\d.]+)\s*,\s*([^;]+)',text):
        schedule.append(dict(date=match.group(1),weekday=match.group(2).strip(),time=match.group(3).strip(),count=int(match.group(4)),unit_price=float(match.group(5)),place=match.group(6).strip()))
    if not schedule and months:
        time_match=re.search(r'\b([01]?\d|2[0-3]):[0-5]\d\b',t); count_match=re.search(r'(\d+)\s*трениров',t); price_match=re.search(r'по\s*([\d ]+)\s*(?:руб|₽)',t)
        weekdays=[word for word in ('понедельник','вторник','сред','четверг','пятниц','суббот','воскресень') if word in t]
        if time_match or count_match or price_match or weekdays:
            schedule=[{'date':None,'weekday':' и '.join(weekdays) or None,'time':time_match.group(0) if time_match else None,'count':int(count_match.group(1)) if count_match else None,'unit_price':float(price_match.group(1).replace(' ','')) if price_match else None,'place':None}]
    if schedule: put('schedule.entries',schedule)
    return patch


RED_SCHEMA = [
 ('RF001','contract.form_label','BOTH',True),('RF002','appendices.other','BOTH',False),('RF003','contractor.variant_label','BOTH',True),('RF004','party.status_clause','BOTH',True),('RF005','service.name','BOTH',True),
 ('RF006','price.tax_clause','BOTH',True),('RF007','payment.model','BOTH',True),('RF008','payment.bank_guarantee_clause','BOTH',True),('RF009','physical.ndfl_clause','PHYSICAL_PERSON',True),
 ('RF010','ip.invoice_clause','INDIVIDUAL_ENTREPRENEUR',True),('RF011','npd.receipt_clause','NONE',False),('RF012','npd.loss_status_clause','NONE',False),('RF013','ip.additional_warranties','INDIVIDUAL_ENTREPRENEUR',True),
 ('RF014','ip.vat_penalty','INDIVIDUAL_ENTREPRENEUR',True),('RF015','dispute.forum','BOTH',True),('RF016','ip.long_term_resources','INDIVIDUAL_ENTREPRENEUR',True),('RF017','act.party_status_clause','BOTH',True),
 ('RF018','act.price_tax_clause','BOTH',True),('RF019','act.ndfl_clause','PHYSICAL_PERSON',True),('RF020','schedule.place_column','BOTH',True),('RF021','schedule.months','BOTH',True),
 ('RF022','extra_option.enabled','BOTH',True),('RF023','extra_option.limit','BOTH',True),('RF024','physical.passport_snils','PHYSICAL_PERSON',True),('RF025','ip.ogrnip','INDIVIDUAL_ENTREPRENEUR',True),
]


@dataclass(frozen=True)
class RedFieldDefinition:
    field_id: str
    path: str
    field_type: str
    scope: str
    required: bool
    acquisition: str


# This registry mirrors RED_FIELD_SCHEMA from applicability matrix v2.0.  It is
# deliberately complete: dialogue code never invents an RF outside this table.
RED_FIELD_REGISTRY = {
    field_id: RedFieldDefinition(field_id,path,field_type,scope,required,acquisition)
    for field_id,path,field_type,scope,required,acquisition in (
      ('RF001','contract.form_label','ENUM','BOTH',True,'DEFAULT'),('RF002','appendices.other','LIST','BOTH',False,'DEFAULT'),
      ('RF003','contractor.variant_label','ENUM','BOTH',True,'DERIVED'),('RF004','party.status_clause','STRING','BOTH',True,'DERIVED'),
      ('RF005','service.name','STRING','BOTH',True,'USER_INPUT'),('RF006','price.tax_clause','CONDITIONAL_BLOCK','BOTH',True,'DERIVED'),
      ('RF007','payment.model','ENUM','BOTH',True,'USER_INPUT'),('RF008','payment.bank_guarantee_clause','BOOLEAN','BOTH',True,'USER_INPUT'),
      ('RF009','physical.ndfl_clause','CONDITIONAL_BLOCK','PHYSICAL_PERSON',True,'DERIVED'),('RF010','ip.invoice_clause','CONDITIONAL_BLOCK','INDIVIDUAL_ENTREPRENEUR',True,'DERIVED'),
      ('RF011','npd.receipt_clause','CONDITIONAL_BLOCK','NONE',False,'NOT_APPLICABLE'),('RF012','npd.loss_status_clause','CONDITIONAL_BLOCK','NONE',False,'NOT_APPLICABLE'),
      ('RF013','ip.additional_warranties','CONDITIONAL_BLOCK','INDIVIDUAL_ENTREPRENEUR',True,'DEFAULT'),('RF014','ip.vat_penalty','CONDITIONAL_BLOCK','INDIVIDUAL_ENTREPRENEUR',True,'DERIVED'),
      ('RF015','dispute.forum','ENUM','BOTH',True,'DERIVED'),('RF016','ip.long_term_resources','BOOLEAN','INDIVIDUAL_ENTREPRENEUR',True,'USER_INPUT'),
      ('RF017','act.party_status_clause','STRING','BOTH',True,'DERIVED'),('RF018','act.price_tax_clause','CONDITIONAL_BLOCK','BOTH',True,'DERIVED'),
      ('RF019','act.ndfl_clause','CONDITIONAL_BLOCK','PHYSICAL_PERSON',True,'DERIVED'),('RF020','schedule.place_column','BOOLEAN','BOTH',True,'USER_INPUT'),
      ('RF021','schedule.months','SCHEDULE','BOTH',True,'USER_INPUT'),('RF022','extra_option.enabled','BOOLEAN','BOTH',True,'USER_INPUT'),
      ('RF023','extra_option.limit','INTEGER','BOTH',True,'USER_INPUT'),('RF024','physical.passport_snils','PERSON_DETAILS','PHYSICAL_PERSON',True,'USER_INPUT'),
      ('RF025','ip.ogrnip','STRING','INDIVIDUAL_ENTREPRENEUR',True,'USER_INPUT'),
    )
}


def apply_training_defaults(facts: dict, template: str) -> dict:
    """Apply only values made unambiguous by the approved template or facts."""
    result=deepcopy(facts)
    if template!='NONE':
        if _get(result,'contract.form_label') is None: _set(result,'contract.form_label','TYPICAL_FORM')
        if _get(result,'appendices.other') is None: _set(result,'appendices.other',[])
    start,end=_get(result,'term.start'),_get(result,'term.end')
    if start and end and _get(result,'term.more_than_one_year') is None:
        try: _set(result,'term.more_than_one_year',(date.fromisoformat(end)-date.fromisoformat(start)).days>366)
        except ValueError: pass
    return result


def _input(field_id,path,parser,question,quick=None,mapping=None,allowed=None):
    return {'red_field_id':field_id,'path':path,'parser_type':parser,'question':question,'quick_replies':quick or [],'answer_mapping':mapping or {},'allowed_values':allowed or []}


def missing_red_inputs(decision: TrainingDecision, facts: dict) -> list[dict[str,Any]]:
    """Return atomic USER_INPUT requirements in deterministic business order."""
    template=decision.template; model=_get(facts,'payment.model'); vat=_get(facts,'tax.vat_status')
    specs=[
      _input('RF005','service.name','STRING','Как точно называются услуги по проведению тренировок?'),
      _input('RF005','term.start','DATE','Укажите дату начала оказания услуг.',allowed=['YYYY-MM-DD']),
      _input('RF005','term.end','DATE','Укажите дату окончания оказания услуг.',allowed=['YYYY-MM-DD']),
      _input('RF006','price.total','MONEY','Укажите общую стоимость услуг.'),
    ]
    if template=='PHYSICAL_PERSON': specs += [
      _input('RF009','tax.ndfl_rate','PERCENT','Укажите ставку НДФЛ.'),
      _input('RF009','tax.ndfl_amount','MONEY','Укажите сумму НДФЛ.'),
    ]
    if template=='INDIVIDUAL_ENTREPRENEUR':
      specs += [_input('RF010','tax.vat_status','ENUM','ИП является плательщиком НДС?', ['С НДС','Без НДС'],{'с ндс':'VAT_PAYER','плательщик ндс':'VAT_PAYER','да':'VAT_PAYER','без ндс':'NO_VAT','ндс не облагается':'NO_VAT','нет':'NO_VAT'},['VAT_PAYER','NO_VAT'])]
      if vat=='VAT_PAYER': specs += [_input('RF006','tax.vat_rate','PERCENT','Укажите ставку НДС.'),_input('RF006','tax.vat_amount','MONEY','Укажите сумму НДС.')]
      elif vat=='NO_VAT': specs += [_input('RF006','tax.no_vat_basis','STRING','Укажите основание, по которому НДС не облагается.')]
    specs += [_input('RF007','payment.model','ENUM','Как будет производиться оплата?', ['Предоплата','Смешанная','Постоплата'],{'предоплата':'PREPAY','аванс':'PREPAY','100% аванс':'PREPAY','смешанная':'MIXED','аванс и остаток':'MIXED','часть авансом, остаток после оказания услуг':'MIXED','постоплата':'POSTPAY','оплата после оказания услуг':'POSTPAY','после оказания услуг':'POSTPAY'},['PREPAY','MIXED','POSTPAY'])]
    if model in {'PREPAY','MIXED'}: specs += [_input('RF007','payment.advance_percent','PERCENT','Какой процент составляет аванс?'),_input('RF007','payment.advance_days','INTEGER','В течение скольких дней перечисляется аванс?')]
    if model in {'MIXED','POSTPAY'}: specs += [_input('RF007','payment.postpay_percent','PERCENT','Какой процент оплачивается после оказания услуг?'),_input('RF007','payment.postpay_days','INTEGER','В течение скольких дней производится окончательная оплата?')]
    if model in {'PREPAY','MIXED'}: specs += [_input('RF008','payment.bank_guarantee','BOOLEAN','Требуется банковская гарантия?', ['Да','Нет'])]
    specs += [
      _input('RF020','places.multiple_or_not_in_contract','BOOLEAN','Мест проведения несколько или адрес будет указан только в графике?', ['Да','Нет']),
      _input('RF021','schedule.months','LIST','За какие месяцы нужен график? Укажите в формате ГГГГ-ММ.'),
      _input('RF021','schedule.entries','SCHEDULE','Укажите строки графика: дата, день недели, время, количество, цена одной тренировки, место. Строки разделите точкой с запятой.'),
      _input('RF022','extra_option.enabled','BOOLEAN','Нужны дополнительные тренировки или сопровождение?', ['Да','Нет']),
    ]
    if _get(facts,'extra_option.enabled') is True: specs += [_input('RF023','extra_option.limit','INTEGER','Каково максимальное количество дополнительных услуг?')]
    if template=='PHYSICAL_PERSON': specs += [_input('RF024','physical.passport','PERSON_DETAILS','Укажите паспортные данные исполнителя.'),_input('RF024','physical.snils','PERSON_DETAILS','Укажите СНИЛС исполнителя.')]
    if template=='INDIVIDUAL_ENTREPRENEUR':
      if _get(facts,'term.more_than_one_year') is True: specs += [_input('RF016','ip.long_term_resources_clause','BOOLEAN','Включить запрос документов о ресурсах ИП?', ['Да','Нет'])]
      specs += [_input('RF025','ip.ogrnip','OGRNIP','Укажите 15-значный ОГРНИП.')]
    return [spec for spec in specs if _get(facts,spec['path']) in (None,'',[])]


def _normalized_text(text: str) -> str:
    return ' '.join(text.lower().replace('ё','е').strip(' .!?').split())


def _number(text: str) -> float | None:
    match=re.search(r'(?<!\d)(\d[\d\s]*(?:[,.]\d+)?)',text)
    if not match: return None
    value=float(match.group(1).replace(' ','').replace(',','.'))
    if re.search(r'\bтыс(?:яч[аи]?)?\b',text.lower()): value*=1000
    return value


def parse_pending_red_answer(text: str, pending: dict | None) -> tuple[dict,str,str|None]:
    """Parse any RED_FIELDS pending question using its declared parser contract."""
    if not pending or pending.get('phase')!='RED_FIELDS': return {},'free_text',None
    parser=pending.get('parser_type'); path=(pending.get('target_paths') or [None])[0]; normalized=_normalized_text(text); value=None
    mapping=pending.get('answer_mapping') or {}
    for phrase,mapped in sorted(mapping.items(),key=lambda item:len(item[0]),reverse=True):
        if normalized==phrase or (len(phrase)>3 and phrase in normalized): value=mapped; break
    if parser=='BOOLEAN':
        if normalized in {'да','верно','подтверждаю','есть','нужно','требуется'}: value=True
        elif normalized in {'нет','не требуется','отсутствует','не нужно','не требуется'}: value=False
    elif parser in {'MONEY','PERCENT','INTEGER'}:
        value=_number(normalized)
        if value is not None and parser=='INTEGER': value=int(value)
    elif parser=='DATE':
        match=re.search(r'\b(20\d{2})[-./](\d{1,2})[-./](\d{1,2})\b',normalized) or re.search(r'\b(\d{1,2})[./](\d{1,2})[./](20\d{2})\b',normalized)
        if match:
            parts=list(map(int,match.groups())); y,m,d=parts if parts[0]>2000 else (parts[2],parts[1],parts[0])
            try: value=date(y,m,d).isoformat()
            except ValueError: value=None
    elif parser=='LIST':
        value=list(dict.fromkeys(re.findall(r'\b20\d{2}-(?:0[1-9]|1[0-2])\b',normalized))) or None
    elif parser=='SCHEDULE':
        extracted=extract_training_facts(text)
        value=_get(extracted,'schedule.entries')
    elif parser=='OGRNIP':
        match=re.search(r'\b\d{15}\b',normalized); value=match.group(0) if match else None
    elif parser in {'STRING','PERSON_DETAILS'}:
        value=text.strip(' .') if len(text.strip(' .'))>=2 else None
    if value is None: return {},parser.lower() if parser else 'unknown',f'Не удалось распознать ответ. {pending.get("format_hint") or pending.get("question","")}'
    if parser=='PERCENT' and not 0<=float(value)<=100: return {},'percent','Процент должен быть от 0 до 100.'
    if parser=='INTEGER' and int(value)<=0: return {},'integer','Укажите целое число больше нуля.'
    patch={}; _set(patch,path,value); return patch,parser.lower(),None


def validate_training_facts(facts: dict) -> str | None:
    if _get(facts,'payment.model')=='MIXED':
        advance,postpay=_get(facts,'payment.advance_percent'),_get(facts,'payment.postpay_percent')
        if advance is not None and postpay is not None and abs(float(advance)+float(postpay)-100)>0.001:
            return 'Сумма аванса и окончательной оплаты должна составлять 100%.'
    start,end=_get(facts,'term.start'),_get(facts,'term.end')
    if start and end and start>end: return 'Дата окончания должна быть позже даты начала.'
    return None


def _money(value): return f'{float(value):,.2f}'.replace(',',' ').replace('.',',') if value is not None else None


def red_field_values(facts: dict, template: str) -> list[dict[str,Any]]:
    physical=template=='PHYSICAL_PERSON'; ip=template=='INDIVIDUAL_ENTREPRENEUR'
    values: dict[str,tuple[Any,str]]={}
    direct={'RF001':'contract.form_label','RF002':'appendices.other','RF005':'service.name','RF007':'payment.model','RF020':'places.multiple_or_not_in_contract','RF021':'schedule.months','RF022':'extra_option.enabled','RF023':'extra_option.limit','RF025':'ip.ogrnip'}
    for rf,path in direct.items(): values[rf]=(_get(facts,path),'USER')
    if _get(facts,'contract.form_label')=='TYPICAL_FORM': values['RF001']=('TYPICAL_FORM','DEFAULT')
    if _get(facts,'appendices.other')==[]: values['RF002']=([],'DEFAULT')
    values['RF003']=(template,'DERIVED'); values['RF004']=(('Гражданин РФ' if physical else 'Индивидуальный предприниматель'),'DERIVED')
    values['RF015']=(('GENERAL_COURT' if physical else 'ARBITRATION_COURT'),'DERIVED'); values['RF017']=values['RF004']
    if physical:
        ndfl=f"включая НДФЛ {facts['tax'].get('ndfl_rate')}% в размере {_money(facts['tax'].get('ndfl_amount'))} руб." if facts['tax'].get('ndfl_rate') is not None and facts['tax'].get('ndfl_amount') is not None else None
        values['RF006']=(ndfl,'DERIVED'); values['RF009']=(ndfl,'DERIVED'); values['RF018']=(ndfl,'DERIVED'); values['RF019']=(ndfl,'DERIVED')
        passport=f"Паспорт {facts['physical'].get('passport')}; СНИЛС {facts['physical'].get('snils')}" if facts['physical'].get('passport') and facts['physical'].get('snils') else None
        values['RF024']=(passport,'USER')
    if ip:
        vat=facts['tax'].get('vat_status')
        tax_clause=(f"в том числе НДС {facts['tax'].get('vat_rate')}% в размере {_money(facts['tax'].get('vat_amount'))} руб." if vat=='VAT_PAYER' and facts['tax'].get('vat_rate') is not None and facts['tax'].get('vat_amount') is not None else f"НДС не облагается на основании {facts['tax'].get('no_vat_basis')}" if vat=='NO_VAT' and facts['tax'].get('no_vat_basis') else None)
        values['RF006']=(tax_clause,'DERIVED'); values['RF018']=(tax_clause,'DERIVED'); values['RF010']=(('Счета-фактуры выставляются Исполнителем.' if vat=='VAT_PAYER' else 'Счет-фактура не выставляется.') if vat else None,'DERIVED')
        values['RF013']=('INCLUDE','DEFAULT'); values['RF014']=(vat=='VAT_PAYER','DERIVED')
        values['RF016']=(facts['ip'].get('long_term_resources_clause') if facts['term'].get('more_than_one_year') else False,'USER' if facts['term'].get('more_than_one_year') else 'DEFAULT')
    payment=facts['payment'].get('model')
    guarantee=facts['payment'].get('bank_guarantee') if payment in {'PREPAY','MIXED'} else False
    values['RF008']=(guarantee,'USER' if payment in {'PREPAY','MIXED'} else 'DEFAULT')
    result=[]
    for field_id,path,scope,required in RED_SCHEMA:
        applicable=template!='NONE' and (scope=='BOTH' or scope==template)
        if field_id in {'RF011','RF012'}: applicable=False
        if field_id=='RF008' and payment not in {'PREPAY','MIXED'}: applicable=False
        if field_id=='RF014' and not (ip and facts['tax'].get('vat_status')=='VAT_PAYER'): applicable=False
        if field_id=='RF016' and not (ip and facts['term'].get('more_than_one_year') is True): applicable=False
        if field_id=='RF023' and facts['extra_option'].get('enabled') is not True: applicable=False
        if not applicable: result.append({'field_id':field_id,'status':'NOT_APPLICABLE','value':None,'source':'NONE'}); continue
        value,source=values.get(field_id,(None,'NONE'))
        filled=value is not None and value!='' and (value!=[] or not required)
        status='FILLED' if filled else ('MISSING' if required else 'NOT_APPLICABLE')
        result.append({'field_id':field_id,'status':status,'value':value if filled else None,'source':source if filled else 'NONE'})
    return result


def _validation_missing(facts: dict, template: str) -> list[str]:
    """Map non-red source facts to the canonical red field that depends on them."""
    missing=[]
    def require(field_id: str, *paths: str):
        if any(_get(facts,path) in (None,'',[]) for path in paths) and field_id not in missing: missing.append(field_id)
    require('RF005','service.name','term.start','term.end')
    require('RF006','price.total')
    model=_get(facts,'payment.model')
    require('RF007','payment.model')
    if model=='PREPAY': require('RF007','payment.advance_percent','payment.advance_days')
    elif model=='POSTPAY': require('RF007','payment.postpay_percent','payment.postpay_days')
    elif model=='MIXED': require('RF007','payment.advance_percent','payment.advance_days','payment.postpay_percent','payment.postpay_days')
    require('RF021','schedule.months','schedule.entries')
    entries=_get(facts,'schedule.entries') or []
    if entries and any(any(entry.get(key) in (None,'') for key in ('date','weekday','time','count','unit_price','place')) for entry in entries): missing.append('RF021') if 'RF021' not in missing else None
    if template=='PHYSICAL_PERSON': require('RF009','tax.ndfl_rate','tax.ndfl_amount')
    if template=='INDIVIDUAL_ENTREPRENEUR':
        require('RF010','tax.vat_status')
        if _get(facts,'tax.vat_status')=='VAT_PAYER': require('RF006','tax.vat_rate','tax.vat_amount')
        elif _get(facts,'tax.vat_status')=='NO_VAT': require('RF006','tax.no_vat_basis')
    return missing


REASONS={
 'D001':('UNSUPPORTED_CUSTOMER','Шаблон рассчитан на корпоративного заказчика — АО.'),'D002':('WRONG_SUBJECT','Основной предмет — не проведение тренировок.'),'D003':('NOT_PAID','Форма предназначена для возмездных услуг.'),
 'D004':('LEGAL_ENTITY','Нужен шаблон для юридического лица.'),'D005':('NPD_UNSUPPORTED','Для данного статуса исполнителя текущие шаблоны не применяются. Обратитесь в ДПО.'),
 'D006':('FOREIGN_PERSON','Форма физлица рассчитана на гражданина РФ.'),'D007':('MINOR','Специальные правила несовершеннолетнего не покрыты.'),'D008':('THIRD_PARTIES','Шаблон требует личного оказания и запрещает третьих лиц.'),
 'D009':('LABOR_RELATION','Есть признаки трудовой модели.'),'D010':('STATUS_CONFLICT','Статусные данные противоречат друг другу.'),'D011':('STATUS_UNKNOWN','Исполнитель будет выступать как обычное физическое лицо, ИП или плательщик НПД?'),
 'D012':('SUBJECT_UNKNOWN','Что именно должен сделать исполнитель?'),'D013':('PAID_UNKNOWN','Услуги оплачиваются?'),'D014':('PERFORMANCE_UNKNOWN','Тренировки исполнитель проводит лично и без привлечения других лиц?'),
 'D020':('SELECT_PHYSICAL','Подходит шаблон для физического лица.'),'D021':('SELECT_IP','Подходит шаблон для ИП.'),'D099':('NO_MATCH','Не удалось подобрать разрешённый шаблон. Обратитесь в ДПО.'),
}


def _decision(rule, applicability, template, status, facts) -> TrainingDecision:
    red=red_field_values(facts,template) if template!='NONE' else red_field_values(facts,'NONE')
    missing=list(dict.fromkeys([x['field_id'] for x in red if x['status']=='MISSING']+(_validation_missing(facts,template) if template!='NONE' else [])))
    if template!='NONE' and not missing: rule='D030' if template=='PHYSICAL_PERSON' else 'D031'; status='READY'
    code,reason=REASONS.get(rule,('READY',f'Документ по шаблону {template} готов к формированию.'))
    return TrainingDecision(applicability,template,rule,code,reason,status,missing,red)


def decide_training(facts: dict) -> TrainingDecision:
    status=_get(facts,'contractor.status','UNKNOWN'); subject=_get(facts,'subject.type','UNKNOWN')
    if _get(facts,'customer.type')!='CORPORATE_AO': return _decision('D001','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if subject not in {'TRAINING_SERVICE','UNKNOWN'}: return _decision('D002','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if _get(facts,'subject.paid') is False: return _decision('D003','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if status=='LEGAL_ENTITY': return _decision('D004','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if status in {'NPD_PHYSICAL','IP_NPD'}: return _decision('D005','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if status=='FOREIGN_PHYSICAL' or _get(facts,'contractor.citizenship')=='NON_RU': return _decision('D006','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if status=='MINOR' or _get(facts,'contractor.age_18_plus') is False: return _decision('D007','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if _get(facts,'performance.personal') is False or _get(facts,'performance.third_parties') is True: return _decision('D008','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if any(_get(facts,'labor.'+x) is True for x in ('salary_or_position','internal_rules','fixed_workday','leave_benefits')): return _decision('D009','NOT_APPLICABLE','NONE','BLOCKED',facts)
    if status=='CONFLICT': return _decision('D010','NOT_APPLICABLE','NONE','NEEDS_INPUT',facts)
    if status=='UNKNOWN': return _decision('D011','NOT_APPLICABLE','NONE','NEEDS_INPUT',facts)
    if subject=='UNKNOWN': return _decision('D012','NOT_APPLICABLE','NONE','NEEDS_INPUT',facts)
    if _get(facts,'subject.paid') is None: return _decision('D013','NOT_APPLICABLE','NONE','NEEDS_INPUT',facts)
    if _get(facts,'performance.personal') is None or _get(facts,'performance.third_parties') is None: return _decision('D014','NOT_APPLICABLE','NONE','NEEDS_INPUT',facts)
    if status=='PHYSICAL_NO_IP_NO_NPD': return _decision('D020','APPLICABLE','PHYSICAL_PERSON','NEEDS_INPUT',facts)
    if status=='IP_NO_NPD': return _decision('D021','APPLICABLE','INDIVIDUAL_ENTREPRENEUR','NEEDS_INPUT',facts)
    return _decision('D099','NOT_APPLICABLE','NONE','BLOCKED',facts)


def next_question(decision: TrainingDecision, facts: dict) -> str:
    if decision.template=='NONE': return decision.reason
    questions={'RF001':'Какая форма договора используется: типовая или иная?','RF005':'Как точно называются услуги?','RF006':'Уточните налоговые показатели исполнителя.','RF007':'Оплата: аванс, аванс+остаток или постоплата?','RF008':'Требуется банковская гарантия?','RF009':'Укажите ставку и сумму НДФЛ.','RF010':'ИП является плательщиком НДС?','RF014':'Является ли ИП плательщиком НДС?','RF016':'Включить запрос документов о ресурсах ИП?','RF020':'Мест проведения несколько или адрес не указан в договоре?','RF021':'За какие месяцы составляется график?','RF022':'Нужны дополнительные тренировки или сопровождение?','RF023':'Каков максимальный объём дополнительной опции?','RF024':'Укажите паспортные данные и СНИЛС исполнителя.','RF025':'Укажите ОГРНИП.'}
    return questions.get(decision.missing_field_ids[0],f'Уточните поле {decision.missing_field_ids[0]}.') if decision.missing_field_ids else 'Данные заполнены. Документ готов к формированию.'


def clarification_for(decision: TrainingDecision, facts: dict, previous: dict | None=None) -> dict | None:
    """Build an explicit, persisted clarification contract for the next answer."""
    personal=_get(facts,'performance.personal'); third=_get(facts,'performance.third_parties')
    if decision.matched_rule_id=='D014':
        if personal is None and third is None:
            item={'question_id':'TRAINING_PERFORMANCE_COMPOSITE','question':'Тренировки исполнитель проводит лично и без привлечения других лиц?','target_fields':['performance.personal','performance.third_parties'],'answer_mapping':{'yes':{'performance':{'personal':True,'third_parties':False}}}}
        elif personal is None:
            item={'question_id':'TRAINING_PERFORMANCE_PERSONAL','question':'Исполнитель будет проводить тренировки лично?','target_fields':['performance.personal'],'answer_mapping':{'yes':{'performance':{'personal':True}},'no':{'performance':{'personal':False}}}}
        else:
            item={'question_id':'TRAINING_PERFORMANCE_THIRD_PARTIES','question':'Будут ли привлекаться другие лица для проведения тренировок?','target_fields':['performance.third_parties'],'answer_mapping':{'yes':{'performance':{'third_parties':True}},'no':{'performance':{'third_parties':False}}}}
    elif decision.template!='NONE' and decision.missing_field_ids:
        specs=missing_red_inputs(decision,facts)
        if not specs: return None
        spec=specs[0]; field_id=spec['red_field_id']
        item={'question_id':f'training:{field_id}:{spec["path"]}','phase':'RED_FIELDS','red_field_id':field_id,'question':spec['question'],'target_fields':[field_id],'target_paths':[spec['path']],'parser_type':spec['parser_type'],'allowed_values':spec['allowed_values'],'answer_mapping':spec['answer_mapping'],'quick_replies':spec['quick_replies']}
    else:
        targets={'D011':['contractor.status'],'D012':['subject.type'],'D013':['subject.paid']}.get(decision.matched_rule_id,[])
        if not targets: return None
        item={'question_id':'TRAINING_'+decision.matched_rule_id,'question':decision.reason,'target_fields':targets,'answer_mapping':{}}
    same=bool(previous and previous.get('question_id')==item['question_id'] and previous.get('target_fields')==item['target_fields'])
    item['repeat_count']=(int(previous.get('repeat_count',0))+1) if same else 0
    if same and item['repeat_count']==1 and item['question_id']=='TRAINING_PERFORMANCE_COMPOSITE':
        item['question']='Подтвердите отдельно: исполнитель сам проводит тренировки и других лиц не привлекает?'
    if same and item['repeat_count']>=2 and len(item['target_fields'])>1:
        item=clarification_for(decision,{**facts,'performance':{**facts['performance'],'personal':facts['performance']['personal'] if facts['performance']['personal'] is not None else True}},None) or item
    elif same and item['repeat_count']==1 and item.get('phase')=='RED_FIELDS':
        choices=item.get('quick_replies') or item.get('allowed_values')
        item['question']=item['question']+(' Выберите: '+', '.join(map(str,choices))+'.' if choices else ' Проверьте формат ответа и укажите значение ещё раз.')
    elif same and item['repeat_count']>=2 and item.get('phase')=='RED_FIELDS':
        item['question_id']=item['question_id']+':validation'
        item['question']='Ответ не прошёл проверку. '+item['question']
    return item


def resolve_contextual_answer(text: str, pending: dict | None) -> tuple[dict, str, str | None]:
    """Resolve short replies only against the clarification that produced them."""
    normalized=' '.join(text.lower().replace('ё','е').strip(' .!?').split())
    answer_type='free_text'; key=None
    if normalized in {'да','верно','именно','подтверждаю','правильно'}: key='yes'; answer_type='contextual_yes'
    elif normalized in {'нет','неверно'}: key='no'; answer_type='contextual_no'
    if not key or not pending: return {},answer_type,None
    mapping=(pending.get('answer_mapping') or {}).get(key)
    if mapping is not None: return mapping,answer_type,None
    if key=='no' and pending.get('question_id')=='TRAINING_PERFORMANCE_COMPOSITE':
        return {},answer_type,'Уточните, пожалуйста: исполнитель сам не проводит тренировки или планируется привлечение других лиц?'
    return {},answer_type,None


def training_root(root: Path) -> Path: return root/'training'


def resolve_source(root: Path, names: tuple[str,...]) -> Path:
    base=training_root(root)
    for name in names:
        candidate=(base/name).resolve()
        if candidate.is_relative_to(base.resolve()) and candidate.is_file(): return candidate
    raise FileNotFoundError(', '.join(names))


def load_occurrences(root: Path) -> list[dict[str,Any]]:
    matrix=resolve_source(root,MATRIX_NAMES); workbook=load_workbook(matrix,data_only=True,read_only=True)
    rows=workbook['RED_OCCURRENCES'].iter_rows(values_only=True); headers=next(rows)
    result=[dict(zip(headers,row)) for row in rows]
    if len(result)!=74 or any(row['mapping_status']!='MAPPED' for row in result): raise ValueError('RED_OCCURRENCES mapping is incomplete')
    return result


def _paragraph_by_location(document, location: str):
    if location.startswith('P'): return document.paragraphs[int(location[1:])]
    match=re.fullmatch(r'T(\d+)R(\d+)C(\d+)P(\d+)',location)
    if not match: raise ValueError(location)
    ti,ri,ci,pi=map(int,match.groups()); return document.tables[ti].rows[ri].cells[ci].paragraphs[pi]


def _replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    full=''.join(run.text for run in paragraph.runs)
    if old not in full and old.strip('[]()') not in full:
        red_runs=[run for run in paragraph.runs if run.font.color.rgb and str(run.font.color.rgb).upper() in {'FF0000','F00000','C00000'}]
        if not red_runs: return False
        red_runs[0].text=new
        red_runs[0].font.color.rgb=None
        for run in red_runs[1:]: run.text=''
        return True
    target=old if old in full else old.strip('[]()'); replaced=full.replace(target,new)
    first=paragraph.runs[0] if paragraph.runs else paragraph.add_run(); first.text=replaced
    first.font.color.rgb=None
    for run in paragraph.runs[1:]: run.text=''
    return True


def render_training_contract(root: Path, facts: dict, decision: TrainingDecision, output: Path) -> dict:
    if decision.render_status!='READY' or decision.template=='NONE': raise ValueError('Document is not READY')
    source=resolve_source(root,TEMPLATES[decision.template]); original=source.read_bytes(); document=Document(source)
    red={x['field_id']:x for x in decision.red_fields}; occurrences=load_occurrences(root); errors=[]
    for occurrence in occurrences:
        field=red[occurrence['canonical_field_id']]; paragraph=_paragraph_by_location(document,occurrence['docx_location'])
        value=field['value']; replacement=''
        if field['status']=='FILLED':
            replacement=', '.join(map(str,value)) if isinstance(value,list) else ('Да' if value is True else 'Нет' if value is False else str(value))
        # The approved physical/IP source files already pre-resolve some alternatives
        # from the universal matrix.  In that case the canonical location remains
        # valid, but neither the universal red wording nor a red run is present.
        _replace_in_paragraph(paragraph,str(occurrence['red_text']),replacement)
    if errors: raise ValueError('Unmapped occurrences: '+','.join(errors))
    # Fill schedule rows conservatively without rebuilding the approved tables.
    entries=facts.get('schedule',{}).get('entries') or []
    for table_index in (1,6):
        if table_index >= len(document.tables): continue
        table=document.tables[table_index]
        for index,entry in enumerate(entries):
            row_index=2+index
            if row_index>=len(table.rows): break
            values=[entry.get(k,'') for k in ('date','weekday','time','count','unit_price','place')]
            for cell,value in zip(table.rows[row_index].cells,values): cell.text=str(value)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.color.rgb and str(run.font.color.rgb).upper() in {'FF0000','F00000','C00000'}: run.text=''; run.font.color.rgb=None
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.color.rgb and str(run.font.color.rgb).upper() in {'FF0000','F00000','C00000'}: run.text=''; run.font.color.rgb=None
    output.parent.mkdir(parents=True,exist_ok=True); document.save(output)
    if source.read_bytes()!=original: raise RuntimeError('Source template was modified')
    return {'source':source,'occurrences':len(occurrences),'mapping_errors':errors}


def audit_payload(decision: TrainingDecision) -> dict:
    return {'applicability':decision.applicability,'template':decision.template,'matched_rule_id':decision.matched_rule_id,'reason_code':decision.reason_code,'render_status':decision.render_status,'missing_field_ids':decision.missing_field_ids}
