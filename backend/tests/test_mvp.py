import io
import os
import tempfile
import hashlib
import shutil
import zipfile
from urllib.parse import quote
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

os.environ['DATABASE_URL'] = 'sqlite:///' + tempfile.mktemp(suffix='.db')
os.environ['STORAGE_LOCAL_PATH'] = tempfile.mkdtemp()
os.environ['SEED_DEMO_USERS'] = 'true'
template_directory = Path(tempfile.mkdtemp())
os.environ['CONTRACT_TEMPLATES_PATH'] = str(template_directory)
specification_template_directory = Path(tempfile.mkdtemp())
os.environ['SPECIFICATION_TEMPLATES_PATH'] = str(specification_template_directory)
os.environ['CONTRACT_TEMPLATE_ALLOW_DYNAMIC_CHECKSUMS'] = 'true'
for filename in ('Договор_подряда.docx','договор возмездного оказания услуг (образец).docx','Shablon-agentskogo-dogovora.docx'):
    fixture=Document(); fixture.add_paragraph(filename); fixture.save(template_directory/filename)
(template_directory/'договор поставки.rtf').write_bytes(b'{\\rtf1\\ansi approved supply fixture}')
training_fixture_root=Path('/app/contract_templates/training') if Path('/app/contract_templates/training').is_dir() else Path(__file__).resolve().parents[2]/'contract_templates'/'training'
shutil.copytree(training_fixture_root,template_directory/'training')
supply_specification_source=Path(__file__).resolve().parents[2]/'specification_templates'/'supply'
shutil.copytree(supply_specification_source,specification_template_directory/'supply')

from fastapi.testclient import TestClient
import fitz
import openpyxl
import xlwt
from app.db import Base, engine, SessionLocal, User, QuestionnaireField
from app.documents.service import SpreadsheetSpecificationParser
from app.ai.dialogue import clarification

Base.metadata.create_all(engine)
from app.seed import seed
seed()
from app.main import app


def auth(client: TestClient, name: str = 'ivanov') -> dict[str, str]:
    response = client.post('/api/v1/auth/login', json={'external_id': name, 'password': name.capitalize()})
    assert response.status_code == 200
    return {'x-csrf-token': response.json()['csrf_token']}


def contract_docx_bytes() -> bytes:
    document=Document(); paragraph=document.add_paragraph(); paragraph.add_run('Заказчик: '); paragraph.add_run('АО '); paragraph.add_run('«НСПК»')
    table=document.add_table(rows=2,cols=2); table.cell(0,0).merge(table.cell(0,1)).text='Условия договора'; table.cell(1,0).text='Срок'; table.cell(1,1).text='30 дней'
    document.sections[0].header.paragraphs[0].text='Корпоративный договор'; document.sections[0].footer.paragraphs[0].text='Конфиденциально'
    payload=io.BytesIO(); document.save(payload); return payload.getvalue()


def create_platform_template(client,headers,code='TRAINING_POC'):
    response=client.post('/api/v1/document-templates',headers=headers,json={'code':code,'name':'Договор для проверки'})
    assert response.status_code==201,response.text
    return response.json()


def configure_platform_field(client,headers,template,version,field_key='customer_name',label='Наименование заказчика',occurrences=1):
    field=client.post(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/fields",headers=headers,json={'field_key':field_key,'label':label,'field_type':'TEXT','required':True,'position':1,'group':'Стороны'}).json()
    preview=client.get(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/preview",headers=headers).json()
    matches=[]
    for node in preview['nodes']:
        start=0
        while (index:=node['text'].find('АО «НСПК»',start))>=0:
            matches.append((node,index)); start=index+1
    assert len(matches)>=occurrences
    for node,index in matches[:occurrences]:
        response=client.post(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/bindings",headers=headers,json={'field_id':field['id'],'locator':node['locator'],'start':index,'end':index+len('АО «НСПК»'),'selected_text':'АО «НСПК»'})
        assert response.status_code==201,response.text
    return field


def repeating_contract_docx_bytes(vertical_merge=False) -> bytes:
    document=Document(); document.add_paragraph('Заказчик: CUSTOMER')
    table=document.add_table(rows=3,cols=2); table.style='Table Grid'
    table.rows[0].cells[0].text='Дата'; table.rows[0].cells[1].text='Место'
    date_cell=table.rows[1].cells[0]; date_cell.text=''; p=date_cell.paragraphs[0]; p.add_run('DA'); p.add_run('TE'); date_cell.add_paragraph('Дополнительная строка')
    table.rows[1].cells[1].text='PLACE'
    table.rows[2].cells[0].merge(table.rows[2].cells[1]).text='ИТОГО'
    if vertical_merge:
        properties=table.rows[1].cells[0]._tc.get_or_add_tcPr(); marker=OxmlElement('w:vMerge'); marker.set(qn('w:val'),'restart'); properties.append(marker)
    payload=io.BytesIO(); document.save(payload); return payload.getvalue()


def configure_repeating_platform(client,headers,code,vertical_merge=False,min_items=0,max_items=10,visible_when=None):
    template=create_platform_template(client,headers,code)
    raw=repeating_contract_docx_bytes(vertical_merge)
    version=client.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('repeat.docx',raw,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
    preview=client.get(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/preview",headers=headers).json()
    node=lambda text: next(x for x in preview['nodes'] if text in x['text'])
    customer=client.post(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/fields",headers=headers,json={'field_key':'customer_name','label':'Заказчик','field_type':'TEXT','required':True,'position':1}).json(); customer_node=node('CUSTOMER'); customer_start=customer_node['text'].index('CUSTOMER')
    linked=client.post(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/bindings",headers=headers,json={'field_id':customer['id'],'locator':customer_node['locator'],'start':customer_start,'end':customer_start+8,'selected_text':'CUSTOMER'}); assert linked.status_code==201,linked.text
    date_node=node('DATE')
    response=client.post(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/repeating-groups",headers=headers,json={'group_key':'schedule','name':'График тренировок','description':'Добавьте тренировки','min_items':min_items,'max_items':max_items,'required':False,'position':1,'row_locator':date_node['locator'],'visible_when':visible_when})
    if vertical_merge: return template,version,raw,response,None
    assert response.status_code==201,response.text
    group=response.json(); fields=[]
    for position,(key,label,kind,needle) in enumerate((('date','Дата','DATE','DATE'),('place','Место','TEXT','PLACE')),1):
        field=client.post(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/repeating-groups/{group['id']}/fields",headers=headers,json={'field_key':key,'label':label,'field_type':kind,'required':True,'position':position}).json(); target=node(needle); start=target['text'].index(needle)
        bound=client.put(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/repeating-groups/{group['id']}/fields/{field['id']}/binding",headers=headers,json={'locator':target['locator'],'start':start,'end':start+len(needle),'selected_text':needle})
        assert bound.status_code==200,bound.text; fields.append(bound.json())
    return template,version,raw,response,group


def test_clarification_asks_only_one_atomic_field():
    question, fields, _ = clarification(['purpose', 'deadline'], [], 0)
    assert fields == ['purpose']
    assert 'цель' in question.lower() and 'срок' not in question.lower()


def test_clarification_never_exposes_internal_field_name():
    question, _, _ = clarification(['subject'], [], 0)
    assert question == 'Укажите предмет закупки.' and 'subject' not in question.lower()
    fallback, _, _ = clarification(['internal_unknown'], [], 0)
    assert fallback == 'Уточните, пожалуйста, сведения о закупке.' and 'internal_unknown' not in fallback


def test_full_flow_and_versions():
    client = TestClient(app); headers = auth(client)
    case_id = client.post('/api/v1/cases', headers=headers, json={'title': 'Лендинг'}).json()['id']
    response = client.post(f'/api/v1/cases/{case_id}/messages', headers=headers, json={'content': 'Нужно заказать разработку лендинга для мероприятия'})
    assert response.json()['state'] == 'COLLECTING_SPEC_DATA'
    for answer in ['Продвижение мероприятия', 'Создать адаптивный лендинг', '30 ноября 2026 года', 'Соответствие требованиям']:
        response = client.post(f'/api/v1/cases/{case_id}/messages', headers=headers, json={'content': answer})
    assert response.json()['state'] == 'SPEC_REVIEW'
    spec = client.patch(f'/api/v1/cases/{case_id}/technical-specification', headers=headers, json={'changes': [{'field': 'deadline', 'value': '2026-11-30'}]})
    assert spec.json()['version'] == 3 and spec.json()['data']['purpose'] == 'Продвижение мероприятия'
    recommendation = client.post(f'/api/v1/cases/{case_id}/technical-specification/confirm', headers=headers)
    assert recommendation.json()['status'] == 'matched'
    assert client.get(f'/api/v1/cases/{case_id}/technical-specification/download').status_code == 200


def test_permissions_and_escalation():
    ivan = TestClient(app); ivan_headers = auth(ivan)
    case_id = ivan.post('/api/v1/cases', headers=ivan_headers, json={}).json()['id']
    anna = TestClient(app); auth(anna, 'petrova')
    assert anna.get(f'/api/v1/cases/{case_id}').status_code == 404
    response = ivan.post(f'/api/v1/cases/{case_id}/messages', headers=ivan_headers, json={'content': 'Мне нужен личный юридический совет'})
    assert response.json()['state'] == 'ESCALATED_TO_DPO'


def test_upload_docx_analysis_and_contract_download():
    document = Document(); document.add_paragraph('Техническое задание: разработка корпоративного лендинга. Срок 30 ноября 2026 года.')
    buffer = io.BytesIO(); document.save(buffer)
    client = TestClient(app); headers = auth(client)
    case_id = client.post('/api/v1/cases', headers=headers, json={'title': 'Готовое ТЗ'}).json()['id']
    response = client.post(f'/api/v1/cases/{case_id}/attachments', headers=headers, files={'file': ('tz.docx', buffer.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code == 200 and response.json()['analysis_ready']
    recommendation = client.post(f'/api/v1/cases/{case_id}/technical-specification/confirm', headers=headers)
    assert recommendation.json()['status'] == 'matched'
    assert client.get(f'/api/v1/cases/{case_id}/contract-template').status_code == 200


def test_sessions_csrf_ownership_and_switching():
    ivan = TestClient(app); ivan_headers = auth(ivan)
    assert ivan.get('/api/v1/auth/me').json()['user']['email'] == 'ivanov@nspk.local'
    case_id = ivan.post('/api/v1/cases', headers=ivan_headers, json={}).json()['id']
    assert ivan.post('/api/v1/cases', json={}).status_code == 403
    ivan.post(f'/api/v1/cases/{case_id}/messages', headers=ivan_headers, json={'content': 'Нам нужно купить велосипеды'})
    assert ivan.get(f'/api/v1/cases/{case_id}').json()['title'] == 'Закупка велосипедов'

    anna = TestClient(app); anna_headers = auth(anna, 'petrova')
    assert case_id not in [item['id'] for item in anna.get('/api/v1/cases').json()]
    attempts = [
        anna.get(f'/api/v1/cases/{case_id}'),
        anna.patch(f'/api/v1/cases/{case_id}', headers=anna_headers, json={'title': 'чужое'}),
        anna.post(f'/api/v1/cases/{case_id}/messages', headers=anna_headers, json={'content': 'чужое'}),
        anna.post(f'/api/v1/cases/{case_id}/attachments', headers=anna_headers, files={'file': ('x.docx', b'x', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}),
        anna.get(f'/api/v1/cases/{case_id}/technical-specification/download'),
    ]
    assert all(response.status_code == 404 for response in attempts)
    anna_case = anna.post('/api/v1/cases', headers=anna_headers, json={'title': 'Обучение сотрудников'}).json()['id']
    assert anna.post('/api/v1/auth/logout', headers=anna_headers).status_code == 200
    assert anna.get('/api/v1/auth/me').status_code == 401
    assert case_id in [item['id'] for item in ivan.get('/api/v1/cases').json()]
    assert anna_case not in [item['id'] for item in ivan.get('/api/v1/cases').json()]


def test_login_cookie_and_rate_limit():
    client = TestClient(app)
    response = client.post('/api/v1/auth/login', json={'external_id': 'Ivanov', 'password': 'Ivanov'})
    assert response.status_code == 200 and 'session_id' not in response.json()
    cookie = response.headers['set-cookie'].lower()
    assert 'httponly' in cookie and 'samesite=lax' in cookie
    for _ in range(5):
        assert client.post('/api/v1/auth/login', json={'external_id': 'rate-limit-probe', 'password': 'wrong-password'}).status_code == 401
    assert client.post('/api/v1/auth/login', json={'external_id': 'rate-limit-probe', 'password': 'wrong-password'}).status_code == 429


def spreadsheet_bytes() -> bytes:
    workbook = openpyxl.Workbook(); sheet = workbook.active; sheet.title = 'Спецификация'
    sheet['B1'] = 'Спецификация'
    for column,value in enumerate(['Наименование','Характеристики','Кол-во','Стоимость (руб., c учетом НДС)','Срок поставки (рабочих дней)','Гарантия','Ссылки на ресурс, при наличии','Адрес поставки'],2): sheet.cell(2,column,value)
    first=['Ноутбук Альфа','RAM 16 ГБ; SSD 512 ГБ',3,300000,20,'24 месяца','https://example.test/notebook-a\nhttps://example.test/notebook-b','Москва']
    second=['Ноутбук Бета','RAM 32 ГБ; SSD 1 ТБ',2,250000,20,'36 месяцев','','Москва']
    for row_index,values in enumerate((first,second),3):
        sheet.cell(row_index,1,row_index-2)
        for column,value in enumerate(values,2): sheet.cell(row_index,column,value)
    sheet['H3'].hyperlink = 'https://example.test/notebook-a'
    sheet['C5']='Блокирующие требования:\n- Не допускается восстановленное оборудование\n- Без замены на аналоги'
    requirements = workbook.create_sheet('Требования')
    requirements.append(['Общие требования','Только поставка, без установки'])
    requirements.append(['Срок поставки','20 рабочих дней'])
    buffer = io.BytesIO(); workbook.save(buffer); return buffer.getvalue()


def test_contract_only_without_specification_and_download():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    response=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Хотим купить 5 ноутбуков у юридического лица, только поставка.'})
    body=response.json(); assert body['technical_specification'] is None
    assert body['procurement_context']['intent']=='GET_CONTRACT' and body['procurement_context']['installation_required'] is False
    assert body['recommendation']['status']=='MATCHED'
    assert client.get(f'/api/v1/cases/{case_id}/contract-template').status_code==200


def test_contract_ambiguity_then_clarification_and_legal_review():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    first=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Хотим купить 5 ноутбуков у юридического лица.'}).json()
    assert first['technical_specification'] is None and first['recommendation']['template_code']=='SUPPLY'
    second=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Только поставка, без установки.'}).json()
    assert second['recommendation']['status']=='MATCHED'
    review_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    review=client.post(f'/api/v1/cases/{review_id}/messages',headers=headers,json={'content':'Нужен договор, требуется юридическая экспертиза.'}).json()
    assert review['state']=='ESCALATED_TO_DPO'


def test_spreadsheet_parser_xlsx_and_xls():
    with tempfile.TemporaryDirectory() as directory:
        xlsx=Path(directory)/'sample.xlsx'; xlsx.write_bytes(spreadsheet_bytes())
        parsed=SpreadsheetSpecificationParser().parse(xlsx,'.xlsx')
        assert len(parsed['items'])==2 and parsed['items'][0]['quantity']==3
        assert '16 ГБ' in parsed['items'][0]['characteristics']
        assert len(parsed['items'][0]['references'].splitlines())==2
        assert any('восстановленное' in item for item in parsed['blocking_requirements'])
        assert parsed['delivery_terms']=='20 рабочих дней'
        book=xlwt.Workbook(); sheet=book.add_sheet('Спецификация')
        for column,value in enumerate(['Наименование','Количество','Единица измерения','Характеристики']): sheet.write(0,column,value)
        for column,value in enumerate(['Ноутбук',4,'шт.','RAM 16 ГБ']): sheet.write(1,column,value)
        xls=Path(directory)/'sample.xls'; book.save(str(xls))
        assert SpreadsheetSpecificationParser().parse(xls,'.xls')['items'][0]['quantity']==4


def test_xlsx_upload_one_file_and_corruption():
    client=TestClient(app); headers=auth(client)
    bad_id=client.post('/api/v1/cases',headers=headers,json={}).json()['id']
    bad=client.post(f'/api/v1/cases/{bad_id}/attachments',headers=headers,files={'file':('broken.xlsx',b'not-a-workbook','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert bad.status_code==422
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'ANALYZE_SPECIFICATION'}).json()['id']
    upload=client.post(f'/api/v1/cases/{case_id}/attachments',headers=headers,files={'file':('sample.xlsx',spreadsheet_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert upload.status_code==200
    spec=upload.json()['case']['technical_specification']['data']; assert len(spec['items'])==2
    second=client.post(f'/api/v1/cases/{case_id}/attachments',headers=headers,files={'file':('second.xlsx',spreadsheet_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert second.status_code==409


def test_xlsx_context_selects_same_approved_supply_template():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'ANALYZE_SPECIFICATION'}).json()['id']
    upload=client.post(f'/api/v1/cases/{case_id}/attachments',headers=headers,files={'file':('sample.xlsx',spreadsheet_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert upload.status_code==200
    body=upload.json()['case']
    assert body['technical_specification']['data']['kind']=='SUPPLY'
    assert body['procurement_context']['specification_template']=='SUPPLY'


def test_consistent_docx_pdf_xlsx_exports():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'ANALYZE_SPECIFICATION'}).json()['id']
    client.post(f'/api/v1/cases/{case_id}/attachments',headers=headers,files={'file':('sample.xlsx',spreadsheet_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    # Supply output is intentionally limited to the approved DOCX and cannot
    # be generated before the imported data has passed validation.
    assert client.get(f'/api/v1/cases/{case_id}/technical-specification/download?format=docx').status_code==409


def test_search_delete_and_ownership():
    ivan=TestClient(app); ih=auth(ivan); anna=TestClient(app); ah=auth(anna,'petrova')
    own_case=ivan.post('/api/v1/cases',headers=ih,json={'title':'Секретные велосипеды'}).json()['id']
    foreign=anna.post('/api/v1/cases',headers=ah,json={'title':'Секретные ноутбуки'}).json()['id']
    results=ivan.get('/api/v1/cases?q=Секретные').json()
    assert own_case in [item['id'] for item in results] and foreign not in [item['id'] for item in results]
    assert ivan.delete(f'/api/v1/cases/{foreign}',headers=ih).status_code==404
    assert ivan.delete(f'/api/v1/cases/{own_case}',headers=ih).status_code==204
    assert ivan.get(f'/api/v1/cases/{own_case}').status_code==404


def test_four_approved_contracts_and_original_download_bytes():
    scenarios={
        'SUPPLY':('Нужен договор: купить 10 ноутбуков, только поставка без установки','договор поставки.rtf','application/rtf'),
        'SERVICES':('Нужен договор: заказать консультационные услуги','договор возмездного оказания услуг (образец).docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
        'WORKS':('Нужен договор: выполнить ремонт помещения','Договор_подряда.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
        'AGENCY':('Нужен договор: агент должен по поручению закупить материалы у третьих лиц, доставить их и предоставить отчёт агента за вознаграждение','Shablon-agentskogo-dogovora.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
    }
    client=TestClient(app); headers=auth(client)
    for code,(message,filename,mime) in scenarios.items():
        case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':message}).json()
        assert body['recommendation']['status']=='MATCHED'
        assert body['recommendation']['template_code']==code
        response=client.get(f'/api/v1/cases/{case_id}/contract-template')
        assert response.status_code==200 and response.headers['content-type'].startswith(mime)
        assert quote(filename) in response.headers['content-disposition']
        assert hashlib.sha256(response.content).hexdigest()==hashlib.sha256((template_directory/filename).read_bytes()).hexdigest()


def test_rules_fallbacks_are_deterministic():
    from app.contracts import select_contract
    from app.db import SessionLocal
    db=SessionLocal()
    base={'subject':'Закупка'}
    assert select_contract(db,{**base,'requires_transfer_of_goods':True,'requires_tangible_work_result':True}).status=='AMBIGUOUS'
    assert select_contract(db,{}).status=='AMBIGUOUS'
    assert select_contract(db,{**base,'category':'software_license','licensing_required':True}).status=='NO_MATCH'
    assert select_contract(db,{**base,'exceptions':['Нестандартные условия']}).status=='LEGAL_REVIEW_REQUIRED'
    db.close()


def test_training_dialogue_persists_facts_and_selects_physical_template():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    first=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нужен тренер для проведения платных тренировок, лично, без третьих лиц.'}).json()
    assert first['state']=='TRAINING_APPLICABILITY' and first['procurement_context']['training_facts']['subject']['type']=='TRAINING_SERVICE'
    assert first['procurement_context']['contract_domain']=='training'
    second=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Обычное физлицо без ИП, гражданин РФ, совершеннолетний.'}).json()
    assert second['state']=='TRAINING_TEMPLATE_SELECTED' and second['recommendation']['template_code']=='PHYSICAL_PERSON'
    assert second['procurement_context']['training_decision']['matched_rule_id']=='D020'


def test_contract_router_keeps_supply_and_services_in_generic_engine():
    client=TestClient(app); headers=auth(client)
    supply_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    supply=client.post(f'/api/v1/cases/{supply_id}/messages',headers=headers,json={'content':'Нужно купить ноутбуки, только поставка без установки.'}).json()
    assert supply['recommendation']['template_code']=='SUPPLY' and supply['procurement_context'].get('contract_domain')!='training'
    service_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    service=client.post(f'/api/v1/cases/{service_id}/messages',headers=headers,json={'content':'Нужно заказать консультационные услуги.'}).json()
    assert service['recommendation']['template_code']=='SERVICES' and service['procurement_context'].get('contract_domain')!='training'


def test_contract_router_training_ip_npd_and_clarification():
    client=TestClient(app); headers=auth(client)
    for phrase,template,status in [
        ('Нужен тренер. ИП без НПД, лично, без третьих лиц.','INDIVIDUAL_ENTREPRENEUR','NEEDS_INPUT'),
        ('Нужен тренер. Исполнитель самозанятый плательщик НПД.','', 'BLOCKED'),
    ]:
        case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':phrase}).json()
        assert body['procurement_context']['contract_domain']=='training'
        assert body['recommendation']['template_code']==template and body['recommendation']['status']==status
    unclear_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    unclear=client.post(f'/api/v1/cases/{unclear_id}/messages',headers=headers,json={'content':'Нужно организовать занятия сотрудникам.'}).json()
    assert unclear['procurement_context']['pending_contract_domain']=='training'
    confirmed=client.post(f'/api/v1/cases/{unclear_id}/messages',headers=headers,json={'content':'Да'}).json()
    assert confirmed['procurement_context']['contract_domain']=='training'


def test_training_composite_yes_and_explicit_answers_via_api():
    client=TestClient(app); headers=auth(client)
    for answer in ('Да','ЛИЧНО БЕЗ ПРИВЛЕЧЕНИЯ ДРУГИХ ЛИЦ'):
        case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
        first=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нужен тренер для проведения тренировок, исполнитель физлицо без ИП и НПД'}).json()
        pending=first['procurement_context']['pending_clarification']
        assert pending['question_id']=='TRAINING_PERFORMANCE_COMPOSITE'
        assert first['procurement_context']['training_facts']['contractor']['status']=='PHYSICAL_NO_IP_NO_NPD'
        previous_question=first['messages'][-1]['content']
        second=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':answer}).json()
        persisted=client.get(f'/api/v1/cases/{case_id}').json()
        performance=persisted['procurement_context']['training_facts']['performance']
        assert persisted['procurement_context']['training_facts']['contractor']['status']=='PHYSICAL_NO_IP_NO_NPD'
        assert performance['personal'] is True and performance['third_parties'] is False
        assert persisted['messages'][-1]['content']!=previous_question
        assert not {'performance.personal','performance.third_parties'} & set(persisted['procurement_context']['pending_clarification']['target_fields'])
        assert persisted['recommendation']['contract_status_text']=='Подходит шаблон для физического лица',persisted['procurement_context']['training_decision']


def test_training_composite_no_requests_disambiguation_without_inference():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нужен тренер для тренировок, исполнитель физлицо без ИП и НПД'})
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нет'}).json()
    assert body['procurement_context']['training_facts']['performance']=={'personal':None,'third_parties':None}
    assert body['procurement_context']['pending_clarification']['question_id']=='TRAINING_PERFORMANCE_DISAMBIGUATION'
    assert 'Уточните' in body['messages'][-1]['content']


def run_training_dialogue(client,headers,initial,vat='NO_VAT',schedule_answer='2026-11-10, вторник, 10:00, 1, 120000, Москва'):
    answers={
      'service.name':'Проведение тренировок по плаванию','term.start':'2026-11-01','term.end':'2027-01-31',
      'price.total':'120 000 рублей','tax.ndfl_rate':'13%','tax.ndfl_amount':'15 600 рублей',
      'tax.vat_status':'Без НДС' if vat=='NO_VAT' else 'С НДС','tax.no_vat_basis':'Применение УСН',
      'tax.vat_rate':'20%','tax.vat_amount':'20 000 рублей','payment.model':'Постоплата',
      'payment.postpay_percent':'100%','payment.postpay_days':'10 дней','places.multiple_or_not_in_contract':'Нет',
      'schedule.months':'2026-10','schedule.entries':schedule_answer,
      'extra_option.enabled':'Нет','physical.passport':'45 01 123456, выдан 01.01.2020',
      'physical.snils':'123-456-789 00','ip.ogrnip':'123456789012345','ip.long_term_resources_clause':'Нет',
    }
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':initial}).json()
    seen=[]; turns=1
    while body['state']!='TRAINING_READY':
        assert body['state']!='TRAINING_BLOCKED',body['messages'][-1]['content']
        pending=body['procurement_context']['pending_clarification']; question_id=pending['question_id']
        assert question_id not in seen,(question_id,body['messages'][-1]['content'])
        seen.append(question_id)
        path=pending['target_paths'][0]
        assert path in answers,(path,pending)
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':answers[path]}).json(); turns+=1
        assert turns<30
    assert body['procurement_context']['training_decision']['render_status']=='READY'
    assert not body['procurement_context']['missing_fields']
    download=client.get(f'/api/v1/cases/{case_id}/contract-template')
    assert download.status_code==200 and download.content[:2]==b'PK'
    return body,turns


def test_sequential_full_physical_flow_reaches_ready_and_generates_docx():
    client=TestClient(app); headers=auth(client)
    body,turns=run_training_dialogue(client,headers,'Нужен тренер, обычное физлицо без ИП и НПД, гражданин РФ, совершеннолетний, проводит лично без привлечения других лиц')
    assert body['recommendation']['template_code']=='PHYSICAL_PERSON' and turns==15


def test_sequential_full_ip_no_vat_and_vat_flows_reach_ready():
    client=TestClient(app); headers=auth(client)
    base='Нужен тренер ИП без НПД, проводит лично без привлечения других лиц'
    no_vat,no_vat_turns=run_training_dialogue(client,headers,base,'NO_VAT')
    vat,vat_turns=run_training_dialogue(client,headers,base,'VAT_PAYER')
    assert no_vat['recommendation']['template_code']=='INDIVIDUAL_ENTREPRENEUR' and no_vat_turns==14
    assert vat['recommendation']['template_code']=='INDIVIDUAL_ENTREPRENEUR' and vat_turns==15


def test_training_negative_flows_stop_before_red_fields():
    client=TestClient(app); headers=auth(client)
    for phrase in ('Нужен тренер самозанятый НПД','Нужен тренер юридическое лицо ООО Ромашка','Нужен иностранный тренер','Нужен несовершеннолетний тренер','Нужен тренер, будут ещё другие тренеры','Нужен тренер в штат на оклад'):
        case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':phrase}).json()
        assert body['state']=='TRAINING_BLOCKED',phrase
        assert body['procurement_context']['pending_clarification'] is None


def test_full_api_schedule_with_three_rows_and_address_commas_progresses():
    client=TestClient(app); headers=auth(client)
    schedule='20.10.2026, вторник, 19:00, 1, 5000 руб., Москва, ул. Большая Татарская, 11; 22.10.2026, четверг, 19:00, 1, 5000 руб., Москва, ул. Большая Татарская, 11; 27.10.2026, вторник, 19:00, 1, 5000 руб., Москва, ул. Большая Татарская, 11'
    body,_=run_training_dialogue(client,headers,'Нужен тренер, обычное физлицо без ИП и НПД, гражданин РФ, совершеннолетний, проводит лично без привлечения других лиц',schedule_answer=schedule)
    entries=body['procurement_context']['training_facts']['schedule']['entries']
    assert len(entries)==3 and [entry['date'] for entry in entries]==['20.10.2026','22.10.2026','27.10.2026']
    assert all(entry['place']=='Москва, ул. Большая Татарская, 11' for entry in entries)
    assert body['procurement_context']['pending_clarification'] is None


def test_mixed_contract_flow_returns_clarification():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нужен договор: купить оборудование и выполнить большой комплекс монтажа'}).json()
    assert body['recommendation']['status']=='AMBIGUOUS'
    assert 'основным предметом' in body['recommendation']['reason']
    assert client.get(f'/api/v1/cases/{case_id}/contract-template').status_code==404


def contract_conversation(client,headers,messages):
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'GET_CONTRACT'}).json()['id']
    states=[]
    for message in messages:
        states.append(client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':message}).json())
    return case_id,states


def test_equivalent_service_answers_resolve_goods_services_question():
    client=TestClient(app); headers=auth(client)
    for answer in ('Оказание услуги','Самостоятельное оказание услуги','Услуга'):
        _,states=contract_conversation(client,headers,['Нужен договор: передача товара или оказание услуг',answer])
        final=states[-1]
        assert final['procurement_context']['procurement_object_type']=='services'
        assert final['recommendation']['template_code']=='SERVICES'
        assert 'передача товара или самостоятельное оказание услуг' not in final['messages'][-1]['content']


def test_full_supply_specification_api_dialogue_reaches_ready_and_downloads_docx():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'CREATE_SPECIFICATION'}).json()['id']
    body=client.get(f'/api/v1/cases/{case_id}',headers=headers).json(); seen=[]
    answers={
      'equipment_category':'компьютерное оборудование','subject':'Поставка ноутбуков','item_count':'1','items[0].description':'Ноутбук Lenovo ThinkPad',
      'items[0].part_number':'21ABC','items[0].manufacturer':'Lenovo','items[0].warranty_term':'3 года',
      'items[0].unit_price':'100 тыс. руб.','items[0].quantity':'10 шт.','delivery_address':'Москва, ул. Большая Татарская, 11',
      'currency':'рубли','vat_mode':'С НДС','vat_rate':'20%','delivery_term':'10 рабочих дней',
      'support_required':'Нет','warranty_service_description':'ремонт и замена неисправных компонентов',
      'warranty_terms_url':'https://example.test/warranty',
    }
    while pending:=body['procurement_context'].get('pending_specification_question'):
        seen.append(pending['question_id'])
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':answers[pending['target']]}).json()
    assert len(seen)==len(set(seen)) and body['state']=='SPEC_REVIEW'
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Подтверждаю'}).json()
    assert body['state']=='SPEC_READY' and body['technical_specification']['status']=='READY'
    response=client.get(f'/api/v1/cases/{case_id}/technical-specification/download?format=docx',headers=headers)
    assert response.status_code==200 and response.content.startswith(b'PK')


def test_text_confirmation_uses_same_transition_as_button_after_xlsx():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'ANALYZE_SPECIFICATION'}).json()['id']
    uploaded=client.post(f'/api/v1/cases/{case_id}/attachments',headers=headers,files={'file':('sample.xlsx',spreadsheet_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json()['case']
    assert uploaded['procurement_context']['procurement_object_type']=='goods'
    assert uploaded['procurement_context']['requires_tangible_work_result'] is False
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Всё верно'}).json()
    assert body['technical_specification']['data']['kind']=='SUPPLY'
    assert body['technical_specification']['status']=='EDITING'
    assert body['procurement_context']['pending_specification_question'] is not None


def test_vague_water_procurement_clarifies_then_selects_supply():
    client=TestClient(app); headers=auth(client)
    _,states=contract_conversation(client,headers,['Нужно купить много воды','Да'])
    first,second=states
    assert first['state']=='COLLECTING_CONTRACT_DATA' and first['recommendation']['status']=='AMBIGUOUS'
    assert 'обычная поставка' in first['messages'][-1]['content'] and first['state']!='ESCALATED_TO_DPO'
    assert second['recommendation']['template_code']=='SUPPLY' and second['state']=='COMPLETED'


def test_full_technical_support_description_selects_services():
    client=TestClient(app); headers=auth(client)
    message='Нужен договор на оказание услуг по технической поддержке информационной системы в течение 12 месяцев. Исполнитель консультирует пользователей, устраняет ошибки и обеспечивает сопровождение. Передача товаров и создание нового программного продукта не предусмотрены.'
    _,states=contract_conversation(client,headers,[message])
    body=states[0]
    assert body['procurement_context']['requires_transfer_of_goods'] is False
    assert body['procurement_context']['procurement_object_type']=='services'
    assert body['procurement_context']['new_software_development'] is False
    assert body['recommendation']['template_code']=='SERVICES'


def test_context_contradiction_replaces_domain_not_marks_missing():
    client=TestClient(app); headers=auth(client)
    _,states=contract_conversation(client,headers,['Нужен договор: только поставка оборудования','Нет, всё-таки это услуга'])
    body=states[-1]
    assert body['procurement_context']['requires_transfer_of_goods'] is False
    assert body['procurement_context']['requires_service_activity'] is True
    assert body['recommendation']['template_code']=='SERVICES'


def test_invalid_short_and_test_placeholders_are_not_saved_or_returned():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'CREATE_SPECIFICATION'}).json()['id']
    client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нужно разработать лендинг'})
    for value in ('М','ТЕст'):
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':value}).json()
    assert 'М' not in body['technical_specification']['data'].values()
    assert 'ТЕст' not in body['technical_specification']['data'].values()


def test_supply_dialogue_does_not_repeat_question_after_valid_answer():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'CREATE_SPECIFICATION'}).json()['id']
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'компьютерное оборудование'}).json()
    first=body['procurement_context']['pending_specification_question']['question_id']
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Поставка компьютерного оборудования'}).json()
    second=body['procurement_context']['pending_specification_question']['question_id']
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'2'}).json()
    third=body['procurement_context']['pending_specification_question']['question_id']
    assert first=='SUPPLY_SUBJECT' and second=='SUPPLY_ITEM_COUNT' and third=='SUPPLY_ITEM_1_DESCRIPTION'
    assert len({first,second,third})==3


def complete_supply_case(client,headers,confirm=True):
    case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'CREATE_SPECIFICATION'}).json()['id']
    body=client.get(f'/api/v1/cases/{case_id}',headers=headers).json()
    answers={'equipment_category':'компьютерное оборудование','subject':'Поставка ноутбуков','item_count':'1','items[0].description':'Ноутбук Lenovo ThinkPad',
      'items[0].manufacturer':'Lenovo','items[0].warranty_term':'3 года','items[0].unit_price':'100 тыс. руб.',
      'items[0].quantity':'5 шт.','delivery_address':'Москва, ул. Большая Татарская, 11','currency':'рубли',
      'vat_mode':'С НДС','vat_rate':'22%','delivery_term':'10 рабочих дней','support_required':'Нет',
      'warranty_service_description':'ремонт и замена','warranty_terms_url':'https://example.test/warranty'}
    while pending:=body['procurement_context'].get('pending_specification_question'):
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':answers[pending['target']]}).json()
    if confirm: body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Подтверждаю'}).json()
    return case_id,body


def test_general_new_case_enters_canonical_supply_flow():
    client=TestClient(app); headers=auth(client)
    case_id=client.post('/api/v1/cases',headers=headers,json={}).json()['id']
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нужно купить ноутбуки.'}).json()
    assert body['technical_specification']['data']['kind']=='SUPPLY'
    assert body['supply_specification']['data']['kind']=='SUPPLY'
    assert body['procurement_context']['procurement_type']=='SUPPLY'
    assert body['procurement_context']['workflow_type']=='SUPPLY_SPECIFICATION'
    assert body['procurement_context']['specification_template']=='SUPPLY'
    assert body['procurement_context']['pending_specification_question']['question_id'].startswith('SUPPLY_')


def test_supply_reject_confirmation_enters_editing_without_repeating_summary():
    client=TestClient(app); headers=auth(client); case_id,body=complete_supply_case(client,headers,confirm=False)
    previous=body['messages'][-1]['content']; assert 'Всё верно?' in previous
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'нет'}).json()
    assert body['technical_specification']['status']=='EDITING'
    assert body['messages'][-1]['content']=='Хорошо. Что нужно изменить в ТЗ?'
    assert body['messages'][-1]['content']!=previous


def test_supply_reject_then_select_vat_field_asks_specific_question():
    client=TestClient(app); headers=auth(client); case_id,body=complete_supply_case(client,headers,confirm=False)
    client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'нет'})
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'НДС'}).json()
    assert body['messages'][-1]['content']=='Укажите новую ставку НДС, например 22%.'
    assert 'Проверьте ТЗ' not in body['messages'][-1]['content']
    assert body['procurement_context']['pending_specification_question']['target']=='vat_rate'
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'18%'}).json()
    assert body['technical_specification']['data']['vat_rate']==18


def test_supply_reject_confirmation_with_sparse_patch():
    client=TestClient(app); headers=auth(client); case_id,body=complete_supply_case(client,headers,confirm=False)
    subject=body['technical_specification']['data']['subject']; items=len(body['technical_specification']['data']['items'])
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Нет, срок поставки должен быть 30 дней.'}).json()
    assert body['technical_specification']['status']=='EDITING'
    assert body['technical_specification']['data']['delivery_term']=='30 дней'
    assert body['technical_specification']['data']['subject']==subject and len(body['technical_specification']['data']['items'])==items
    assert 'Всё верно?' not in body['messages'][-1]['content']


def test_supply_chat_does_not_leak_internal_enum_names():
    client=TestClient(app); headers=auth(client); _,body=complete_supply_case(client,headers,confirm=False)
    visible='\n'.join(message['content'] for message in body['messages'] if message['role']=='assistant')
    for internal in ('SUBJECT','SUPPLY','READY','EDITING','COMPLETED','NEEDS_INPUT','MANUAL_REVIEW','question_id'):
        assert internal not in visible


def test_supply_owner_can_patch_draft_and_foreign_user_cannot():
    owner=TestClient(app); owner_headers=auth(owner); case_id=owner.post('/api/v1/cases',headers=owner_headers,json={'initial_intent':'CREATE_SPECIFICATION'}).json()['id']
    response=owner.patch(f'/api/v1/cases/{case_id}/technical-specification',headers=owner_headers,json={'changes':[{'field':'delivery_term','value':'30 дней'}]})
    assert response.status_code==200 and response.json()['data']['delivery_term']=='30 дней'
    foreign=TestClient(app); foreign_headers=auth(foreign,'petrova')
    response=foreign.patch(f'/api/v1/cases/{case_id}/technical-specification',headers=foreign_headers,json={'changes':[{'field':'delivery_term','value':'1 день'}]})
    assert response.status_code==404


def test_supply_ready_reopens_then_right_panel_patch_persists():
    client=TestClient(app); headers=auth(client); case_id,_=complete_supply_case(client,headers)
    reopened=client.post(f'/api/v1/cases/{case_id}/technical-specification/reopen',headers=headers)
    assert reopened.status_code==200 and reopened.json()['status']=='EDITING'
    patched=client.patch(f'/api/v1/cases/{case_id}/technical-specification',headers=headers,json={'changes':[{'field':'items[0].quantity','value':'10'}]})
    assert patched.status_code==200 and patched.json()['data']['items'][0]['quantity']==10
    persisted=client.get(f'/api/v1/cases/{case_id}',headers=headers).json()
    assert persisted['technical_specification']['data']['items'][0]['quantity']==10


def test_supply_chat_sparse_edits_subject_lock_and_undo():
    client=TestClient(app); headers=auth(client); case_id,body=complete_supply_case(client,headers)
    subject=body['technical_specification']['data']['subject']; original_address=body['technical_specification']['data']['delivery_address']
    edits=('Поменяй количество первой позиции на 10','Исправь срок поставки на 30 дней','Добавь критерий приёмки: оборудование должно быть новым','Добавь критерий приёмки: упаковка не повреждена','Поменяй гарантию на 36 месяцев')
    for message in edits:
        body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':message}).json()
        assert body['technical_specification']['data']['subject']==subject
    data=body['technical_specification']['data']
    assert data['items'][0]['quantity']==10 and data['delivery_term']=='30 дней'
    assert 'оборудование должно быть новым' in data['acceptance_criteria'] and data['items'][0]['warranty_term']=='36 месяцев'
    assert data['delivery_address']==original_address
    body=client.post(f'/api/v1/cases/{case_id}/messages',headers=headers,json={'content':'Отмени последнее изменение'}).json()
    assert body['technical_specification']['data']['items'][0]['warranty_term']=='3 года'


def test_supply_without_part_number_ready_download_completed_and_contract_handoff():
    client=TestClient(app); headers=auth(client); case_id,body=complete_supply_case(client,headers)
    assert body['technical_specification']['status']=='READY' and body['technical_specification']['data']['items'][0].get('part_number','')==''
    assert body['recommendation']['template_code']=='SUPPLY' and body['procurement_context']['contract_status']=='READY'
    response=client.get(f'/api/v1/cases/{case_id}/technical-specification/download?format=docx',headers=headers)
    assert response.status_code==200
    document=Document(io.BytesIO(response.content)); assert document.tables[0].rows[1].cells[1].text==''
    body=client.get(f'/api/v1/cases/{case_id}',headers=headers).json()
    assert body['technical_specification']['status']=='COMPLETED' and body['state']=='COMPLETED'


def test_uploaded_tz_extracts_delivery_term():
    client=TestClient(app); headers=auth(client); case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'ANALYZE_SPECIFICATION'}).json()['id']
    source=Document(); source.add_paragraph('Предмет закупки: поставка ноутбуков'); source.add_paragraph('Срок поставки — 22 календарных дня'); payload=io.BytesIO(); source.save(payload)
    body=client.post(f'/api/v1/cases/{case_id}/attachments',headers=headers,files={'file':('existing.docx',payload.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()['case']
    assert body['technical_specification']['data']['delivery_term']=='22 календарных дня'
    assert body['supply_specification']['data']['kind']=='SUPPLY'
    assert body['procurement_context']['workflow_type']=='SUPPLY_SPECIFICATION'


def test_uploaded_supply_tz_keeps_vat_and_delivery_term_separate():
    client=TestClient(app); headers=auth(client); case_id=client.post('/api/v1/cases',headers=headers,json={'initial_intent':'ANALYZE_SPECIFICATION'}).json()['id']
    source=Document(); source.add_paragraph('Предмет закупки: поставка ноутбуков'); source.add_paragraph('НДС: 22%.'); source.add_paragraph('Срок поставки: 10 дней.'); payload=io.BytesIO(); source.save(payload)
    body=client.post(f'/api/v1/cases/{case_id}/attachments',headers=headers,files={'file':('supply.docx',payload.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()['case']
    data=body['supply_specification']['data']
    assert data['vat_rate']==22 and data['delivery_term']=='10 дней'
    assert data['delivery_term'] not in {'22','22.0'}


def test_contract_platform_rbac_ownership_versions_and_immutable_source():
    employee=TestClient(app); employee_headers=auth(employee,'petrova')
    assert employee.get('/api/v1/document-templates',headers=employee_headers).status_code==403
    assert employee.post('/api/v1/document-templates',headers=employee_headers,json={'code':'NO','name':'Нельзя'}).status_code==403

    manager=TestClient(app); manager_headers=auth(manager); template=create_platform_template(manager,manager_headers,'POC_VERSIONS')
    source=contract_docx_bytes(); digest=hashlib.sha256(source).hexdigest()
    first=manager.post(f"/api/v1/document-templates/{template['id']}/versions",headers=manager_headers,files={'file':('contract.docx',source,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    second=manager.post(f"/api/v1/document-templates/{template['id']}/versions",headers=manager_headers,files={'file':('contract-v2.docx',source,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert first.status_code==201 and second.status_code==201
    assert first.json()['version_number']==1 and second.json()['version_number']==2
    assert first.json()['source_sha256']==digest
    downloaded=manager.get(f"/api/v1/document-templates/{template['id']}/versions/{first.json()['id']}/source",headers=manager_headers)
    assert downloaded.status_code==200 and downloaded.content==source and hashlib.sha256(downloaded.content).hexdigest()==digest
    assert len(manager.get(f"/api/v1/document-templates/{template['id']}/versions",headers=manager_headers).json())==2
    configure_platform_field(manager,manager_headers,template,second.json())
    published=manager.post(f"/api/v1/document-templates/{template['id']}/versions/{second.json()['id']}/publish",headers=manager_headers)
    assert published.status_code==200 and published.json()['status']=='PUBLISHED'
    metadata=manager.get(f"/api/v1/document-templates/{template['id']}",headers=manager_headers).json()
    assert metadata['current_version_id']==second.json()['id'] and metadata['status']=='ACTIVE'

    database=SessionLocal(); petrova=database.query(User).filter_by(external_id='petrova').one(); petrova.roles=['employee','dpo_specialist']; database.commit(); database.close()
    employee_headers=auth(employee,'petrova')
    assert employee.get(f"/api/v1/document-templates/{template['id']}",headers=employee_headers).status_code==404


def test_contract_platform_rejects_invalid_uploads_and_separates_publish_permission():
    client=TestClient(app); headers=auth(client,'petrova'); template=create_platform_template(client,headers,'POC_SECURITY')
    text=client.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('contract.txt',b'plain text','text/plain')})
    fake=client.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('contract.docx',b'not a zip','application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert text.status_code==415 and fake.status_code==422
    uploaded=client.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('contract.docx',contract_docx_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
    denied=client.post(f"/api/v1/document-templates/{template['id']}/versions/{uploaded['id']}/publish",headers=headers)
    assert denied.status_code==403


def test_docx_preview_and_binding_use_stable_structural_locator():
    client=TestClient(app); headers=auth(client); template=create_platform_template(client,headers,'POC_LOCATORS')
    uploaded=client.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('contract.docx',contract_docx_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
    preview=client.get(f"/api/v1/document-templates/{template['id']}/versions/{uploaded['id']}/preview",headers=headers).json()
    assert preview['adapter']=='DOCX_XML_READ_ONLY' and preview['structure_fingerprint']==uploaded['structure_fingerprint']
    assert any(node['part'].startswith('word/header') for node in preview['nodes'])
    assert any(node['part'].startswith('word/footer') for node in preview['nodes'])
    merged=next(node for node in preview['nodes'] if node['text']=='Условия договора')
    assert merged['kind']=='table_cell' and merged['metadata']['grid_span']==2
    node=next(node for node in preview['nodes'] if 'АО «НСПК»' in node['text']); start=node['text'].index('АО «НСПК»'); end=start+len('АО «НСПК»')
    binding=client.post(f"/api/v1/document-templates/{template['id']}/versions/{uploaded['id']}/bindings",headers=headers,json={'binding_key':'contract.customer','locator':node['locator'],'start':start,'end':end,'selected_text':'АО «НСПК»'})
    assert binding.status_code==201
    locator=binding.json()['locator']; assert locator['start_run']!=locator['end_run'] and locator['structure_fingerprint']==preview['structure_fingerprint']
    invalid=client.post(f"/api/v1/document-templates/{template['id']}/versions/{uploaded['id']}/bindings",headers=headers,json={'binding_key':'contract.invalid','locator':node['locator'],'start':start,'end':end,'selected_text':'другой текст'})
    assert invalid.status_code==422


def test_contract_platform_end_to_end_multiple_occurrences_generation_and_ownership():
    document=Document()
    for prefix in ('Заказчик: ','Подпись заказчика: '):
        paragraph=document.add_paragraph(); paragraph.add_run(prefix); paragraph.add_run('АО '); paragraph.add_run('«НСПК»'); paragraph.add_run(', действует на основании устава.')
    source_buffer=io.BytesIO(); document.save(source_buffer); source=source_buffer.getvalue(); original_hash=hashlib.sha256(source).hexdigest()

    manager=TestClient(app); manager_headers=auth(manager); template=create_platform_template(manager,manager_headers,'E2E_CONTRACT')
    uploaded=manager.post(f"/api/v1/document-templates/{template['id']}/versions",headers=manager_headers,files={'file':('e2e.docx',source,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert uploaded.status_code==201,uploaded.text
    version=uploaded.json()
    field=configure_platform_field(manager,manager_headers,template,version,occurrences=2)
    bindings=manager.get(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/bindings",headers=manager_headers).json()
    assert [item['occurrence_index'] for item in bindings]==[1,2]
    assert manager.post(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/publish",headers=manager_headers).status_code==200
    assert manager.patch(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/fields/{field['id']}",headers=manager_headers,json={'field_key':'customer_name','label':'Изменить нельзя','field_type':'TEXT','required':True,'position':1}).status_code==409

    employee=TestClient(app); employee_headers=auth(employee,'petrova')
    catalog=employee.get('/api/v1/contract-catalog',headers=employee_headers).json(); assert any(item['id']==template['id'] for item in catalog)
    draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']})
    assert draft.status_code==201,draft.text
    draft_id=draft.json()['id']
    assert employee.post(f'/api/v1/contract-drafts/{draft_id}/generate',headers=employee_headers).status_code==422
    assert employee.patch(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers,json={'answers':{'unknown':'x'}}).status_code==422
    saved=employee.patch(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers,json={'answers':{'customer_name':'ООО «Ромашка»'}})
    assert saved.status_code==200
    generated=employee.post(f'/api/v1/contract-drafts/{draft_id}/generate',headers=employee_headers)
    assert generated.status_code==201,generated.text
    download=employee.get(generated.json()['download_url'],headers=employee_headers)
    assert download.status_code==200 and 'attachment' in download.headers['content-disposition'].lower()
    rendered=Document(io.BytesIO(download.content)); text='\n'.join(p.text for p in rendered.paragraphs)
    assert text.count('ООО «Ромашка»')==2 and 'действует на основании устава' in text and 'АО «НСПК»' not in text
    original=manager.get(f"/api/v1/document-templates/{template['id']}/versions/{version['id']}/source",headers=manager_headers)
    assert hashlib.sha256(original.content).hexdigest()==original_hash
    outsider=TestClient(app); outsider_headers=auth(outsider)
    assert outsider.get(f'/api/v1/contract-drafts/{draft_id}',headers=outsider_headers).status_code==404
    assert outsider.get(generated.json()['download_url'],headers=outsider_headers).status_code==404


def test_contract_platform_conditional_payment_question_end_to_end():
    document=Document(); document.add_paragraph('Порядок оплаты: PAYMENT'); document.add_paragraph('Процент предоплаты: ADVANCE')
    payload=io.BytesIO(); document.save(payload)
    manager=TestClient(app); manager_headers=auth(manager); template=create_platform_template(manager,manager_headers,'R31_PAYMENT')
    version=manager.post(f"/api/v1/document-templates/{template['id']}/versions",headers=manager_headers,files={'file':('payment.docx',payload.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
    base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"
    payment=manager.post(base+'/fields',headers=manager_headers,json={'field_key':'payment.model','label':'Какой порядок оплаты предусмотрен?','field_type':'SINGLE_SELECT','required':True,'position':1,'options':['После оказания услуг','Полная предоплата','Частичная предоплата']})
    assert payment.status_code==201,payment.text
    condition={'operator':'EQUALS','field_key':'payment.model','value':'Частичная предоплата'}
    percent=manager.post(base+'/fields',headers=manager_headers,json={'field_key':'payment.advance_percent','label':'Какой процент предоплаты предусмотрен?','field_type':'INTEGER','required':False,'position':2,'visible_when':condition,'required_when':condition,'validation':{'min':1,'max':99}})
    assert percent.status_code==201,percent.text
    preview=manager.get(base+'/preview',headers=manager_headers).json()
    for marker,field_id in [('PAYMENT',payment.json()['id']),('ADVANCE',percent.json()['id'])]:
        node=next(x for x in preview['nodes'] if marker in x['text']); start=node['text'].index(marker)
        bound=manager.post(base+'/bindings',headers=manager_headers,json={'field_id':field_id,'locator':node['locator'],'start':start,'end':start+len(marker),'selected_text':marker})
        assert bound.status_code==201,bound.text
    assert manager.post(base+'/publish',headers=manager_headers).status_code==200

    employee=TestClient(app); employee_headers=auth(employee,'petrova')
    draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); draft_id=draft['id']
    bypass=employee.patch(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers,json={'answers':{'payment.model':'Частичная предоплата'}})
    assert bypass.status_code==422 and 'payment.advance_percent' in bypass.text
    first=employee.patch(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers,json={'answers':{'payment.model':'Частичная предоплата','payment.advance_percent':30}})
    assert first.status_code==200
    hidden=employee.patch(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers,json={'answers':{'payment.model':'После оказания услуг'}})
    assert hidden.status_code==200 and hidden.json()['answers']['payment.advance_percent']==30
    generated=employee.post(f'/api/v1/contract-drafts/{draft_id}/generate',headers=employee_headers)
    assert generated.status_code==201,generated.text
    rendered=Document(io.BytesIO(employee.get(generated.json()['download_url'],headers=employee_headers).content))
    text='\n'.join(x.text for x in rendered.paragraphs)
    assert 'После оказания услуг' in text and 'ADVANCE' in text and '30' not in text
    restored=employee.get(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers).json()
    assert restored['answers']['payment.advance_percent']==30
    again=employee.patch(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers,json={'answers':{'payment.model':'Частичная предоплата'}})
    assert again.status_code==200 and again.json()['answers']['payment.advance_percent']==30


def test_contract_platform_rejects_invalid_condition_configuration():
    client=TestClient(app); headers=auth(client); template=create_platform_template(client,headers,'R31_INVALID')
    uploaded=client.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('conditions.docx',contract_docx_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json()
    base=f"/api/v1/document-templates/{template['id']}/versions/{uploaded['id']}"
    a=client.post(base+'/fields',headers=headers,json={'field_key':'field.a','label':'Поле А','field_type':'TEXT','position':1}).json()
    missing=client.post(base+'/fields',headers=headers,json={'field_key':'field.bad','label':'Ошибка ссылки','field_type':'TEXT','position':2,'visible_when':{'operator':'EQUALS','field_key':'missing','value':'x'}})
    assert missing.status_code==422
    b=client.post(base+'/fields',headers=headers,json={'field_key':'field.b','label':'Поле Б','field_type':'TEXT','position':2,'visible_when':{'operator':'EQUALS','field_key':'field.a','value':'x'}})
    assert b.status_code==201
    cycle=client.patch(base+f"/fields/{a['id']}",headers=headers,json={'field_key':'field.a','label':'Поле А','field_type':'TEXT','position':1,'visible_when':{'operator':'EQUALS','field_key':'field.b','value':'x'}})
    assert cycle.status_code==422 and 'циклическая' in cycle.text
    wrong=client.patch(base+f"/fields/{b.json()['id']}",headers=headers,json={'field_key':'field.b','label':'Поле Б','field_type':'TEXT','position':2,'visible_when':{'operator':'GREATER_THAN','field_key':'field.a','value':'x'}})
    assert wrong.status_code==422 and 'несовместим' in wrong.text
    assert client.delete(base+f"/fields/{a['id']}",headers=headers).status_code==422
    database=SessionLocal(); broken=database.query(QuestionnaireField).filter_by(id=b.json()['id']).one(); broken.config={**broken.config,'visible_when':{'operator':'EQUALS','field_key':'deleted.field','value':'x'}}; database.commit(); database.close()
    rejected=client.post(base+'/publish',headers=headers)
    assert rejected.status_code==422 and 'не найдено' in rejected.text


def test_r32_conditional_legal_payment_alternatives_generate_only_selected_branch():
    source_doc=Document(); source_doc.add_paragraph('Вариант А: оплата после оказания услуг.'); source_doc.add_paragraph('Вариант Б: полная предоплата.')
    source_buffer=io.BytesIO(); source_doc.save(source_buffer); source=source_buffer.getvalue(); digest=hashlib.sha256(source).hexdigest()
    manager=TestClient(app); manager_headers=auth(manager); template=create_platform_template(manager,manager_headers,'R32_PAYMENT')
    version=manager.post(f"/api/v1/document-templates/{template['id']}/versions",headers=manager_headers,files={'file':('legal-payment.docx',source,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json(); base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"
    field=manager.post(base+'/fields',headers=manager_headers,json={'field_key':'payment.model','label':'Порядок оплаты','field_type':'SINGLE_SELECT','position':1,'options':['После оказания услуг','Предоплата']})
    assert field.status_code==201,field.text
    nodes=manager.get(base+'/preview',headers=manager_headers).json()['nodes']; after=next(x for x in nodes if 'Вариант А' in x['text']); advance=next(x for x in nodes if 'Вариант Б' in x['text'])
    rules=[('after_services',after,{'operator':'AND','conditions':[{'operator':'EQUALS','field_key':'payment.model','value':'После оказания услуг'},{'operator':'NOT','condition':{'operator':'EQUALS','field_key':'payment.model','value':'Предоплата'}}]}),('prepayment',advance,{'operator':'EQUALS','field_key':'payment.model','value':'Предоплата'})]
    for key,node,condition in rules:
        created=manager.post(base+'/conditional-regions',headers=manager_headers,json={'region_key':key,'region_type':'CONDITIONAL_PARAGRAPH','start_locator':node['locator'],'include_when':condition})
        assert created.status_code==201,created.text
    assert manager.post(base+'/publish',headers=manager_headers).status_code==200
    employee=TestClient(app); employee_headers=auth(employee,'petrova')
    for choice,present,absent in [('После оказания услуг','Вариант А','Вариант Б'),('Предоплата','Вариант Б','Вариант А')]:
        draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); draft_id=draft['id']
        assert employee.patch(f'/api/v1/contract-drafts/{draft_id}',headers=employee_headers,json={'answers':{'payment.model':choice}}).status_code==200
        generated=employee.post(f'/api/v1/contract-drafts/{draft_id}/generate',headers=employee_headers); assert generated.status_code==201,generated.text
        rendered=Document(io.BytesIO(employee.get(generated.json()['download_url'],headers=employee_headers).content)); text='\n'.join(x.text for x in rendered.paragraphs)
        assert present in text and absent not in text
    original=manager.get(base+'/source',headers=manager_headers)
    assert hashlib.sha256(original.content).hexdigest()==digest


def test_r32_supports_all_region_types_and_rejects_overlap_and_stale_locator():
    document=Document(); document.add_paragraph('Префикс: условный фрагмент.'); document.add_paragraph('Отдельный пункт.'); document.add_paragraph('Начало группы.'); document.add_paragraph('Конец группы.')
    first=document.add_table(rows=2,cols=2); first.cell(0,0).merge(first.cell(0,1)).text='Объединённый заголовок'; first.cell(1,0).text='Удаляемая'; first.cell(1,1).text='строка'
    second=document.add_table(rows=1,cols=1); second.cell(0,0).text='Удаляемая таблица'
    payload=io.BytesIO(); document.save(payload)
    client=TestClient(app); headers=auth(client); template=create_platform_template(client,headers,'R32_TYPES')
    version=client.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('regions.docx',payload.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json(); base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"
    assert client.post(base+'/fields',headers=headers,json={'field_key':'enabled','label':'Включить условия','field_type':'BOOLEAN','position':1}).status_code==201
    nodes=client.get(base+'/preview',headers=headers).json()['nodes']; condition={'operator':'EQUALS','field_key':'enabled','value':True}
    run=next(x for x in nodes if 'условный фрагмент' in x['text']); start=run['text'].index('условный фрагмент')
    requests=[{'region_key':'run','region_type':'CONDITIONAL_RUN','start_locator':run['locator'],'start':start,'end':start+len('условный фрагмент'),'selected_text':'условный фрагмент'}, {'region_key':'paragraph','region_type':'CONDITIONAL_PARAGRAPH','start_locator':next(x for x in nodes if x['text']=='Отдельный пункт.')['locator']}, {'region_key':'group','region_type':'CONDITIONAL_PARAGRAPH_GROUP','start_locator':next(x for x in nodes if x['text']=='Начало группы.')['locator'],'end_locator':next(x for x in nodes if x['text']=='Конец группы.')['locator']}, {'region_key':'row','region_type':'CONDITIONAL_TABLE_ROW','start_locator':next(x for x in nodes if x['text']=='Удаляемая')['locator']}, {'region_key':'table','region_type':'CONDITIONAL_TABLE','start_locator':next(x for x in nodes if x['text']=='Удаляемая таблица')['locator']}]
    for body in requests:
        response=client.post(base+'/conditional-regions',headers=headers,json={**body,'include_when':condition}); assert response.status_code==201,response.text
        assert response.json()['locator']['structure_fingerprint']==version['structure_fingerprint']
    overlap=client.post(base+'/conditional-regions',headers=headers,json={'region_key':'overlap','region_type':'CONDITIONAL_PARAGRAPH','start_locator':run['locator'],'include_when':condition})
    assert overlap.status_code==422 and 'пересекаются' in overlap.text
    listed=client.get(base+'/conditional-regions',headers=headers).json(); assert {x['region_type'] for x in listed}=={'CONDITIONAL_RUN','CONDITIONAL_PARAGRAPH','CONDITIONAL_PARAGRAPH_GROUP','CONDITIONAL_TABLE_ROW','CONDITIONAL_TABLE'}
    database=SessionLocal(); from app.db import ConditionalDocumentRegion
    stale=database.query(ConditionalDocumentRegion).filter_by(region_key='table',template_version_id=version['id']).one(); stale.locator={**stale.locator,'structure_fingerprint':'0'*64}; database.commit(); database.close()
    rejected=client.post(base+'/publish',headers=headers); assert rejected.status_code==422 and 'устарел' in rejected.text


def test_r33_full_repeating_table_flow_preserves_order_static_rows_and_source():
    manager=TestClient(app); manager_headers=auth(manager); template,version,source,_,_=configure_repeating_platform(manager,manager_headers,'R33_FULL')
    base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"; assert manager.post(base+'/publish',headers=manager_headers).status_code==200
    assert hashlib.sha256(manager.get(base+'/source',headers=manager_headers).content).hexdigest()==hashlib.sha256(source).hexdigest()
    employee=TestClient(app); employee_headers=auth(employee,'petrova'); draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); did=draft['id']
    rows=[{'_id':f'row-{i}','date':f'2026-11-{i:02d}','place':f'Зал {i}'} for i in range(1,6)]; rows[2]['place']='Изменённый зал'; rows.pop(1); rows=[rows[2],rows[0],rows[1],rows[3]]
    saved=employee.patch(f'/api/v1/contract-drafts/{did}/autosave',headers=employee_headers,json={'answers':{'customer_name':'АО Заказчик','schedule':rows}}); assert saved.status_code==200,saved.text
    restored=employee.get(f'/api/v1/contract-drafts/{did}',headers=employee_headers).json(); assert restored['answers']['schedule']==rows
    assert employee.patch(f'/api/v1/contract-drafts/{did}',headers=employee_headers,json={'answers':restored['answers']}).status_code==200
    generated=employee.post(f'/api/v1/contract-drafts/{did}/generate',headers=employee_headers); assert generated.status_code==201,generated.text
    rendered=Document(io.BytesIO(employee.get(generated.json()['download_url'],headers=employee_headers).content)); table=rendered.tables[0]
    assert len(table.rows)==6 and table.rows[0].cells[0].text=='Дата' and table.rows[-1].cells[0].text=='ИТОГО'
    assert [row.cells[1].text for row in table.rows[1:-1]]==[x['place'] for x in rows]
    assert all('Дополнительная строка' in row.cells[0].text and len(row.cells[0].paragraphs)==2 for row in table.rows[1:-1])


def test_r33_zero_one_three_and_ten_rows_and_limits():
    manager=TestClient(app); headers=auth(manager); template,version,_,_,_=configure_repeating_platform(manager,headers,'R33_COUNTS',min_items=0,max_items=10); base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"; assert manager.post(base+'/publish',headers=headers).status_code==200
    employee=TestClient(app); employee_headers=auth(employee,'petrova')
    for count in (0,1,3,10):
        draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); rows=[{'_id':f'{count}-{i}','date':'2026-12-01','place':f'Зал {i}'} for i in range(count)]; answers={'customer_name':'Заказчик','schedule':rows}
        assert employee.patch(f"/api/v1/contract-drafts/{draft['id']}",headers=employee_headers,json={'answers':answers}).status_code==200
        made=employee.post(f"/api/v1/contract-drafts/{draft['id']}/generate",headers=employee_headers); assert made.status_code==201,made.text
        result=Document(io.BytesIO(employee.get(made.json()['download_url'],headers=employee_headers).content)); assert len(result.tables[0].rows)==count+2
    overflow=[{'_id':f'x-{i}','date':'2026-12-01','place':'Зал'} for i in range(11)]; assert employee.patch(f"/api/v1/contract-drafts/{draft['id']}",headers=employee_headers,json={'answers':{'schedule':overflow}}).status_code==422


def test_r33_minimum_and_required_nested_values_are_enforced_server_side():
    manager=TestClient(app); headers=auth(manager); template,version,_,_,_=configure_repeating_platform(manager,headers,'R33_REQUIRED',min_items=2,max_items=4); base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"; assert manager.post(base+'/publish',headers=headers).status_code==200
    employee=TestClient(app); employee_headers=auth(employee,'petrova'); draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); did=draft['id']
    one=employee.patch(f'/api/v1/contract-drafts/{did}',headers=employee_headers,json={'answers':{'customer_name':'Заказчик','schedule':[{'_id':'one','date':'2026-12-01','place':'Зал'}]}}); assert one.status_code==422 and 'минимум 2' in one.text
    missing=employee.patch(f'/api/v1/contract-drafts/{did}',headers=employee_headers,json={'answers':{'customer_name':'Заказчик','schedule':[{'_id':'one','date':'2026-12-01','place':''},{'_id':'two','date':'2026-12-02','place':'Зал'}]}}); assert missing.status_code==422 and 'заполните поле' in missing.text.lower() and 'Место' in missing.text


def test_r33_horizontal_gridspan_and_row_formatting_survive_cloning():
    document=Document(); document.add_paragraph('Служебное поле')
    table=document.add_table(rows=3,cols=2); table.style='Table Grid'; table.rows[0].cells[0].text='Заголовок'; table.rows[0].cells[1].text='Колонка'; merged=table.rows[1].cells[0].merge(table.rows[1].cells[1]); merged.text='VALUE'; merged.paragraphs[0].alignment=1; table.rows[2].cells[0].merge(table.rows[2].cells[1]).text='ИТОГО'
    raw=io.BytesIO(); document.save(raw); manager=TestClient(app); headers=auth(manager); template=create_platform_template(manager,headers,'R33_GRIDSPAN'); version=manager.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('gridspan.docx',raw.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json(); base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"
    assert manager.post(base+'/fields',headers=headers,json={'field_key':'note','label':'Примечание','field_type':'TEXT','required':False,'position':1}).status_code==201
    node=next(x for x in manager.get(base+'/preview',headers=headers).json()['nodes'] if x['text']=='VALUE'); group=manager.post(base+'/repeating-groups',headers=headers,json={'group_key':'lines','name':'Строки','min_items':0,'max_items':3,'position':1,'row_locator':node['locator']}).json(); field=manager.post(base+f"/repeating-groups/{group['id']}/fields",headers=headers,json={'field_key':'value','label':'Значение','field_type':'TEXT','required':True,'position':1}).json(); assert manager.put(base+f"/repeating-groups/{group['id']}/fields/{field['id']}/binding",headers=headers,json={'locator':node['locator'],'start':0,'end':5,'selected_text':'VALUE'}).status_code==200; assert manager.post(base+'/publish',headers=headers).status_code==200
    employee=TestClient(app); employee_headers=auth(employee,'petrova'); draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); rows=[{'_id':str(i),'value':f'Строка {i}'} for i in range(3)]; assert employee.patch(f"/api/v1/contract-drafts/{draft['id']}",headers=employee_headers,json={'answers':{'lines':rows}}).status_code==200; made=employee.post(f"/api/v1/contract-drafts/{draft['id']}/generate",headers=employee_headers).json(); rendered=Document(io.BytesIO(employee.get(made['download_url'],headers=employee_headers).content)); result=rendered.tables[0]
    assert len(result.rows)==5 and all(row._tr.find('.//'+qn('w:gridSpan')).get(qn('w:val'))=='2' for row in result.rows[1:4]); assert all(row.cells[0].paragraphs[0].alignment==1 for row in result.rows[1:4]); assert result.rows[-1].cells[0].text=='ИТОГО'


def test_r33_rejects_vertical_merge_stale_locator_conflict_and_foreign_draft():
    manager=TestClient(app); headers=auth(manager); _,_,_,response,_=configure_repeating_platform(manager,headers,'R33_VMERGE',vertical_merge=True); assert response.status_code==422 and 'вертикально' in response.text
    template,version,_,_,group=configure_repeating_platform(manager,headers,'R33_SAFETY'); base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"
    database=SessionLocal(); from app.db import RepeatingGroup
    stored=database.query(RepeatingGroup).filter_by(id=group['id']).one(); original=stored.locator; stored.locator={**original,'row_hash':'0'*64}; database.commit(); database.close(); stale=manager.post(base+'/publish',headers=headers); assert stale.status_code==422 and 'устарела' in stale.text
    database=SessionLocal(); stored=database.query(RepeatingGroup).filter_by(id=group['id']).one(); stored.locator=original; database.commit(); database.close()
    preview=manager.get(base+'/preview',headers=headers).json(); date_node=next(x for x in preview['nodes'] if 'DATE' in x['text']); conflict=manager.post(base+'/bindings',headers=headers,json={'binding_key':'conflict','locator':date_node['locator'],'start':0,'end':4,'selected_text':'DATE'}); assert conflict.status_code==201
    assert manager.post(base+'/publish',headers=headers).status_code==422
    manager.delete(base+f"/bindings/{conflict.json()['id']}",headers=headers)
    conditional=manager.post(base+'/conditional-regions',headers=headers,json={'region_key':'same_row','region_type':'CONDITIONAL_TABLE_ROW','start_locator':date_node['locator'],'include_when':{'operator':'IS_NOT_EMPTY','field_key':'customer_name'}}); assert conditional.status_code==201,conditional.text
    assert manager.post(base+'/publish',headers=headers).status_code==422
    manager.delete(base+f"/conditional-regions/{conditional.json()['id']}",headers=headers); assert manager.post(base+'/publish',headers=headers).status_code==200
    owner=TestClient(app); owner_headers=auth(owner,'petrova'); draft=owner.post('/api/v1/contract-drafts',headers=owner_headers,json={'template_id':template['id']}).json(); foreign=TestClient(app); foreign_headers=auth(foreign)
    assert foreign.patch(f"/api/v1/contract-drafts/{draft['id']}/autosave",headers=foreign_headers,json={'answers':{}}).status_code==404


def test_r33_inactive_group_keeps_answers_but_does_not_render_or_require_rows():
    manager=TestClient(app); headers=auth(manager); template,version,_,_,group=configure_repeating_platform(manager,headers,'R33_CONDITION',visible_when={'operator':'EQUALS','field_key':'customer_name','value':'Показывать'})
    # The referenced text field is condition-compatible; publishing also proves shared ConditionEngine validation.
    base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"; assert manager.post(base+'/publish',headers=headers).status_code==200
    employee=TestClient(app); employee_headers=auth(employee,'petrova'); draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); hidden=[{'_id':'kept','date':'','place':''}]
    saved=employee.patch(f"/api/v1/contract-drafts/{draft['id']}",headers=employee_headers,json={'answers':{'customer_name':'Скрыть','schedule':hidden}}); assert saved.status_code==200,saved.text
    made=employee.post(f"/api/v1/contract-drafts/{draft['id']}/generate",headers=employee_headers); assert made.status_code==201,made.text
    result=Document(io.BytesIO(employee.get(made.json()['download_url'],headers=employee_headers).content)); assert len(result.tables[0].rows)==2
    restored=employee.get(f"/api/v1/contract-drafts/{draft['id']}",headers=employee_headers).json()
    assert restored['answers']['schedule']==hidden


def test_r33_real_training_contract_empty_schedule_cells_are_bindable():
    source=next((template_directory/'training').glob('*физлицо.docx')).read_bytes(); digest=hashlib.sha256(source).hexdigest()
    manager=TestClient(app); headers=auth(manager); template=create_platform_template(manager,headers,'R33_REAL_TRAINING')
    version=manager.post(f"/api/v1/document-templates/{template['id']}/versions",headers=headers,files={'file':('training.docx',source,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}).json(); base=f"/api/v1/document-templates/{template['id']}/versions/{version['id']}"
    assert manager.post(base+'/fields',headers=headers,json={'field_key':'comment','label':'Комментарий','field_type':'TEXT','required':False,'position':1}).status_code==201
    nodes=manager.get(base+'/preview',headers=headers).json()['nodes']; cells=[x for x in nodes if x['kind']=='table_cell' and x['metadata']['table']==2 and x['metadata']['row']==3 and x['metadata']['cell']<=7 and x['locator'].endswith('/p[1]')]
    assert len(cells)==7 and all(x['text']=='' for x in cells)
    created=manager.post(base+'/repeating-groups',headers=headers,json={'group_key':'schedule','name':'График тренировок','min_items':0,'max_items':10,'required':False,'position':1,'row_locator':cells[0]['locator']}); assert created.status_code==201,created.text; group=created.json()
    for position,(cell,key,label,kind) in enumerate(zip(cells,['date','weekday','time','quantity','price','total','place'],['Дата','День недели','Время','Количество','Цена','Стоимость','Место'],['DATE','TEXT','TEXT','INTEGER','MONEY','MONEY','TEXT']),1):
        field=manager.post(base+f"/repeating-groups/{group['id']}/fields",headers=headers,json={'field_key':key,'label':label,'field_type':kind,'required':False,'position':position}).json()
        bound=manager.put(base+f"/repeating-groups/{group['id']}/fields/{field['id']}/binding",headers=headers,json={'locator':cell['locator'],'start':0,'end':0,'selected_text':''}); assert bound.status_code==200,bound.text
    assert manager.post(base+'/publish',headers=headers).status_code==200
    employee=TestClient(app); employee_headers=auth(employee,'petrova'); draft=employee.post('/api/v1/contract-drafts',headers=employee_headers,json={'template_id':template['id']}).json(); rows=[{'_id':f'real-{i}','date':f'2026-11-{i:02d}','weekday':'Понедельник','time':'19:00','quantity':1,'price':'5000','total':'5000','place':f'Зал {i}'} for i in range(1,5)]
    assert employee.patch(f"/api/v1/contract-drafts/{draft['id']}",headers=employee_headers,json={'answers':{'schedule':rows}}).status_code==200
    made=employee.post(f"/api/v1/contract-drafts/{draft['id']}/generate",headers=employee_headers); assert made.status_code==201,made.text
    rendered=Document(io.BytesIO(employee.get(made.json()['download_url'],headers=employee_headers).content)); schedule=rendered.tables[1]
    assert len(schedule.rows)==18 and [schedule.rows[i].cells[6].text for i in range(2,6)]==[x['place'] for x in rows]
    assert any('Итого за общий период' in row.cells[0].text for row in schedule.rows)
    assert hashlib.sha256(manager.get(base+'/source',headers=headers).content).hexdigest()==digest
