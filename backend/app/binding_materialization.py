"""Fail-closed materialization of reviewed structural candidates."""
from __future__ import annotations

import copy
import hashlib
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from docx import Document

from app.contract_platform import adapter, safe_source
from app.calculations import CalculationError, TextExpressionEngine
from app.db import (
    ConditionalDocumentRegion, DocumentBinding, QuestionnaireField,
    QuestionnaireSchema, RepeatingGroup, TemplateAuditEvent, TemplateVersion, settings,
)


class MaterializationError(ValueError):
    pass


DIRECT_BINDING_OPERATIONS = {"REPLACE", "DERIVED"}


def _overlaps(a_start: int, a_end: int, b_start: int, b_end: int) -> bool:
    return max(a_start, b_start) < min(a_end, b_end)


def candidate_is_resolved(session, version, candidate, candidates=None):
    if candidate.get("status") == "REGION_BINDING_COVERAGE":
        return verified_list_region_coverage(session, version, candidate)
    if candidate.get("status") == "BINDING_COVERAGE":
        return verified_binding_coverage(session, version, candidate)
    if candidate.get("status") in {"NOT_APPLICABLE", "SOURCE_PRESERVED"}:
        if not verified_training_exception(version, candidate):
            return False
        if candidate["status"] == "SOURCE_PRESERVED":
            required = set(candidate["resolution_proof"]["locators"])
            for region in session.query(ConditionalDocumentRegion).filter_by(template_version_id=version.id, status="ACTIVE"):
                locator = region.locator or {}
                if required.intersection(locator.get("nodes") or [locator.get("node")]):
                    return False
            for binding in session.query(DocumentBinding).filter_by(template_version_id=version.id, status="ACTIVE"):
                if (binding.locator or {}).get("node") in required:
                    return False
        return True
    if candidate.get("status") in {
        "CONFIGURED",
        "GENERATED_AND_VERIFIED",
    }:
        return True

    if candidate.get("status") != "STRUCTURAL_ALIAS":
        return False

    anchor_id = candidate.get("alias_of_candidate_id")
    if not anchor_id or anchor_id == candidate.get("candidate_id"):
        return False

    if candidates is None:
        candidates = (version.preview_metadata or {}).get(
            "binding_candidates", []
        )

    anchor = next(
        (
            item for item in candidates
            if item.get("candidate_id") == anchor_id
        ),
        None,
    )

    if anchor is None or anchor.get("status") not in {
        "CONFIGURED",
        "GENERATED_AND_VERIFIED",
    }:
        return False

    if candidate.get("rf_id") != anchor.get("rf_id"):
        return False

    if candidate.get("field_path") != anchor.get("field_path"):
        return False

    if candidate.get("structure_fingerprint") != anchor.get(
        "structure_fingerprint"
    ):
        return False

    if candidate.get("materialized_id"):
        return False

    binding_id = anchor.get("materialized_id")
    if not binding_id:
        return False

    if anchor.get("materialized_type") == "CONDITIONAL_DOCUMENT_REGION":
        region = session.get(ConditionalDocumentRegion, binding_id)
        row = re.fullmatch(r"word/document\.xml#/tbl\[(\d+)\]/tr\[(\d+)\]/tc\[\d+\]/p\[\d+\]", candidate.get("locator") or "")
        loc = region.locator or {} if region else {}
        return bool(
            region and row and region.status == "ACTIVE"
            and region.template_version_id == version.id
            and region.region_type == "CONDITIONAL_TABLE_ROW"
            and (int(row.group(1)), int(row.group(2))) == (loc.get("table"), loc.get("row"))
            and region.include_when == {"operator": "EQUALS", "field_key": anchor.get("field_path"), "value": True}
        )
    if anchor.get("materialized_type") != "DOCUMENT_BINDING":
        return False
    binding = session.get(DocumentBinding, binding_id)
    return bool(binding and binding.status == "ACTIVE" and binding.template_version_id == version.id and binding.binding_key == anchor.get("field_path"))


def verified_binding_coverage(session, version, candidate):
    """A multi-placeholder RF is covered only while every exact binding exists."""
    if version.status != "DRAFT" or candidate.get("rf_id") not in {"RF004", "RF017"}:
        return False
    proof = candidate.get("coverage_proof") or {}
    if proof.get("source_sha256") != version.source_sha256 or proof.get("structure_fingerprint") != version.structure_fingerprint:
        return False
    if hashlib.sha256(safe_source(version).read_bytes()).hexdigest() != version.source_sha256:
        return False
    expected_node = {
        ("RF004", "INDIVIDUAL_ENTREPRENEUR"): "word/document.xml#/p[47]",
        ("RF004", "PHYSICAL_PERSON"): "word/document.xml#/p[47]",
        ("RF017", "INDIVIDUAL_ENTREPRENEUR"): "word/document.xml#/p[294]",
        ("RF017", "PHYSICAL_PERSON"): "word/document.xml#/p[286]",
    }.get((candidate.get("rf_id"), candidate.get("variant")))
    required_keys = {"party.performer_name", "ip.ogrnip"} if candidate.get("variant") == "INDIVIDUAL_ENTREPRENEUR" else {"party.performer_name"}
    entries = proof.get("bindings") or []
    if not expected_node or len(entries) != len(required_keys) or {entry.get("binding_key") for entry in entries} != required_keys:
        return False
    schema = session.query(QuestionnaireSchema).filter_by(template_version_id=version.id).first()
    if not schema:
        return False
    fields = {field.field_key: field for field in session.query(QuestionnaireField).filter_by(schema_id=schema.id)}
    preview = adapter.preview(safe_source(version))
    if preview["structure_fingerprint"] != version.structure_fingerprint:
        return False
    nodes = {node["locator"]: node["text"] for node in preview["nodes"]}
    for entry in entries:
        binding = session.get(DocumentBinding, entry.get("id"))
        field = fields.get(entry.get("binding_key"))
        if not binding or not field or not field.required or binding.field_id != field.id or binding.status != "ACTIVE" or binding.template_version_id != version.id:
            return False
        locator = binding.locator or {}
        start, end = locator.get("text_start"), locator.get("text_end")
        text = nodes.get(expected_node)
        if locator.get("node") != expected_node or locator.get("structure_fingerprint") != version.structure_fingerprint or not isinstance(start, int) or not isinstance(end, int) or text is None or text[start:end] != binding.selected_text:
            return False
        if binding.binding_key != entry["binding_key"] or binding.selected_text != entry.get("selected_text"):
            return False
    return True


def verified_list_region_coverage(session, version, candidate):
    if version.status != "DRAFT" or candidate.get("rf_id") != "RF002" or candidate.get("operation") != "REPLACE_OR_DELETE":
        return False
    proof = candidate.get("coverage_proof") or {}
    if proof.get("source_sha256") != version.source_sha256 or proof.get("structure_fingerprint") != version.structure_fingerprint:
        return False
    if hashlib.sha256(safe_source(version).read_bytes()).hexdigest() != version.source_sha256:
        return False
    region = session.get(ConditionalDocumentRegion, proof.get("region_id"))
    binding = session.get(DocumentBinding, proof.get("binding_id"))
    schema = session.query(QuestionnaireSchema).filter_by(template_version_id=version.id).first()
    field = session.query(QuestionnaireField).filter_by(schema_id=schema.id, field_key="appendices.other").first() if schema else None
    node = "word/document.xml#/p[17]"
    if not region or not binding or not field or field.field_type != "LIST":
        return False
    if region.status != "ACTIVE" or region.template_version_id != version.id or region.region_type != "CONDITIONAL_PARAGRAPH" or (region.locator or {}).get("node") != node:
        return False
    if region.include_when != {"operator": "IS_NOT_EMPTY", "field_key": "appendices.other"}:
        return False
    if binding.status != "ACTIVE" or binding.template_version_id != version.id or binding.binding_key != "appendices.other" or binding.field_id != field.id or not (binding.format_config or {}).get("render_list"):
        return False
    preview = adapter.preview(safe_source(version))
    nodes = {n["locator"]: n["text"] for n in preview["nodes"]}
    loc = binding.locator or {}
    start, end = loc.get("text_start"), loc.get("text_end")
    if preview["structure_fingerprint"] != version.structure_fingerprint or loc.get("node") != node or not isinstance(start, int) or not isinstance(end, int) or nodes.get(node, "")[start:end] != " (указываются при наличии)":
        return False
    numbering = (field.config or {}).get("numbering") or {}
    return numbering == {"start": 3, "item_prefix": "Приложение № ", "name_separator": " — ", "prefix": ": "} and (field.config or {}).get("separator") == "; "


def verified_training_exception(version, candidate):
    """Re-check approved source and matrix; never treat an absent clause as DELETE."""
    if version.status != "DRAFT":
        return False
    proof = candidate.get("resolution_proof") or {}
    if proof.get("source_sha256") != version.source_sha256 or proof.get("structure_fingerprint") != version.structure_fingerprint:
        return False
    source = safe_source(version)
    if hashlib.sha256(source.read_bytes()).hexdigest() != version.source_sha256:
        return False
    matrix = Path(settings.contract_templates_path) / "training" / "Матрица_применимости_LLM_v2.xlsx"
    if not matrix.is_file() or hashlib.sha256(matrix.read_bytes()).hexdigest() != proof.get("matrix_sha256"):
        return False
    workbook = load_workbook(matrix, read_only=True, data_only=True)
    try:
        row = next((r for r in workbook["RED_FIELD_SCHEMA"].iter_rows(min_row=2, values_only=True) if r[0] == candidate.get("rf_id")), None)
        occurrence = next((r for r in workbook["RED_OCCURRENCES"].iter_rows(min_row=2, values_only=True) if r[0] == candidate.get("occurrence_id")), None)
    finally:
        workbook.close()
    if not row or not occurrence or occurrence[3] != candidate.get("rf_id") or candidate.get("structure_fingerprint") != version.structure_fingerprint:
        return False
    nodes = {n["locator"]: n["text"] for n in adapter.preview(source)["nodes"]}
    if candidate.get("status") == "NOT_APPLICABLE":
        if candidate.get("rf_id") not in {"RF011", "RF012"} or candidate.get("operation") != "DELETE" or row[3] != "NONE" or row[4] != "always false for allowed templates" or row[6] != "NOT_APPLICABLE":
            return False
        # Matrix clause must genuinely be absent, not merely at a different locator.
        marked = str(occurrence[5] or "").strip("[]")
        return bool(marked and not any(marked in text for text in nodes.values()) and proof.get("decision") == "PILOT_NOT_APPLICABLE")
    if candidate.get("status") == "SOURCE_PRESERVED":
        if candidate.get("rf_id") != "RF013" or candidate.get("variant") != "INDIVIDUAL_ENTREPRENEUR" or candidate.get("operation") != "INCLUDE" or row[3] != "INDIVIDUAL_ENTREPRENEUR" or row[5] != "YES":
            return False
        if proof.get("locators") != [f"word/document.xml#/p[{i}]" for i in range(119, 128)]:
            return False
        if not all(nodes.get(loc) and hashlib.sha256(nodes[loc].encode()).hexdigest() == proof.get("node_sha256", {}).get(loc) for loc in proof["locators"]):
            return False
        diagnostics = proof.get("diagnostic_sha256") or {}
        expected = {"ip-guarantee-0.docx", "ip-guarantee-1.docx"}
        if set(diagnostics) != expected:
            return False
        for name, digest in diagnostics.items():
            path = Path(settings.storage) / "training-combined-diagnostics" / name
            if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                return False
            paragraphs = {p.text for p in Document(path).paragraphs}
            if any(nodes[loc] not in paragraphs for loc in proof["locators"]):
                return False
        return True
    return False


class BindingMaterializationService:
    def __init__(self, session, actor):
        self.session = session
        self.actor = actor

    def _candidate(self, version: TemplateVersion, candidate_id: str):
        metadata = copy.deepcopy(version.preview_metadata or {})
        candidates = metadata.get("binding_candidates", [])
        candidate = next((item for item in candidates if item.get("candidate_id") == candidate_id), None)
        if candidate is None:
            raise MaterializationError("Кандидат разметки не найден")
        return metadata, candidates, candidate

    def _validate(self, version: TemplateVersion, candidate: dict, manual: bool):
        if version.status != "DRAFT":
            raise MaterializationError("Материализация разрешена только в черновой версии")
        if candidate.get("operation") not in DIRECT_BINDING_OPERATIONS:
            raise MaterializationError("Для этой операции требуется conditional/repeating configuration")
        if candidate.get("operation") == "DERIVED" and not manual:
            raise MaterializationError("Вычисляемый текст требует ручной проверки диапазона и выражения")
        if manual:
            if candidate.get("review_decision") != "CONFIRMED":
                raise MaterializationError("Кандидат должен быть явно подтверждён специалистом")
        elif candidate.get("status") != "CANDIDATE_FOUND" or candidate.get("review_decision") == "REJECTED":
            raise MaterializationError("Автоматически материализуются только точные кандидаты")
        preview = adapter.preview(safe_source(version))
        if preview["structure_fingerprint"] != version.structure_fingerprint:
            raise MaterializationError("Структурный locator устарел")
        if candidate.get("structure_fingerprint") != version.structure_fingerprint:
            raise MaterializationError("Fingerprint кандидата не соответствует версии")
        try:
            selection = adapter.select_range(
                safe_source(version), candidate["locator"], int(candidate["start"]),
                int(candidate["end"]), candidate.get("selected_text") or "",
            )
        except (KeyError, TypeError, ValueError) as exc:
            raise MaterializationError(f"Диапазон кандидата не подтверждён: {exc}") from exc
        schema = self.session.query(QuestionnaireSchema).filter_by(template_version_id=version.id).first()
        field = self.session.query(QuestionnaireField).filter_by(schema_id=schema.id, field_key=candidate["field_path"]).first() if schema else None
        if field is None:
            raise MaterializationError("Поле анкеты для кандидата не настроено")
        if candidate.get("operation") == "DERIVED":
            config = field.config or {}
            if field.field_type != "TEXT" or not config.get("calculated"):
                raise MaterializationError("DERIVED требует настроенного вычисляемого текстового поля")
            fields = self.session.query(QuestionnaireField).filter_by(schema_id=schema.id).all()
            try:
                TextExpressionEngine.validate(config.get("text_expression"), fields)
            except CalculationError as exc:
                raise MaterializationError(f"Некорректная текстовая формула: {exc}") from exc
        locator = adapter.get_structural_locator(safe_source(version), selection)
        for binding in self.session.query(DocumentBinding).filter_by(template_version_id=version.id, status="ACTIVE"):
            existing = binding.locator or {}
            if existing.get("part") == locator.get("part") and existing.get("node") == locator.get("node"):
                if _overlaps(existing.get("text_start", 0), existing.get("text_end", 0), locator["text_start"], locator["text_end"]):
                    if candidate.get("materialized_id") == binding.id:
                        return field, locator, binding
                    raise MaterializationError("Диапазон пересекается с существующей привязкой")
        region_nodes = {
            node for region in self.session.query(ConditionalDocumentRegion).filter_by(template_version_id=version.id, status="ACTIVE")
            for node in (region.locator or {}).get("nodes", [(region.locator or {}).get("node")]) if node
        }
        if locator["node"] in region_nodes:
            raise MaterializationError("Узел уже используется условной областью")
        for group in self.session.query(RepeatingGroup).filter_by(template_version_id=version.id, status="ACTIVE"):
            if (group.locator or {}).get("table") == selection["node"].get("metadata", {}).get("table") and (group.locator or {}).get("row") == selection["node"].get("metadata", {}).get("row"):
                raise MaterializationError("Узел уже используется повторяемой группой")
        return field, locator, None

    def materialize(self, version: TemplateVersion, candidate_id: str, *, manual: bool = False) -> dict:
        metadata, candidates, candidate = self._candidate(version, candidate_id)
        field, locator, existing = self._validate(version, candidate, manual)
        if existing is None:
            binding = DocumentBinding(
                template_version_id=version.id, field_id=field.id, binding_key=field.field_key,
                occurrence_index=int(str(candidate["occurrence_id"])[1:]), locator=locator,
                selected_text=candidate.get("selected_text") or "", format_config={}, created_by_id=self.actor.id,
            )
            self.session.add(binding); self.session.flush()
        else:
            binding = existing
        candidate.update({
            "status": "CONFIGURED", "materialized_type": "DOCUMENT_BINDING",
            "materialized_id": binding.id, "materialization_source": "MANUAL_REVIEW" if manual else "AUTO_EXACT",
            "materialized_at": datetime.now(timezone.utc).isoformat(),
        })
        metadata["binding_candidates"] = candidates; version.preview_metadata = metadata
        self.session.add(TemplateAuditEvent(
            template_id=version.template_id, template_version_id=version.id, actor_id=self.actor.id,
            event_type="BINDING_CANDIDATE_MATERIALIZED",
            metadata_json={"candidate_id": candidate_id, "binding_id": binding.id, "source": candidate["materialization_source"], "operation": candidate["operation"], "locator": locator["node"]},
        ))
        self.session.commit()
        return candidate

    def rollback(self, version: TemplateVersion, candidate_id: str) -> dict:
        metadata, candidates, candidate = self._candidate(version, candidate_id)
        if version.status != "DRAFT":
            raise MaterializationError("Откат разрешён только в черновой версии")
        binding_id = candidate.get("materialized_id")
        binding = self.session.get(DocumentBinding, binding_id) if binding_id else None
        if binding and binding.template_version_id == version.id:
            self.session.delete(binding)
        candidate.pop("materialized_id", None); candidate.pop("materialized_type", None)
        candidate.pop("materialization_source", None); candidate.pop("materialized_at", None)
        candidate["status"] = "CANDIDATE_FOUND" if candidate.get("selected_text") else "REVIEW_REQUIRED"
        metadata["binding_candidates"] = candidates; version.preview_metadata = metadata
        self.session.add(TemplateAuditEvent(template_id=version.template_id, template_version_id=version.id, actor_id=self.actor.id, event_type="BINDING_CANDIDATE_ROLLED_BACK", metadata_json={"candidate_id": candidate_id}))
        self.session.commit()
        return candidate
