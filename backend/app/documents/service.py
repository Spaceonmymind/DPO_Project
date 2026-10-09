from pathlib import Path
from typing import Any
import re

import fitz
import openpyxl
import xlrd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

LABELS={'subject':'Предмет закупки','purpose':'Цель закупки','scope':'Объём работ / поставки','functional_requirements':'Функциональные требования','nonfunctional_requirements':'Нефункциональные требования','deadline':'Срок выполнения','acceptance_criteria':'Критерии приёмки','other_conditions':'Дополнительные условия'}
ITEM_HEADERS={'name':['наименование','товар','позиция','предмет'],'description':['описание'],'quantity':['количество','кол-во','количество, шт'],'unit':['единица','ед. изм','единица измерения'],'unit_price':['цена за единицу','цена'],'total_price':['стоимость','сумма'],'characteristics':['характеристики','технические характеристики','требования'],'delivery_term':['срок поставки'],'warranty':['гарантия'],'delivery_address':['адрес поставки'],'references':['ссылка','ссылки'],'notes':['примечание','комментарий']}


def _text(value: Any) -> str:
    if value is None: return ''
    if isinstance(value,float) and value.is_integer(): return str(int(value))
    return str(value).strip()


class SpreadsheetSpecificationParser:
    def parse(self, path: Path, ext: str) -> dict:
        sheets = self._xlsx(path) if ext == '.xlsx' else self._xls(path)
        if not sheets: raise ValueError('В книге нет данных')
        result={'subject':'','purpose':'','items':[],'common_requirements':[],'blocking_requirements':[],'additional_conditions':[],'delivery_terms':'','total_amount':None,'source_document':path.name,'sheets':[name for name,_ in sheets]}
        for sheet_name, rows in sheets:
            self._parse_rows(rows,result,sheet_name)
        if not result['subject'] and result['items']: result['subject']='; '.join(item['name'] for item in result['items'][:3] if item.get('name'))
        if result['items'] and not result.get('scope'): result['scope']='; '.join(f"{item.get('name','Позиция')} — {item.get('quantity','')} {item.get('unit','')}".strip() for item in result['items'])
        result['functional_requirements']='; '.join(result['common_requirements'])
        result['other_conditions']='; '.join(result['additional_conditions'])
        return result

    def _xlsx(self,path):
        workbook=openpyxl.load_workbook(path,data_only=False,read_only=False)
        sheets=[]
        for ws in workbook.worksheets:
            merged={str(cell):_text(ws.cell(rng.min_row,rng.min_col).value) for rng in ws.merged_cells.ranges for row in ws[rng.coord] for cell in row}
            rows=[]
            for row in ws.iter_rows():
                values=[]
                for cell in row:
                    value=merged.get(cell.coordinate,_text(cell.value)); link=cell.hyperlink.target if cell.hyperlink else ''
                    values.append({'value':value,'link':link})
                if any(cell['value'] or cell['link'] for cell in values): rows.append(values)
            sheets.append((ws.title,rows))
        return sheets

    def _xls(self,path):
        workbook=xlrd.open_workbook(path)
        return [(sheet.name,[[{'value':_text(sheet.cell_value(r,c)),'link':''} for c in range(sheet.ncols)] for r in range(sheet.nrows) if any(_text(sheet.cell_value(r,c)) for c in range(sheet.ncols))]) for sheet in workbook.sheets()]

    def _parse_rows(self,rows,result,sheet_name):
        values=[[cell['value'] for cell in row] for row in rows]
        header_index=None; mapping={}
        for index,row in enumerate(values):
            normalized=[re.sub(r'\s+',' ',value.lower()).strip(' .:') for value in row]
            candidate={field:column for field,aliases in ITEM_HEADERS.items() for column,value in enumerate(normalized) if value in aliases or any(alias in value for alias in aliases)}
            if 'name' in candidate and ('quantity' in candidate or 'characteristics' in candidate): header_index=index; mapping=candidate; break
        if header_index is not None:
            for row_index,row in enumerate(rows[header_index+1:],header_index+2):
                item={field:(row[column]['link'] or row[column]['value']) if column<len(row) else '' for field,column in mapping.items()}
                if not item.get('name') or item['name'].lower().startswith(('итого','всего')): continue
                item['source_sheet']=sheet_name; item['source_row']=row_index
                for number_field in ('quantity','unit_price','total_price'):
                    raw=item.get(number_field,'')
                    try: item[number_field]=float(str(raw).replace(' ','').replace(',','.'))
                    except (ValueError,TypeError): pass
                result['items'].append(item)
        current_section=None
        section_targets={'общие требования':'common_requirements','блокирующие требования':'blocking_requirements','дополнительные условия':'additional_conditions'}
        for row in rows:
            nonempty=[cell for cell in row if cell['value']]
            line=' '.join(cell['value'] for cell in nonempty)
            low=line.lower()
            is_section_heading=low.strip(' :') in {'общие требования','блокирующие требования','дополнительные условия'}
            if is_section_heading:
                current_section=section_targets[low.strip(' :')]
                continue
            if len(nonempty)>=2:
                key=nonempty[0]['value'].lower().strip(' .:'); value=' '.join(cell['value'] for cell in nonempty[1:])
                if 'предмет закупки' in key: result['subject']=value
                elif 'цель закупки' in key: result['purpose']=value
                elif 'срок поставки' in key or 'срок выполнения' in key: result['delivery_terms']=value; result['deadline']=value
                elif 'адрес поставки' in key: result['delivery_address']=value
            if current_section and line.lstrip().startswith(('•','-')):
                result[current_section].append(line.lstrip('•- ').strip())
            else:
                if any(word in low for word in ('блокирующ','не допускается','обязательно отсутствие')): result['blocking_requirements'].append(line)
                elif any(word in low for word in ('обязательн','общие требования')) and line not in result['common_requirements']: result['common_requirements'].append(line)
                elif any(word in low for word in ('дополнительные условия','примечание')) and line not in result['additional_conditions']: result['additional_conditions'].append(line)


def extract(path:Path,ext:str):
    if ext=='.docx': return '\n'.join([*(p.text for p in Document(path).paragraphs),*( ' | '.join(cell.text for cell in row.cells) for table in Document(path).tables for row in table.rows)])
    if ext=='.pdf': return ''.join(page.get_text() for page in fitz.open(path))
    if ext in {'.xlsx','.xls'}:
        parsed=SpreadsheetSpecificationParser().parse(path,ext)
        lines=[parsed.get('subject','')]
        lines.extend(' | '.join(_text(item.get(key,'')) for key in ('name','quantity','unit','characteristics','warranty','delivery_address')) for item in parsed['items'])
        lines.extend(parsed['common_requirements']+parsed['blocking_requirements']+parsed['additional_conditions'])
        return '\n'.join(line for line in lines if line)
    raise ValueError('Unsupported document type')


def _base(document):
    section=document.sections[0]; section.top_margin=Cm(2); section.bottom_margin=Cm(2); section.left_margin=Cm(2.5); section.right_margin=Cm(2)
    style=document.styles['Normal']; style.font.name='Arial'; style.font.size=Pt(10); style.paragraph_format.space_after=Pt(6)


def normalized_sections(data):
    return [(LABELS[key],data.get(key)) for key in LABELS if data.get(key)]


def item_details(item):
    parts=[]
    for key,label in [('description','Описание'),('characteristics','Характеристики'),('delivery_term','Срок'),('warranty','Гарантия'),('delivery_address','Адрес'),('references','Ссылка'),('notes','Примечание')]:
        if item.get(key): parts.append(f'{label}: {item[key]}')
    return '\n'.join(parts)


def generate_spec(data,path):
    document=Document(); _base(document); title=document.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=title.add_run('ТЕХНИЧЕСКОЕ ЗАДАНИЕ'); run.bold=True; run.font.size=Pt(16)
    for index,(label,value) in enumerate(normalized_sections(data),1): document.add_paragraph(f'{index}. {label}').runs[0].bold=True; document.add_paragraph(str(value))
    items=data.get('items',[])
    if items:
        document.add_paragraph('Позиции').runs[0].bold=True; table=document.add_table(rows=1,cols=6); table.style='Table Grid'
        for cell,value in zip(table.rows[0].cells,['№','Наименование','Количество','Ед. изм.','Стоимость','Сведения и требования']): cell.text=value
        for index,item in enumerate(items,1):
            cells=table.add_row().cells
            for cell,value in zip(cells,[index,item.get('name',''),item.get('quantity',''),item.get('unit',''),item.get('total_price',''),item_details(item)]): cell.text=str(value or '')
    for key,label in [('common_requirements','Общие требования'),('blocking_requirements','Блокирующие требования'),('additional_conditions','Дополнительные условия')]:
        if data.get(key): document.add_paragraph(label).runs[0].bold=True; [document.add_paragraph(str(value),style='List Bullet') for value in data[key]]
    document.save(path)


def generate_xlsx(data,path):
    workbook=openpyxl.Workbook(); ws=workbook.active; ws.title='Техническое задание'; ws.sheet_view.showGridLines=False
    ws.merge_cells('A1:F1'); ws['A1']='ТЕХНИЧЕСКОЕ ЗАДАНИЕ'; ws['A1'].font=Font(size=16,bold=True,color='FFFFFF'); ws['A1'].fill=PatternFill('solid',fgColor='008B76'); ws['A1'].alignment=Alignment(horizontal='center')
    row=3
    for label,value in normalized_sections(data): ws.cell(row,1,label).font=Font(bold=True); ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=6); ws.cell(row,2,str(value)).alignment=Alignment(wrap_text=True); row+=1
    row+=1; headers=['№','Наименование','Описание','Количество','Ед. изм.','Цена за единицу','Стоимость','Характеристики','Срок поставки','Гарантия','Адрес поставки','Ссылка','Примечание']
    for column,value in enumerate(headers,1): cell=ws.cell(row,column,value); cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='008B76'); cell.alignment=Alignment(wrap_text=True)
    for index,item in enumerate(data.get('items',[]),1):
        row+=1; values=[index,item.get('name',''),item.get('description',''),item.get('quantity',''),item.get('unit',''),item.get('unit_price',''),item.get('total_price',''),item.get('characteristics',''),item.get('delivery_term',''),item.get('warranty',''),item.get('delivery_address',''),item.get('references',''),item.get('notes','')]
        for column,value in enumerate(values,1): ws.cell(row,column,value).alignment=Alignment(vertical='top',wrap_text=True)
        for column in (4,6,7): ws.cell(row,column).number_format='#,##0.00'
        if item.get('references'): ws.cell(row,12).hyperlink=str(item['references']); ws.cell(row,12).style='Hyperlink'
    for key,label in [('common_requirements','Общие требования'),('blocking_requirements','Блокирующие требования'),('additional_conditions','Дополнительные условия')]:
        if data.get(key):
            row+=2; ws.cell(row,1,label).font=Font(bold=True,color='008B76')
            for value in data[key]: row+=1; ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=13); ws.cell(row,1,'• '+str(value)).alignment=Alignment(wrap_text=True)
    widths=[7,30,28,12,12,16,16,45,18,18,28,32,28]
    for column,width in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width=width
    ws.freeze_panes='A4'; workbook.save(path)


def generate_pdf(data,path):
    font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'; pdfmetrics.registerFont(TTFont('DejaVu',font_path))
    styles=getSampleStyleSheet()
    for style in styles.byName.values(): style.fontName='DejaVu'
    story=[Paragraph('ТЕХНИЧЕСКОЕ ЗАДАНИЕ',styles['Title']),Spacer(1,12)]
    for label,value in normalized_sections(data): story.extend([Paragraph(f'<b>{label}</b>',styles['BodyText']),Paragraph(str(value),styles['BodyText']),Spacer(1,6)])
    items=data.get('items',[])
    if items:
        rows=[['№','Наименование','Кол-во','Ед.','Стоимость','Сведения и требования']]+[[str(index),str(item.get('name','')),str(item.get('quantity','')),str(item.get('unit','')),str(item.get('total_price','')),item_details(item)] for index,item in enumerate(items,1)]
        table=Table(rows,colWidths=[22,100,42,32,55,220],repeatRows=1); table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'DejaVu'),('FONTSIZE',(0,0),(-1,-1),7),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#008B76')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.4,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP')])) ; story.extend([Spacer(1,10),table])
    SimpleDocTemplate(str(path),pagesize=A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36).build(story)


def render_specification(data,path,fmt):
    {'docx':generate_spec,'xlsx':generate_xlsx,'pdf':generate_pdf}[fmt](data,path)


def generate_contract(name,path):
    document=Document(); _base(document); title=document.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=title.add_run(name.upper()); run.bold=True; run.font.size=Pt(15)
    document.add_paragraph('г. Москва                                                   «__» ______ 20__ г.')
    for number,title in enumerate(['Предмет договора','Права и обязанности сторон','Стоимость и порядок расчётов','Порядок сдачи и приёмки','Ответственность сторон','Срок действия договора','Реквизиты сторон'],1): document.add_paragraph(f'{number}. {title}').runs[0].bold=True; document.add_paragraph('Условия определяются сторонами в соответствии с предметом договора.')
    document.save(path)
